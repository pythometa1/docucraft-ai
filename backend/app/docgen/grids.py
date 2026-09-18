"""A computed table, as the document will print it.

Shared by every module that renders a table from data rather than letting a
model type one: CMC's specifications and stability summaries, and the Safety
module's case tabulations and line listings. Both rest on the same rule --
the model writes the prose around a table and never the table -- and both
need the same four things from this file:

* one string for an absent value, so a reader scanning a column sees the gap;
* a check that the grid is rectangular, run where the builder that made it can
  still be named, rather than when `emit` refuses the whole document;
* the grid as blueprint blocks, from the same cells the caller asserts on, so
  the printed table and the tested one cannot drift;
* the grid as a ruled `.docx`, because `emit` writes no table style and a
  table without ruling is a page of floating numbers.

Extracted from `app.cmc.tables` when the second module needed it, rather than
copied. `app.cmc.tables` keeps its own names for these as aliases, so nothing
that imported them from there changes.
"""

import docx

from app.generation.reproducibility import normalise_docx
from app.templates import blueprint as bp
from app.templates.emit_docx import emit

#: What an absent value prints as. One character, the same in every table, so
#: a reader sees the gap -- and the caller records which cell it was, because a
#: dash in a printed document names nothing.
HOLE = "-"

#: The ruling a rendered table is given.
TABLE_STYLE = "Table Grid"


class TableError(Exception):
    """A table cannot be rendered as asked."""


class TableUnavailable(TableError):
    """There is no data for this table.

    Distinct from a table with holes in it. A hole is rendered, because the
    reviewer needs to see which cell is empty; nothing at all is refused,
    because a grid of dashes under a heading reads as a finding rather than as
    an absence of data.
    """


class UnknownTable(TableError):
    """No builder is registered under this key.

    Kept apart from `TableUnavailable`. `[TABLE: spec_tabel]` is a typo, and
    reporting it as "no data" sends somebody to look for a row that was never
    the problem.
    """


class RaggedTable(TableError):
    """A builder produced rows of unequal width.

    `emit` refuses a ragged grid -- Word has no representation for one -- and
    it refuses it after the whole document has been assembled. Caught at the
    builder, the message can name the table that is wrong.
    """


def text(value) -> str:
    """A stored value as a cell, or a hole where there is none.

    Not a transformation of a value: `None` and `""` are not values.
    """
    if value is None:
        return HOLE
    rendered = str(value)
    return rendered if rendered.strip() else HOLE


def cell(value: str, *, bold: bool = False) -> list:
    """One table cell: a paragraph holding one static segment.

    Static because a placeholder is a fill slot the compiler would try to write
    into, and a rendered value is already the answer.
    """
    return [bp.paragraph([bp.segment("static", value, bold=bold)])]


def heading(value: str) -> dict:
    return bp.paragraph([bp.segment("static", value, bold=True)])


def grid(columns, rows) -> dict:
    header = [cell(text(name), bold=True) for name in columns]
    return bp.table([header] + [[cell(value) for value in row] for row in rows])


def check_rectangular(key: str, columns, rows, *, where: str = "") -> None:
    """Refuse a ragged grid, naming the table and the row."""
    for index, row in enumerate(rows):
        if len(row) != len(columns):
            raise RaggedTable(
                f"{key}{where}: row {index} has {len(row)} cell(s) and the table has "
                f"{len(columns)} column(s); Word has no representation for a ragged "
                "grid")


def write_docx(blocks, output_path: str) -> str:
    """Write blocks to a ruled, byte-reproducible `.docx`.

    Ruling is applied after `emit`, not instead of it: `emit` round-trips the
    file it wrote and proves the text in it is the text in the body, and a
    style applied by a second renderer would be a second renderer. The archive
    is normalised on the way out, because emitting the same data twice has to
    give the same bytes.
    """
    body = bp.normalise_body({"blocks": list(blocks), "sect_pr_from": None})
    emit(body, output_path)
    document = docx.Document(output_path)
    for table in document.tables:
        table.style = TABLE_STYLE
    document.save(output_path)
    return normalise_docx(output_path)
