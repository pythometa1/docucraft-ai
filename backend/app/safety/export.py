"""§12: a periodic safety report, written from what was approved.

Three things this module will not do.

**It never exports the original case text.** Nothing here names
`pv_case_originals`. Every word comes from a section draft -- masked text,
scanned on every save -- or from a table builder reading structured fields.
The finished file is scanned again anyway, every part of it including deleted
text in a tracked-changes copy, and a hit deletes the file and refuses the
export. A leak found after the file is written is still a leak found before
anybody downloads it.

**It never produces a submission payload.** No E2B XML, no CIOMS form, no
gateway message. Transmission belongs to the safety system of record, and a
module that half-did it would produce something that looks transmittable.

**It never resolves a table at draft time.** `[TABLE: key]` is resolved here,
from the case store as it is now, by the same builders the Tables screen and QC
use. An event confirmed after a section was written appears in the export
without the section being regenerated.

Assembly is split from writing. `assemble` is pure: section text and resolved
tables in, blocks out, with a parallel list of revision marks when a
tracked-changes copy is wanted. `write` turns blocks into a `.docx` and adds
what the blueprint cannot express -- header, page numbers, table of contents,
watermark, revision marks -- after `emit` has verified the text.
"""

import difflib
import re
import zipfile
from dataclasses import dataclass, field

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from lxml import etree

from app.docgen.assembly import paragraph_texts, split_on_tables
from app.safety import deident
from app.templates import blueprint as bp

#: What an export may carry after the report body, and the tables each is
#: built from. The same builders as the body: an appendix that counted its own
#: way would be a second answer to the same question.
APPENDICES = {
    "line_listings": ("Line listings", ("line_listing_sar",)),
    "tabulations": ("Summary tabulations",
                    ("summary_tab_soc_pt", "summary_tab_trials", "exposure_table")),
    "rsi": ("Reference safety information", ("rsi_listed_terms",)),
    "study_inventory": ("Study inventory", ("study_inventory",)),
    "signal_log": ("Signal log", ("signal_overview",)),
}

#: "strip" removes `[S#]` markers; "keep" leaves them for an internal review
#: copy, where a reader wants to check a sentence against its source.
CITATION_MODES = ("strip", "keep")

#: A citation, and the spaces in front of it -- never the newline. A pattern
#: that ate the line break pulled the next line up onto a `[TABLE: key]`
#: marker, which then stopped matching and was printed as text (CMC, G8).
_CITATION_RE = re.compile(r"[ \t]*\[S\d+(?:,[^\]]*)?\]")

#: The paragraph `write` turns into a table-of-contents field. Word fills the
#: field when the document is opened; until then this is what shows.
TOC_PLACEHOLDER = "Update this field to build the table of contents."

#: The author Word shows on a revision. The changes were computed against the
#: baseline report, not typed by a person, and the name says so.
REVISION_AUTHOR = "Comparison with the previous report"


class ExportError(Exception):
    """The report cannot be written as asked."""


@dataclass
class SectionText:
    """One section as the export sees it."""

    code: str
    title: str
    is_container: bool
    content: str = ""
    #: The previous report's text for this section, for tracked changes.
    #: None when the section is new in this report.
    baseline: str | None = None


@dataclass
class Assembled:
    blocks: list = field(default_factory=list)
    #: One per block: None, "ins", "del", or a list of (text, None|"ins"|"del")
    #: word runs for a paragraph that changed.
    marks: list = field(default_factory=list)
    #: Table keys the body asked for, in order.
    table_keys: list = field(default_factory=list)

    def add(self, block, mark=None) -> None:
        self.blocks.append(block)
        self.marks.append(mark)


# ----------------------------------------------------------------- the text

def heading_level(section_code: str) -> int:
    """Depth from the numbering: 7.3 sits under 7, II.SVIII under II."""
    return min(1 + (section_code or "").count("."), 4)


def body_of(content: str, section: SectionText, *, citations: str = "strip") -> str:
    """A draft without the heading it opens with, and with its citations
    handled.

    Drafts open with "<code> <title>" (the drafting prompt's rule 11). The
    export writes that heading itself, as a real Heading style the table of
    contents can find, so the draft's copy of it would print twice.
    """
    lines = (content or "").splitlines()
    index = 0
    while index < len(lines) and not lines[index].strip():
        index += 1
    if index < len(lines) and _is_heading(lines[index], section):
        lines = lines[index + 1:]
    text = "\n".join(lines)
    if citations == "strip":
        text = _CITATION_RE.sub("", text)
    return text.strip("\n")


def _is_heading(line: str, section: SectionText) -> bool:
    """"4 Case Series Review" is the heading; "4 cases were reviewed." is the
    first sentence, and dropping it would delete a finding from the report.
    The title's first word must match as a whole word, and a heading does not
    end in a full stop."""
    stripped = line.strip().lstrip("#").strip()
    if not stripped.startswith(f"{section.code} ") or stripped.endswith("."):
        return False
    rest = stripped[len(section.code):].strip()
    first = (section.title or "").split()
    return bool(first) and re.match(rf"{re.escape(first[0])}\b", rest,
                                    re.IGNORECASE) is not None


def items_of(body: str) -> list:
    """The body as ("text", paragraph) and ("table", key) items, in order."""
    out = []
    for kind, value in split_on_tables(body):
        if kind == "table":
            out.append(("table", value))
        else:
            out.extend(("text", paragraph) for paragraph in paragraph_texts(value))
    return out


def _item_key(item) -> str:
    kind, value = item
    return f"\x00table:{value}" if kind == "table" else " ".join(value.split())


def word_diff(old: str, new: str) -> list:
    """(text, None|"ins"|"del") runs turning `old` into `new`, word by word."""
    a = re.findall(r"\S+|\s+", old)
    b = re.findall(r"\S+|\s+", new)
    runs = []
    for op, a0, a1, b0, b1 in difflib.SequenceMatcher(
            a=a, b=b, autojunk=False).get_opcodes():
        if op == "equal":
            runs.append(("".join(b[b0:b1]), None))
            continue
        if a1 > a0:
            runs.append(("".join(a[a0:a1]), "del"))
        if b1 > b0:
            runs.append(("".join(b[b0:b1]), "ins"))
    return runs


def diff_items(old: list, new: list) -> list:
    """This report's items with revision marks against the previous report's.

    Returns (kind, value, mark). A paragraph unchanged but for whitespace is
    unchanged. A table is never marked: it is computed from this interval's
    data every time, so "the table changed" is always true and says nothing --
    the tracked changes are for what a person wrote.
    """
    matcher = difflib.SequenceMatcher(
        a=[_item_key(i) for i in old], b=[_item_key(i) for i in new], autojunk=False)
    out = []
    for op, a0, a1, b0, b1 in matcher.get_opcodes():
        if op == "equal":
            out.extend((kind, value, None) for kind, value in new[b0:b1])
        elif op == "insert":
            out.extend((kind, value, None if kind == "table" else "ins")
                       for kind, value in new[b0:b1])
        elif op == "delete":
            out.extend((kind, value, "del") for kind, value in old[a0:a1]
                       if kind == "text")
        else:
            olds = [value for kind, value in old[a0:a1] if kind == "text"]
            for kind, value in new[b0:b1]:
                if kind == "table":
                    out.append((kind, value, None))
                elif olds:
                    out.append((kind, value, word_diff(olds.pop(0), value)))
                else:
                    out.append((kind, value, "ins"))
            out.extend(("text", value, "del") for value in olds)
    return out


# ------------------------------------------------------------------ assembly

def _paragraph(text: str, style: str | None = None, **extra) -> dict:
    return bp.paragraph([bp.segment("static", text, **extra)], style=style)


def assemble(sections: list, tables: dict, *, front: list, title: str,
             appendices: list = (), citations: str = "strip",
             tracked: bool = False) -> Assembled:
    """The whole report as blocks.

    `tables` maps a table key to its blocks, or to a string saying why there is
    no table. `appendices` is a list of (heading, [table keys]).
    """
    out = Assembled()
    out.add(_paragraph(title, "Title"))
    for line in front:
        out.add(_paragraph(line))
    out.add(_paragraph("Contents", bold=True))
    out.add(_paragraph(TOC_PLACEHOLDER))

    for section in sections:
        level = heading_level(section.code)
        out.add(_paragraph(f"{section.code} {section.title}", f"Heading {level}"))
        if section.is_container and not (section.content or "").strip():
            continue
        new = items_of(body_of(section.content, section, citations=citations))
        if tracked:
            old = (items_of(body_of(section.baseline, section, citations=citations))
                   if section.baseline is not None else [])
            marked = diff_items(old, new)
        else:
            marked = [(kind, value, None) for kind, value in new]
        for kind, value, mark in marked:
            if kind == "table":
                out.table_keys.append(value)
                _add_table(out, tables, value)
            elif mark == "del":
                out.add(_paragraph(value), "del")
            else:
                out.add(_paragraph(value), mark)

    if appendices:
        out.add(_paragraph("Appendices", "Heading 1"))
        for number, (heading, keys) in enumerate(appendices, start=1):
            out.add(_paragraph(f"Appendix {number}: {heading}", "Heading 2"))
            for key in keys:
                _add_table(out, tables, key, appendix=True)
    return out


def _add_table(out: Assembled, tables: dict, key: str, *, appendix: bool = False):
    built = tables.get(key)
    if isinstance(built, list):
        for block in built:
            out.add(block)
        return
    if appendix:
        # An empty appendix says so. "No signal was open" is information; a
        # heading followed by nothing reads as a table that failed to print.
        out.add(_paragraph(f"No {key.replace('_', ' ')}: {built or 'no data'}.",
                           italic=True))
        return
    # QC refuses an export with an unresolved body marker, so reaching here is
    # a caller that skipped the gate. Refuse rather than print a hole.
    raise ExportError(f"[TABLE: {key}] has no table behind it: {built or 'not built'}")


# ------------------------------------------------------------------- writing

def write(assembled: Assembled, output_path: str, *, header: str,
          draft: bool = False, revision_date: str | None = None) -> str:
    """Blocks to a finished `.docx`.

    `emit` writes and verifies the text first. Everything after it -- ruling,
    header and footer, the contents field, the watermark, revision marks -- is
    structure the blueprint has no vocabulary for, added to the verified file.
    """
    from app.generation.reproducibility import normalise_docx
    from app.templates.emit_docx import emit

    body = bp.normalise_body({"blocks": list(assembled.blocks), "sect_pr_from": None})
    if len(body["blocks"]) != len(assembled.blocks):
        raise ExportError("the document changed shape while it was normalised")
    emit(body, output_path)

    document = docx.Document(output_path)
    for table in document.tables:
        table.style = "Table Grid"
    children = [child for child in document.element.body
                if child.tag in (qn("w:p"), qn("w:tbl"))]
    if len(children) != len(assembled.marks):
        raise ExportError(
            f"the document has {len(children)} block(s) and the assembly "
            f"{len(assembled.marks)}; revision marks would land on the wrong text")
    ids = iter(range(1, 1_000_000))
    stamp = revision_date or "2000-01-01T00:00:00Z"
    for element, mark in zip(children, assembled.marks):
        if mark is None or element.tag != qn("w:p"):
            continue
        _mark_paragraph(element, mark, ids, stamp)
    for paragraph in document.paragraphs:
        if paragraph.text == TOC_PLACEHOLDER:
            _replace_with_field(paragraph._p, 'TOC \\o "1-4" \\h \\z \\u', TOC_PLACEHOLDER)
            break
    _header_and_footer(document, header, draft)
    # After the header, because each landscape section copies the body's
    # section properties -- header and footer references included -- and a
    # copy taken earlier would point at no header at all.
    _lay_out_tables(document)
    settings = document.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update = settings.makeelement(qn("w:updateFields"), {qn("w:val"): "true"})
        settings.append(update)
    document.save(output_path)
    return normalise_docx(output_path)


def _revision(tag: str, ids, stamp: str):
    return etree.Element(qn(tag), {qn("w:id"): str(next(ids)),
                                   qn("w:author"): REVISION_AUTHOR,
                                   qn("w:date"): stamp})


def _plain_run(text: str, *, deleted: bool = False):
    run = etree.Element(qn("w:r"))
    node = etree.SubElement(run, qn("w:delText" if deleted else "w:t"))
    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    node.text = text
    return run


def _mark_paragraph_mark(paragraph_el, tag: str, ids, stamp: str) -> None:
    """Mark the paragraph's own end as inserted or deleted, so accepting or
    rejecting the change removes the paragraph rather than leaving it empty."""
    ppr = paragraph_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("w:pPr"))
        paragraph_el.insert(0, ppr)
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = etree.SubElement(ppr, qn("w:rPr"))
    rpr.append(_revision(tag, ids, stamp))


def _mark_paragraph(paragraph_el, mark, ids, stamp: str) -> None:
    runs = [child for child in paragraph_el if child.tag == qn("w:r")]
    if mark in ("ins", "del"):
        tag = "w:ins" if mark == "ins" else "w:del"
        for run in runs:
            if mark == "del":
                for node in run.findall(qn("w:t")):
                    node.tag = qn("w:delText")
            wrapper = _revision(tag, ids, stamp)
            run.addprevious(wrapper)
            wrapper.append(run)
        _mark_paragraph_mark(paragraph_el, tag, ids, stamp)
        return
    for run in runs:
        paragraph_el.remove(run)
    for text, kind in mark:
        if kind is None:
            paragraph_el.append(_plain_run(text))
            continue
        wrapper = _revision("w:ins" if kind == "ins" else "w:del", ids, stamp)
        wrapper.append(_plain_run(text, deleted=kind == "del"))
        paragraph_el.append(wrapper)


def _replace_with_field(paragraph_el, instruction: str, placeholder: str) -> None:
    for child in list(paragraph_el):
        if child.tag != qn("w:pPr"):
            paragraph_el.remove(child)
    field_el = etree.SubElement(paragraph_el, qn("w:fldSimple"),
                                {qn("w:instr"): f" {instruction} "})
    field_el.append(_plain_run(placeholder))


def _field_run(paragraph_el, instruction: str, shown: str) -> None:
    field_el = etree.SubElement(paragraph_el, qn("w:fldSimple"),
                                {qn("w:instr"): f" {instruction} "})
    field_el.append(_plain_run(shown))


_VML = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:o="urn:schemas-microsoft-com:office:office"')

#: Word's own text-watermark shape: a VML text path, rotated, behind the text.
_WATERMARK = (
    f'<w:r {_VML}><w:pict>'
    '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
    'path="m@7,l@8,m@5,21600l@6,21600e">'
    '<v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/>'
    '<v:f eqn="sum 21600 0 @1"/><v:f eqn="sum 0 0 @2"/><v:f eqn="sum 21600 0 @3"/>'
    '<v:f eqn="if @0 @3 0"/><v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/>'
    '<v:f eqn="if @0 @4 21600"/><v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/>'
    '<v:f eqn="mid @7 @8"/><v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/></v:formulas>'
    '<v:path textpathok="t" o:connecttype="custom" '
    'o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" '
    'o:connectangles="270,180,90,0"/>'
    '<v:textpath on="t" fitshape="t"/>'
    '<o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>'
    '<v:shape id="PowerPlusWaterMarkObject" o:spid="_x0000_s2049" type="#_x0000_t136" '
    'style="position:absolute;margin-left:0;margin-top:0;width:468pt;height:117pt;'
    'rotation:315;z-index:-251655168;mso-position-horizontal:center;'
    'mso-position-horizontal-relative:margin;mso-position-vertical:center;'
    'mso-position-vertical-relative:margin" o:allowincell="f" fillcolor="silver" '
    'stroked="f"><v:fill opacity=".5"/>'
    '<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="DRAFT"/>'
    '</v:shape></w:pict></w:r>')


def _header_and_footer(document, header: str, draft: bool) -> None:
    section = document.sections[0]
    paragraph = section.header.paragraphs[0]
    paragraph.text = header + (" | DRAFT" if draft else "")
    if draft:
        paragraph._p.append(parse_xml(_WATERMARK))
    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer._p.append(_plain_run("Page "))
    _field_run(footer._p, "PAGE", "1")
    footer._p.append(_plain_run(" of "))
    _field_run(footer._p, "NUMPAGES", "1")


# ------------------------------------------------------------ table layout

#: At this many columns a table's text is set smaller...
SMALL_TEXT_COLUMNS = 7
#: ...and at this many it gets a landscape page. A summary tabulation has
#: fourteen columns; in portrait each is under half an inch and its headers
#: wrap a letter at a time.
LANDSCAPE_COLUMNS = 9


def _lay_out_tables(document) -> None:
    """Repeat every table's header row on each page it spans, set wide tables
    in smaller text, and give the widest their own landscape pages.

    Layout only: no text is added, removed or reordered, so what `emit`
    verified is still what the document says.
    """
    for table in document.tables:
        tbl = table._tbl
        for index, row in enumerate(tbl.findall(qn("w:tr"))):
            trpr = row.find(qn("w:trPr"))
            if trpr is None:
                trpr = etree.Element(qn("w:trPr"))
                # `w:trPr` follows `w:tblPrEx` if there is one, else it leads.
                row.insert(1 if row.find(qn("w:tblPrEx")) is not None else 0, trpr)
            # A case's row is read whole: one split across a page break reads
            # as two cases, the second with no identifier.
            if trpr.find(qn("w:cantSplit")) is None:
                etree.SubElement(trpr, qn("w:cantSplit"))
            if index == 0 and trpr.find(qn("w:tblHeader")) is None:
                etree.SubElement(trpr, qn("w:tblHeader"))
        columns = len(table.columns)
        if columns >= SMALL_TEXT_COLUMNS:
            for run in tbl.iter(qn("w:r")):
                rpr = run.find(qn("w:rPr"))
                if rpr is None:
                    rpr = etree.Element(qn("w:rPr"))
                    run.insert(0, rpr)
                etree.SubElement(rpr, qn("w:sz"), {qn("w:val"): "16"})
        if columns >= LANDSCAPE_COLUMNS and tbl.getparent() is document.element.body:
            _landscape(document, tbl)


def _landscape(document, tbl) -> None:
    """Put one table -- and the title paragraph above it -- on landscape pages.

    A Word section ends at the paragraph carrying its properties. So the
    paragraph before the title ends a portrait section, and a new empty
    paragraph after the table ends the landscape one; what follows returns to
    the body's own portrait section.
    """
    import copy

    body_sectpr = document.element.body.find(qn("w:sectPr"))
    if body_sectpr is None:  # pragma: no cover - python-docx always writes one
        return
    title = tbl.getprevious()
    start = title if title is not None and title.tag == qn("w:p") else tbl
    before = start.getprevious()
    if before is None or before.tag != qn("w:p"):
        before = etree.Element(qn("w:p"))
        start.addprevious(before)
    _attach_sectpr(before, copy.deepcopy(body_sectpr))

    wide = copy.deepcopy(body_sectpr)
    size = wide.find(qn("w:pgSz"))
    if size is not None:
        width, height = size.get(qn("w:w")), size.get(qn("w:h"))
        if width and height:
            size.set(qn("w:w"), height)
            size.set(qn("w:h"), width)
        size.set(qn("w:orient"), "landscape")
    after = etree.Element(qn("w:p"))
    tbl.addnext(after)
    _attach_sectpr(after, wide)
    _fit_width(tbl, wide)


def _fit_width(tbl, sectpr) -> None:
    """Spread the table across the landscape text width. python-docx fixed
    every column at a share of the PORTRAIT width when it made the table, and
    a landscape page holding a portrait-width table has gained nothing."""
    size, margins = sectpr.find(qn("w:pgSz")), sectpr.find(qn("w:pgMar"))
    try:
        width = (int(size.get(qn("w:w"))) - int(margins.get(qn("w:left")))
                 - int(margins.get(qn("w:right"))))
    except (AttributeError, TypeError, ValueError):  # pragma: no cover - defensive
        return
    grid = tbl.find(qn("w:tblGrid"))
    columns = grid.findall(qn("w:gridCol")) if grid is not None else []
    if not columns or width <= 0:  # pragma: no cover - a table always has a grid
        return
    share = str(width // len(columns))
    for column in columns:
        column.set(qn("w:w"), share)
    for cell_width in tbl.iter(qn("w:tcW")):
        cell_width.set(qn("w:w"), share)
        cell_width.set(qn("w:type"), "dxa")
    tblpr = tbl.find(qn("w:tblPr"))
    if tblpr is not None:
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is None:
            tblw = etree.SubElement(tblpr, qn("w:tblW"))
        tblw.set(qn("w:w"), str(width))
        tblw.set(qn("w:type"), "dxa")


def _attach_sectpr(paragraph_el, sectpr) -> None:
    ppr = paragraph_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("w:pPr"))
        paragraph_el.insert(0, ppr)
    existing = ppr.find(qn("w:sectPr"))
    if existing is not None:
        ppr.remove(existing)
    # `w:sectPr` belongs after the paragraph's other properties but before a
    # revision record; there is none here, so it goes last.
    ppr.append(sectpr)


# ---------------------------------------------------------- the leakage scan

#: Every part of a `.docx` a reader can see text in.
_TEXT_PARTS = re.compile(r"^word/(document|header\d*|footer\d*|footnotes|endnotes|"
                         r"comments)\.xml$")


def document_text(path: str) -> str:
    """Every word in the file, deleted revisions included.

    Deleted text in a tracked-changes copy is still in the file, and one click
    on "Reject all" puts it back on the page; a scan that read only the visible
    text would pass a name sitting in a deletion.
    """
    lines = []
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if not _TEXT_PARTS.match(name):
                continue
            root = etree.fromstring(archive.read(name))
            for paragraph in root.iter(qn("w:p")):
                text = "".join(node.text or "" for node in paragraph.iter()
                               if node.tag in (qn("w:t"), qn("w:delText")))
                if text:
                    lines.append(text)
    return "\n".join(lines)


def leaks(text: str, identifiers=()) -> list:
    """Identifiers in the finished document: the confident patterns, and every
    string a person confirmed as an identifier in the review queue."""
    found = [{"identifier_type": hit.identifier_type, "text": hit.text}
             for hit in deident.scan(text)]
    found.extend({"identifier_type": "confirmed_identifier", "text": value}
                 for value in deident.confirmed_in(text, identifiers))
    return found
