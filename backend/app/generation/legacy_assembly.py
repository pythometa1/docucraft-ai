"""Document assembly (spec §11). Simplification vs. the full spec: we operate at
the python-docx paragraph API rather than hand-rolled raw OOXML anchors, and jinja
substitution replaces a paragraph's runs wholesale rather than preserving every
individual run's formatting -- both are documented, reasonable simplifications for
an MVP that still produces a real, openable, correctly-ordered DOCX.
"""

import docx
from bs4 import BeautifulSoup, NavigableString, Tag


def assemble_from_docx_template(
    template_path: str,
    output_path: str,
    section_ranges: list[tuple[int, int, str]],  # (start_el, end_el, resolved_text) sorted by start_el
    jinja_replacements: dict[str, str],
) -> None:
    document = docx.Document(template_path)
    paragraphs = document.paragraphs

    # Apply jinja `{var}` substitutions paragraph-by-paragraph first (works across
    # both heading- and jinja-kind templates since jinja vars can appear anywhere).
    if jinja_replacements:
        for p in paragraphs:
            text = p.text
            replaced = text
            for var, value in jinja_replacements.items():
                replaced = replaced.replace(var, value)
            if replaced != text:
                for run in list(p.runs):
                    run.text = ""
                if p.runs:
                    p.runs[0].text = replaced
                else:
                    p.add_run(replaced)

    # Replace each mapped section's body paragraphs (start_el+1..end_el) with a
    # single resolved paragraph, preserving the heading paragraph itself. Processed
    # in reverse start_el order so index lookups into the original snapshot stay valid.
    for start_el, end_el, resolved_text in sorted(section_ranges, key=lambda r: r[0], reverse=True):
        heading_p = paragraphs[start_el]
        new_p = document.add_paragraph(resolved_text)  # appended at end of body...
        heading_p._p.addnext(new_p._p)  # ...then relocated right after the heading

        body_start = start_el + 1
        for i in range(body_start, min(end_el, len(paragraphs) - 1) + 1):
            el = paragraphs[i]._p
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)

    document.save(output_path)


_BLOCK_TAGS = {"h1": 1, "h2": 2, "h3": 3, "p": 0, "blockquote": 0}


def _add_runs(paragraph, node, bold=False, italic=False, underline=False):
    for child in node.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if text:
                run = paragraph.add_run(text)
                run.bold, run.italic, run.underline = bold, italic, underline
        elif isinstance(child, Tag):
            name = child.name.lower()
            _add_runs(
                paragraph,
                child,
                bold=bold or name in ("strong", "b"),
                italic=italic or name in ("em", "i"),
                underline=underline or name == "u",
            )


def assemble_from_html(content_html: str, output_path: str) -> None:
    soup = BeautifulSoup(content_html, "lxml")
    document = docx.Document()
    body = soup.find("body") or soup

    for node in body.find_all(["h1", "h2", "h3", "p", "ul", "ol", "blockquote"], recursive=False) or body.children:
        if not isinstance(node, Tag):
            continue
        name = node.name.lower()
        if name in ("h1", "h2", "h3"):
            p = document.add_heading(level=_BLOCK_TAGS[name])
            _add_runs(p, node)
        elif name == "blockquote":
            p = document.add_paragraph(style="Intense Quote")
            _add_runs(p, node)
        elif name in ("ul", "ol"):
            style = "List Bullet" if name == "ul" else "List Number"
            for li in node.find_all("li", recursive=False):
                p = document.add_paragraph(style=style)
                _add_runs(p, li)
        elif name == "p":
            p = document.add_paragraph()
            _add_runs(p, node)

    document.save(output_path)
