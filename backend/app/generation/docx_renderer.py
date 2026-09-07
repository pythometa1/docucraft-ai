"""The Universal Fill Engine (docs/TEMPLATE_COMPILER_RESEARCH.md §4.2) --
deterministic OOXML surgery on a *copy* of the original template, driven
entirely by a compiled manifest + a validated source record. No LLM call
happens here by design (R3): layout fidelity comes from mutating the
original file in place, never regenerating it.
"""

import copy
import re
from dataclasses import dataclass, field
from datetime import date

import docx
from lxml import etree

from app.templates.parsers.docx_prescan import W_NS, MergeField, RunSpan, _run_text, _walk_paragraphs
from app.generation.missing_policy import BLANK, DEFAULT, REMOVE_SENTENCE, field_on_missing
from app.generation.reproducibility import normalise_docx
from app.expressions.token_parser import evaluate_condition
from app.generation.value_format import DEFAULT_LOCALE, format_value

# The QA gates live in `app/qa` as named checks (architecture record §15), so a
# manifest can declare which of them block and which only warn (§6 `qa_policy`)
# instead of that being a hard-coded property of this renderer.
from app.qa.layout_integrity import findings as layout_findings
from app.qa.content_loss import findings as content_loss_findings
from app.qa.value_format_check import findings as value_format_findings
from app.qa.layout_integrity import structural_failures as _structural_failures  # noqa: F401
from app.qa.overflow_check import findings as overflow_findings
from app.qa.placeholder_check import findings as placeholder_findings
from app.qa.policy import BLOCKING, BRANCH_SELECTION, REQUIRED_VALUE_MISSING, ROW_REPEAT, resolve_policy, route
from app.qa.value_lineage_check import EQUALITY_CONDITION_RE, blocks_document  # noqa: F401
from app.qa.value_lineage_check import branch_findings, missing_value_note

PUT_INSTRUCTION_RE = re.compile(r"^\s*(please\s+)?put\b", re.IGNORECASE)

# Sentence boundary for REMOVE_SENTENCE. Deliberately conservative: it splits on
# terminal punctuation followed by whitespace, so an abbreviation mid-sentence
# ("Ltd. of Dublin") does not cut the sentence in half.
_SENTENCE_END_RE = re.compile(r"(?<=[.!?])\s+(?=[A-Z(\"'])")
# Private Use Area, so it is valid XML (NUL is not) and cannot occur in a
# real document. Both ends are marked so a stray half is still matchable.
_MARKER_OPEN, _MARKER_CLOSE = "\uE000", "\uE001"
_MARKER_RE = re.compile(f"{_MARKER_OPEN}RS\\d+{_MARKER_CLOSE}")


def _q(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


@dataclass
class FieldLineage:
    field_id: str
    value: str | None
    source: str  # "source_record" | "system" | "fallback" | "missing"
    # What the manifest said to do when this field has no value, and what the
    # engine actually did. Recorded even when the field resolved, so an audit
    # can tell "there was a value" apart from "the policy hid its absence".
    on_missing: str | None = None
    blocking: bool = False


@dataclass
class ConditionLineage:
    condition_id: str
    expression: str
    result: bool | None  # None == undecided; the document is blocked, not filtered
    reason: str = "evaluated"
    missing_fields: list = field(default_factory=list)


@dataclass
class FillResult:
    field_lineage: list = field(default_factory=list)
    condition_lineage: list = field(default_factory=list)
    # One entry per §6 TABLE_ROW the manifest declares: how many rows were
    # rendered, and which row values were missing under which policy. Kept
    # separate from `field_lineage` because a repeated field has no single
    # value -- it has one per record of the collection.
    repeat_lineage: list = field(default_factory=list)
    qa_passed: bool = True
    qa_notes: list = field(default_factory=list)
    # The same failures as `qa_notes`, structured. §22 asks for every QA failure
    # logged with "the check that fired, the object involved, and whether a
    # human overrode it", and a list of sentences answers none of the three --
    # grouping by check meant matching strings against the registry, and the
    # object could only be read out of the prose by a person.
    qa_findings: list = field(default_factory=list)

    def block(self, note: str, *, check: str, object_id: str | None = None) -> None:
        """Record a blocking QA failure. `qa_passed` never returns to True.

        `check` names the gate in `app/qa/policy.py` this failure belongs to and
        `object_id` the manifest object it happened on, so the finding can be
        counted alongside the ones the policy-routed checks produce.
        """
        self.qa_passed = False
        if note not in self.qa_notes:
            self.qa_notes.append(note)
            self.qa_findings.append(
                {"check": check, "object_id": object_id, "detail": note, "severity": BLOCKING}
            )


def _today() -> date:
    """A named seam so a golden comparison can pin the date.

    `system.today` fields put the generation date into the letter, so a golden
    .docx captures whichever day it was recorded on and every later day fails the
    comparison. A harness that goes red overnight for no reason is a harness
    people learn to ignore.
    """
    return date.today()


def _resolve_field_value(f: dict, record: dict, locale: str = DEFAULT_LOCALE) -> tuple[str | None, str]:
    fid = f["id"]
    if f.get("value_rule") == "system.today" or fid in ("date_dd_mm_yyyy", "current_date"):
        return format_value(_today(), {**f, "type": "date", "format": f.get("format", "DD/MM/YYYY")}, locale), "system"
    if fid in record and record[fid] not in (None, ""):
        return format_value(record[fid], f, locale), "source_record"
    for slot in f.get("slots", []):
        if slot.get("kind") == "mergefield":
            code = slot["code"]
            # Same emptiness test as the direct hit above. Accepting a present-
            # but-null cell here formatted `None` into the letter and reported
            # the field resolved, which is precisely the silent blank the
            # on_missing policy exists to catch.
            if code in record and record[code] not in (None, ""):
                return format_value(record[code], f, locale), "source_record"
    return None, "missing"


def _record_field(result: "FillResult", fid: str, value, source: str, policy: str) -> None:
    """Append one field's lineage, blocking the document when policy says so."""
    blocking = blocks_document(source, policy)
    if blocking:
        result.block(missing_value_note(fid), check=REQUIRED_VALUE_MISSING, object_id=fid)
    result.field_lineage.append(
        FieldLineage(field_id=fid, value=value, source=source, on_missing=policy, blocking=blocking).__dict__
    )


def _sentence_bounds(text: str, at: int) -> tuple[int, int]:
    """Half-open char range of the sentence containing offset `at`."""
    starts = [0] + [m.end() for m in _SENTENCE_END_RE.finditer(text)] + [len(text)]
    for lo, hi in zip(starts, starts[1:]):
        if lo <= at < hi:
            return lo, hi
    return 0, len(text)


def _apply_sentence_removal(paragraph_el, marker: str) -> bool:
    """Delete the sentence holding `marker`, across every run it spans.

    Works on the paragraph's text nodes rather than the single run the
    placeholder landed in, for two reasons: Word fragments a sentence across
    runs at every formatting change, and a MERGEFIELD is replaced by a brand-new
    run that no span index knows about. Editing `<w:t>` directly covers both.
    """
    t_els = list(paragraph_el.iter(_q("t")))
    texts = [t.text or "" for t in t_els]
    full = "".join(texts)
    at = full.find(marker)
    if at < 0:
        return False

    lo, hi = _sentence_bounds(full, at)
    cursor = 0
    for t_el, text in zip(t_els, texts):
        start = cursor
        cursor += len(text)
        kept = text[: max(0, min(len(text), lo - start))] + text[max(0, min(len(text), hi - start)):]
        if kept != text:
            t_el.text = kept
    return True


def _strip_markers(paragraph_el) -> None:
    """Last resort: a marker that outlived its removal must never be rendered."""
    for t_el in paragraph_el.iter(_q("t")):
        if t_el.text and _MARKER_OPEN in t_el.text:
            t_el.text = _MARKER_RE.sub("", t_el.text)


def _clear_run_color(run_el) -> None:
    rpr = run_el.find(_q("rPr"))
    if rpr is not None:
        color_el = rpr.find(_q("color"))
        if color_el is not None:
            rpr.remove(color_el)


def _set_span_text(span: RunSpan, text: str) -> None:
    if not span.run_elements:
        return
    first = span.run_elements[0]
    t_el = first.find(_q("t"))
    if t_el is None:
        t_el = etree.SubElement(first, _q("t"))
    t_el.text = text
    t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    _clear_run_color(first)
    for extra in span.run_elements[1:]:
        for t in extra.findall(_q("t")):
            extra.remove(t)


def _paragraph_table_ancestor(p_el):
    parent = p_el.getparent()
    while parent is not None:
        if etree.QName(parent).localname == "tbl":
            return parent
        parent = parent.getparent()
    return None


def _paragraph_row_ancestor(p_el):
    """The `w:tr` a paragraph sits in, or None outside a table.

    Innermost first, so a paragraph in a nested table drops the nested row
    rather than the row of the table containing it.
    """
    parent = p_el.getparent()
    while parent is not None:
        if etree.QName(parent).localname == "tr":
            return parent
        parent = parent.getparent()
    return None


# ---- §6 TABLE_ROW: one template row, rendered once per record of a collection ----

def _locate_prototype_row(spec: dict, paragraphs: list):
    """`(row_el, paragraph_indices)` for the template row this spec repeats.

    Located by the column *tokens* rather than by a stored paragraph index, for
    the same reason the blueprint model stores no indices: a number goes stale
    the moment a paragraph is inserted above it, silently, while the token
    `<Item Description>` either is in the row or is not. Returns None when no
    table row carries any of the tokens -- a manifest promising repetition the
    document cannot honour, which the caller reports rather than ignores.

    The `is` comparison below is safe because `row_el` is held alive for the
    whole comprehension: lxml returns the same proxy for an element as long as
    one reference to that proxy exists.
    """
    tokens = [c.get("token") for c in (spec.get("columns") or ()) if c.get("token")]
    if not tokens:
        return None
    # Score every candidate row by how many of THIS spec's tokens it carries,
    # and take the best. First-any-token binding was how two repeating tables
    # sharing a column name ("Amount" on both services and expenses) both bound
    # the first table's row -- the second then rendered nothing, quietly.
    candidates: list = []   # (score, order, row_el)
    seen_rows: list = []
    for p_el in paragraphs:
        row = _paragraph_row_ancestor(p_el)
        if row is None or any(row is r for r in seen_rows):
            continue
        seen_rows.append(row)
        text = "".join(t.text or "" for t in row.iter(_q("t")))
        score = sum(1 for token in tokens if token in text)
        if score:
            candidates.append((score, len(candidates), row))
    if not candidates:
        return None
    row_el = max(candidates, key=lambda c: (c[0], -c[1]))[2]
    indices = {i for i, p in enumerate(paragraphs) if _paragraph_row_ancestor(p) is row_el}
    return row_el, indices


def _replace_in_row(row_el, token: str, value: str) -> bool:
    """Replace the first occurrence of `token` across the row's text nodes.

    Works on the concatenation of every `<w:t>` in the row rather than on any
    single run, because Word fragments a placeholder across runs at every
    formatting change. Same offset-walking technique as
    `_apply_sentence_removal`, which exists for the same reason.
    """
    t_els = list(row_el.iter(_q("t")))
    texts = [t.text or "" for t in t_els]
    full = "".join(texts)
    at = full.find(token)
    if at < 0:
        return False
    lo, hi = at, at + len(token)
    cursor = 0
    for t_el, text in zip(t_els, texts):
        start = cursor
        cursor += len(text)
        if cursor <= lo or start >= hi:
            continue
        head = text[: max(0, lo - start)]
        tail = text[max(0, min(len(text), hi - start)):]
        insert = value if start <= lo else ""
        t_el.text = head + insert + tail
        t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return True


def _clear_placeholder_colours(row_el) -> None:
    """Strip the placeholder blue from a cloned row's runs.

    The prototype row's tokens are written in the authoring blue so the
    pre-scanner classifies them; a rendered line item is the reader's data and
    must not arrive coloured like an unfilled slot.
    """
    from app.templates.parsers.docx_prescan import BLUE_RGBS

    for r in row_el.iter(_q("r")):
        rpr = r.find(_q("rPr"))
        if rpr is None:
            continue
        color_el = rpr.find(_q("color"))
        if color_el is not None and (color_el.get(_q("val")) or "").upper() in BLUE_RGBS:
            rpr.remove(color_el)


def _column_field(column: dict) -> dict:
    """A field-shaped dict for `format_value`, from one TABLE_ROW column spec."""
    return {
        "id": column.get("source_key") or column.get("field_id") or "",
        "type": column.get("type") or "string",
        "format": column.get("format"),
    }


def _apply_row_repeats(repeat_rows: list, source_record: dict, locale: str,
                       result: "FillResult", drop_paragraph_indices: set) -> None:
    """Render each prototype row once per record, then remove the prototype.

    Runs *last*, after every coordinate-driven pass, so the clones never enter
    the span space at all -- the coordinate-shift problem is dissolved rather
    than solved. Anything the scalar passes did to the rest of the document is
    already in the tree; anything a clone carries came from the prototype.
    """
    for entry in repeat_rows:
        spec, tr = entry["spec"], entry["row_el"]
        object_id = spec.get("id") or spec.get("object_id") or "table_row"
        iterate_over = str(spec.get("iterate_over") or "")
        columns = list(spec.get("columns") or ())

        if tr.getparent() is None:
            if entry.get("paragraphs", set()) & drop_paragraph_indices:
                # The condition machinery dropped the region holding this
                # table. That is a decision the manifest already recorded.
                result.repeat_lineage.append({
                    "object_id": object_id, "iterate_over": iterate_over,
                    "rows_rendered": 0, "skipped": "row removed by a condition",
                })
            else:
                # Nothing legitimate removed it. Rendering zero rows quietly
                # here is how a second collection disappears from a document
                # with QA green.
                result.block(
                    f"The prototype row for '{object_id}' is no longer in the document, and "
                    "no condition removed it. The template's repeating rows overlap; give "
                    "each its own tokens.",
                    check=ROW_REPEAT, object_id=object_id,
                )
            continue

        items = source_record.get(iterate_over)
        if items is None:
            items = []
        if not isinstance(items, (list, tuple)):
            result.block(
                f"'{iterate_over}' should be a list of rows for the repeating table row "
                f"'{object_id}', but the source record carries {type(items).__name__} instead.",
                check=ROW_REPEAT, object_id=object_id,
            )
            continue

        rendered, misses = 0, []
        broken = False
        for row_index, item in enumerate(items):
            if not isinstance(item, dict):
                result.block(
                    f"Row {row_index + 1} of '{iterate_over}' is {type(item).__name__}, "
                    "not an object of column values, so nothing can be filled from it.",
                    check=ROW_REPEAT, object_id=object_id,
                )
                broken = True
                break
            clone = copy.deepcopy(tr)
            for column in columns:
                token = column.get("token") or ""
                key = column.get("source_key") or column.get("field_id") or ""
                raw = item.get(key)
                if raw in (None, ""):
                    policy = str(column.get("on_missing") or BLANK).upper()
                    if policy == "BLOCK":
                        result.block(
                            f"Row {row_index + 1} of '{iterate_over}' has no value for "
                            f"'{key}', which this template requires.",
                            check=ROW_REPEAT, object_id=object_id,
                        )
                        value = ""
                    elif policy == DEFAULT and column.get("default") is not None:
                        value = format_value(column.get("default"), _column_field(column), locale)
                    else:
                        value = ""
                    misses.append({"row": row_index, "column": key, "policy": policy})
                else:
                    value = format_value(raw, _column_field(column), locale)
                if token:
                    _replace_in_row(clone, token, value)
            _clear_placeholder_colours(clone)
            tr.addprevious(clone)
            rendered += 1

        empty_behaviour = str(spec.get("empty_behaviour") or "REMOVE_ROW").upper()
        parent = tr.getparent()
        if parent is not None:
            if rendered == 0 and not broken and empty_behaviour == "REMOVE_TABLE":
                table = _paragraph_table_ancestor(tr)
                if table is not None and table.getparent() is not None:
                    table.getparent().remove(table)
            else:
                parent.remove(tr)

        if rendered == 0 and not broken and spec.get("required"):
            result.block(
                f"'{iterate_over}' is empty, and this template requires at least one row "
                f"in its '{object_id}' table.",
                check=ROW_REPEAT, object_id=object_id,
            )
        result.repeat_lineage.append({
            "object_id": object_id, "iterate_over": iterate_over,
            "rows_rendered": rendered, "missing": misses,
        })


def fill_template(
    template_path: str,
    output_path: str,
    manifest: dict,
    source_record: dict,
    locale: str = DEFAULT_LOCALE,
    condition_verdicts: dict | None = None,
) -> FillResult:
    """Mutate a copy of `template_path` into `output_path`.

    `condition_verdicts` lets the resolution orchestrator hand down verdicts it
    already computed -- fuzzy and context-dependent conditions cannot be decided
    here, because this layer sees only one record and no prior units. When it is
    omitted every condition is evaluated locally, which is the deterministic path.
    """
    document = docx.Document(template_path)
    body_el = document.element.body
    paragraphs, table_flags_list = _walk_paragraphs(body_el)
    table_paragraph_indices = {i for i, t in enumerate(table_flags_list) if t}

    # Re-derive live spans + mergefields with the SAME structural algorithm the
    # compiler used (docx_prescan.prescan), but against *this* live document
    # tree (not a second, throwaway parse) so (paragraph_index, span_index)
    # coordinates line up exactly and the elements we mutate are the real ones.
    from app.templates.parsers.docx_prescan import HIGHLIGHT_ROLES, _classify_color, _run_style_key

    spans_by_key: dict[tuple[int, int], RunSpan] = {}
    mergefields_live: list[MergeField] = []
    uses_highlight_markup = False

    for p_idx, p in enumerate(paragraphs):
        run_els = []
        for child in p:
            tag = etree.QName(child).localname
            if tag == "r":
                run_els.append((child, False))
            elif tag == "hyperlink":
                for r in child.findall(_q("r")):
                    run_els.append((r, True))

        i, n = 0, len(run_els)
        consumed = set()
        while i < n:
            r, _hl = run_els[i]
            fld_char = r.find(_q("fldChar"))
            if fld_char is not None and fld_char.get(_q("fldCharType")) == "begin":
                instr_text = None
                j = i + 1
                state = "seeking_instr"
                seq = [r]
                while j < n:
                    rj, _ = run_els[j]
                    seq.append(rj)
                    instr_el = rj.find(_q("instrText"))
                    fc = rj.find(_q("fldChar"))
                    if instr_el is not None and state == "seeking_instr":
                        instr_text = (instr_text or "") + (instr_el.text or "")
                    elif fc is not None and fc.get(_q("fldCharType")) == "separate":
                        state = "collecting_result"
                    elif fc is not None and fc.get(_q("fldCharType")) == "end":
                        j += 1
                        break
                    j += 1
                if instr_text and "MERGEFIELD" in instr_text:
                    code = instr_text.replace("MERGEFIELD", "").replace("\\* MERGEFORMAT", "").strip()
                    mergefields_live.append(MergeField(paragraph_index=p_idx, code=code, field_elements=seq))
                for k in range(i, j):
                    consumed.add(k)
                i = j
                continue
            i += 1

        span_idx = 0
        current = None
        for k, (r, is_hl) in enumerate(run_els):
            if k in consumed:
                current = None
                continue
            text = _run_text(r)
            if not text:
                continue
            style = _run_style_key(r)
            color = _classify_color(style[0], style[3])
            if style[3] and style[3].lower() in HIGHLIGHT_ROLES:
                uses_highlight_markup = True
            if current and current.color == color and current.in_hyperlink == is_hl:
                current.text += text
                current.run_elements.append(r)
            else:
                current = RunSpan(paragraph_index=p_idx, span_index=span_idx, color=color, text=text, run_elements=[r], in_hyperlink=is_hl)
                spans_by_key[(p_idx, span_idx)] = current
                span_idx += 1

    result = FillResult()
    # (paragraph_index, marker) for each REMOVE_SENTENCE field that had no value.
    # Collected while filling, applied once the paragraph is whole again.
    sentence_removals: list[tuple[int, str]] = []

    # ---- 0. repeating table rows (§6 TABLE_ROW): locate prototypes now, fill last ----
    # Located before anything is dropped, because the tokens are still where the
    # template put them; *filled* after every other pass (step 5e), because a
    # cloned row must never enter the (paragraph_index, span_index) space the
    # passes in between address.
    table_row_specs = [b for b in manifest["blocks"] if str(b.get("object_type") or "").upper() == "TABLE_ROW"]
    repeat_rows: list[dict] = []
    repeat_paragraph_indices: set[int] = set()
    for spec in table_row_specs:
        located = _locate_prototype_row(spec, paragraphs)
        if located is None:
            result.block(
                f"The manifest declares a repeating table row "
                f"'{spec.get('id') or spec.get('object_id')}', but no table row in this template "
                "carries its column tokens. The template and its reading disagree; read the "
                "template again.",
                check=ROW_REPEAT, object_id=spec.get("id") or spec.get("object_id"),
            )
            continue
        row_el, row_paragraphs = located
        taken = next((entry for entry in repeat_rows if entry["row_el"] is row_el), None)
        if taken is not None:
            result.block(
                f"Repeating rows '{spec.get('id') or spec.get('object_id')}' and "
                f"'{taken['spec'].get('id') or taken['spec'].get('object_id')}' both resolve "
                "to the same table row -- their column tokens overlap too much to tell the "
                "rows apart. Give each table's columns distinct tokens.",
                check=ROW_REPEAT, object_id=spec.get("id") or spec.get("object_id"),
            )
            continue
        repeat_rows.append({"spec": spec, "row_el": row_el, "paragraphs": row_paragraphs})
        repeat_paragraph_indices |= row_paragraphs

    # A field whose every slot sits inside a prototype row has no single value:
    # it is a column, resolved once per record of the collection in step 5e. It
    # is excluded from the scalar passes below, or the default BLANK policy
    # would blank the prototype's tokens before the repeat pass could read them.
    scalar_fields = manifest["fields"]
    if repeat_paragraph_indices:
        scalar_fields = [
            f for f in manifest["fields"]
            if not (f.get("slots") and all(
                slot.get("paragraph_index") in repeat_paragraph_indices
                for slot in f["slots"]))
        ]

    # ---- 1. evaluate conditions ----
    blocks_by_id = {b["id"]: b for b in manifest["blocks"]}
    drop_block_ids: set[str] = set()
    condition_verdicts = condition_verdicts or {}
    for cond in manifest["conditions"]:
        if cond["id"] in condition_verdicts:
            # Handed down by the resolution orchestrator, which saw context this
            # layer cannot. Already a decision, so it is never undecided.
            decided, reason, missing = bool(condition_verdicts[cond["id"]]), "supplied", ()
        else:
            verdict = evaluate_condition(cond["expression"], source_record)
            decided, reason, missing = verdict.value, verdict.reason, verdict.missing_fields

        result.condition_lineage.append(ConditionLineage(
            condition_id=cond["id"], expression=cond["expression"], result=decided,
            reason=reason, missing_fields=list(missing),
        ).__dict__)

        if decided is None:
            # Undecided is not false. Dropping the block here is what silently
            # removes a section the employee was entitled to see, so keep the
            # block and stop the document instead -- a letter that fails loudly
            # is recoverable, one that is quietly missing a clause is not.
            detail = (
                f"needs {', '.join(missing)}, which the source record does not provide"
                if reason == "missing_input"
                else f"could not be parsed: {cond['expression']!r}"
            )
            result.block(
                f"Condition '{cond['id']}' {detail}. The section it governs was kept and the document blocked.",
                check=BRANCH_SELECTION, object_id=cond["id"],
            )
        elif decided is False:
            drop_block_ids.update(cond["keeps_blocks"])

    drop_paragraph_indices: set[int] = set()
    kept_block_paragraphs: set[int] = set()
    # Span-scoped blocks (an inline switch: two branches on one line). Dropping
    # one blanks its spans and leaves the paragraph -- the static text around the
    # switch belongs to neither branch and has to survive. `touched_paragraphs`
    # are the lines this pass edits mid-sentence, and are the only ones whose
    # spacing and final punctuation are repaired afterwards.
    drop_span_keys: set[tuple[int, int]] = set()
    touched_paragraphs: set[int] = set()
    for b in manifest["blocks"]:
        if str(b.get("object_type") or "").upper() == "TABLE_ROW":
            continue  # located by token and filled in step 5e; it has no paragraph range
        if b.get("start_span") is not None:
            p_idx = b["start_paragraph"]
            touched_paragraphs.add(p_idx)
            if b["id"] in drop_block_ids:
                drop_span_keys.update((p_idx, s) for s in range(b["start_span"], b["end_span"] + 1))
            continue
        rng = range(b["start_paragraph"], b["end_paragraph"] + 1)
        if b["id"] in drop_block_ids:
            drop_paragraph_indices.update(rng)
        else:
            kept_block_paragraphs.update(rng)

    # ---- 2. delete_always (instruction markers / prose) not already covered by a
    # dropped block, and NOT inside a *kept* block's own content range -- a red
    # run can be genuine kept content (e.g. this template's red "Dates of Effect"
    # sub-heading inside the Fixed-Term block) rather than an instruction, and
    # only the instruction *marker* paragraph (which sits before the block starts)
    # is meant to be deleted unconditionally. ----
    # Span-scoped removals (control tokens like `[[IF x > 0]]` / `[[ENDIF]]`) are
    # applied first and unconditionally, including inside kept blocks: the marker
    # is scaffolding wherever it sits, and on a line such as
    # `[[IF address_line2 is not blank]]<address_line2>[[ENDIF]]` the address has
    # to survive while both markers go. Paragraphs left empty by this are cleaned
    # up after filling, once their real content is known.
    marker_only_paragraphs: set[int] = set()
    original_paragraph_text: dict[int, str] = {}
    for p_idx in touched_paragraphs:
        if p_idx < len(paragraphs):
            original_paragraph_text[p_idx] = "".join(t.text or "" for t in paragraphs[p_idx].iter(_q("t")))

    for entry in manifest["delete_always"]:
        removals = entry.get("remove")
        span_scoped = entry.get("scope") == "span"
        if not removals and not span_scoped:
            continue
        span = spans_by_key.get((entry["paragraph_index"], entry["span_index"]))
        if span is None:
            continue
        if span_scoped:
            text = ""
        else:
            text = span.text
            for literal in removals:
                text = text.replace(literal, "")
        _set_span_text(span, text)
        span.text = text
        marker_only_paragraphs.add(entry["paragraph_index"])

    # The losing branch of an inline switch. Blanking the span is the whole point
    # of scoping the block to spans: deleting the paragraph instead is what took
    # `Please return one copy to your Manager` out of a client's letter while QA
    # reported it clean.
    for key in drop_span_keys:
        span = spans_by_key.get(key)
        if span is not None:
            _set_span_text(span, "")
            span.text = ""

    for entry in manifest["delete_always"]:
        if entry.get("partial") or entry.get("remove") or entry.get("scope") == "span":
            continue  # partial: handled when its sibling field slot is filled below
        p_idx = entry["paragraph_index"]
        if p_idx in drop_paragraph_indices:
            continue
        if p_idx in kept_block_paragraphs:
            # `kept_block_paragraphs` exists so a red run that is genuine content
            # (this template's "Dates of Effect" sub-heading) survives inside a
            # kept block. A `Put the X from source file` line never is -- it is
            # addressed to whoever assembles the letter. Blank the span rather
            # than drop the line: the instruction often shares its paragraph with
            # the very mergefield it describes, and that has to be filled.
            span = spans_by_key.get((p_idx, entry["span_index"]))
            if span is not None and PUT_INSTRUCTION_RE.match((span.text or "").strip()):
                _set_span_text(span, "")
                span.text = ""
                marker_only_paragraphs.add(p_idx)
            continue
        drop_paragraph_indices.add(p_idx)

    # ---- 3. remove dropped rows and dropped body paragraphs ----
    # A dropped paragraph inside a table removes its ROW, not its table.
    #
    # This used to remove the whole table, which is right only when the table is
    # wholly conditional -- a Full Time remuneration table standing next to a
    # Part Time one. It is catastrophic when a table merely *contains* a
    # conditional row. Measured on the compensation template: a 16-row salary
    # breakdown carrying three `[[IF ...]]` rows (variable pay, joining bonus,
    # retention bonus). An employee with no retention bonus dropped one row and
    # lost the entire table -- basic salary, HRA, conveyance, LTA, PF, gratuity,
    # every one of them -- from a letter whose subject is their compensation.
    # Three golden fixtures had that damage baked in as expected output.
    #
    # Row-level removal subsumes the old behaviour rather than contradicting it:
    # a table whose every paragraph is dropped loses every row, and the empty
    # table is then removed below. So the wholly-conditional case still works,
    # and it now works by construction instead of by a separate rule.
    #
    # Elements are removed through the live reference we already hold rather
    # than by looking `id(el)` up in a second pass. lxml proxies are created on
    # demand and freed when the last reference goes, so `id()` is neither stable
    # across two lookups nor unique -- a freed proxy's address gets reused. That
    # produced a Fixed Term colleague's letter keeping the *Full Time* table,
    # complete with unresolved «LAB__FT_SALARY__38_HR_» codes, on some runs and
    # not others.
    paras_to_remove = []
    rows_to_remove = []
    touched_tables = []
    for p_idx in sorted(drop_paragraph_indices):
        if p_idx >= len(paragraphs):
            continue
        p_el = paragraphs[p_idx]
        row = _paragraph_row_ancestor(p_el)
        if row is not None:
            rows_to_remove.append(row)
            table = _paragraph_table_ancestor(p_el)
            if table is not None:
                touched_tables.append(table)
        else:
            paras_to_remove.append(p_el)

    for row in rows_to_remove:
        parent = row.getparent()
        # Already gone: another dropped paragraph in the same row removed it.
        if parent is not None:
            parent.remove(row)

    # A table that has lost every row is an empty frame -- borders and nothing
    # else. That is the wholly-conditional table, and it goes.
    for table in touched_tables:
        parent = table.getparent()
        if parent is not None and not table.findall(_q("tr")):
            parent.remove(table)

    for p_el in paras_to_remove:
        if p_el.getparent() is not None:
            p_el.getparent().remove(p_el)

    # ---- 4. fill field slots, grouped by the run-span they live in ----
    # A single merged span can carry more than one bracket token (e.g. this
    # template's "<Colleague First Name> <Colleague Last Name> ," address
    # line) -- filling field-by-field with a blind whole-span overwrite would
    # let the second field clobber the first. Instead: gather every
    # (bracket_token -> resolved_value) pair per span, then resolve the span's
    # final text in one pass.
    # A field that is BOTH a repeating column and a scalar slot (the same
    # <Amount> token inside the row and after "Total payable:") has no scalar
    # value -- its values live per record inside the collection -- so BLANK
    # would print an empty total with QA green. Missing + column-owned blocks.
    repeat_column_field_ids = {
        c.get("field_id")
        for entry in repeat_rows
        for c in (entry["spec"].get("columns") or ())
        if c.get("field_id")
    }
    field_values: dict[str, tuple[str | None, str]] = {}
    field_policies: dict[str, str] = {}
    for f in scalar_fields:
        value, source = _resolve_field_value(f, source_record, locale)
        if source == "missing" and f["id"] in repeat_column_field_ids:
            result.block(
                f"Field '{f['id']}' is a repeating-row column but its token also appears "
                "outside the row, where no single value exists for it. Give the outside "
                "slot its own token (a total wants <Grand Total>, not the row's <Amount>).",
                check=ROW_REPEAT, object_id=f["id"],
            )
        policy = field_on_missing(f)
        field_policies[f["id"]] = policy
        if source == "missing" and policy == DEFAULT:
            declared = f.get("default")
            if declared is None:
                result.block(
                    f"Field '{f['id']}' has no value and its on_missing policy is DEFAULT, "
                    "but the manifest declares no default to fall back to.",
                    check=REQUIRED_VALUE_MISSING, object_id=f["id"],
                )
            else:
                value, source = format_value(declared, f, locale), "fallback"
        field_values[f["id"]] = (value, source)

    slots_by_span: dict[tuple, list[dict]] = {}
    for f in scalar_fields:
        for slot in f.get("slots", []):
            if slot.get("kind") == "mergefield":
                continue
            if slot["paragraph_index"] in drop_paragraph_indices:
                continue
            if slot["paragraph_index"] in repeat_paragraph_indices:
                continue  # a slot inside a prototype row belongs to the repeat pass
            key = (slot["paragraph_index"], slot["span_index"])
            if key in drop_span_keys:
                continue  # losing branch of an inline switch, already blanked
            slots_by_span.setdefault(key, []).append({**slot, "field_id": f["id"]})

    filled_field_ids = set()
    for key, slot_list in slots_by_span.items():
        span = spans_by_key.get(key)
        if span is None:
            continue
        original_text = "".join(_run_text(r) for r in span.run_elements)
        is_instruction_hybrid = bool(PUT_INSTRUCTION_RE.match(original_text.strip()))

        resolved_in_order = []
        new_text = original_text
        for slot in slot_list:
            fid = slot["field_id"]
            value, source = field_values[fid]
            if source == "missing" and field_policies[fid] == REMOVE_SENTENCE:
                # Substituting a marker rather than deleting here: the sentence
                # this placeholder sits in usually runs across several runs, so
                # the removal has to happen once the whole paragraph is readable.
                resolved_value = f"{_MARKER_OPEN}RS{len(sentence_removals)}{_MARKER_CLOSE}"
                sentence_removals.append((key[0], resolved_value))
            else:
                resolved_value = value if value is not None else ""
            resolved_in_order.append(resolved_value)
            new_text = new_text.replace(slot["text"], resolved_value, 1)

        final_text = " ".join(v for v in resolved_in_order if v) if is_instruction_hybrid else new_text
        _set_span_text(span, final_text)
        for slot in slot_list:
            filled_field_ids.add(slot["field_id"])

    for fid in filled_field_ids:
        value, source = field_values[fid]
        _record_field(result, fid, value, source, field_policies[fid])

    # ---- 5. resolve mergefields (replace the complex-field run sequence with one literal run) ----
    mf_by_para_code: dict[tuple, MergeField] = {(mf.paragraph_index, mf.code): mf for mf in mergefields_live}
    for f in scalar_fields:
        for slot in f.get("slots", []):
            if slot.get("kind") != "mergefield":
                continue
            if slot["paragraph_index"] in drop_paragraph_indices:
                continue
            if slot["paragraph_index"] in repeat_paragraph_indices:
                continue
            mf = mf_by_para_code.get((slot["paragraph_index"], slot["code"]))
            if mf is None or not mf.field_elements:
                continue
            value, source = field_values[f["id"]]
            policy = field_policies[f["id"]]
            if source == "missing" and policy == REMOVE_SENTENCE:
                written = f"{_MARKER_OPEN}RS{len(sentence_removals)}{_MARKER_CLOSE}"
                sentence_removals.append((slot["paragraph_index"], written))
            else:
                written = value if value is not None else ""
            anchor = mf.field_elements[0]
            new_run = etree.SubElement(anchor.getparent(), _q("r"))
            anchor.addprevious(new_run)
            t_el = etree.SubElement(new_run, _q("t"))
            t_el.text = written
            for el in mf.field_elements:
                if el.getparent() is not None:
                    el.getparent().remove(el)
            _record_field(result, f["id"], value, source, policy)

    # ---- 5a. REMOVE_SENTENCE: cut the sentences whose field never resolved ----
    # Deferred to here because the substitution above had to finish first: the
    # sentence is only readable once every other placeholder on the line has
    # become its value.
    for p_idx, marker in sentence_removals:
        if p_idx >= len(paragraphs):
            continue
        p_el = paragraphs[p_idx]
        if p_el.getparent() is None:
            continue  # its block was dropped anyway
        _apply_sentence_removal(p_el, marker)
        _strip_markers(p_el)
        # An entirely-conditional line is now empty; hand it to the cleanup below
        # rather than leaving a blank gap where the sentence used to be.
        marker_only_paragraphs.add(p_idx)

    # A line that held nothing but control tokens is now blank. Left in place it
    # becomes a stray empty paragraph between clauses, so the letter reads with
    # a gap wherever the template had an `[[IF]]` or `[[ENDIF]]` on its own line.
    # Only paragraphs that actually carried a marker are considered, so genuine
    # blank lines in the template are preserved.
    for p_idx in sorted(marker_only_paragraphs, reverse=True):
        if p_idx >= len(paragraphs):
            continue
        p_el = paragraphs[p_idx]
        if p_el.getparent() is None:
            continue
        if "".join(t.text or "" for t in p_el.iter(_q("t"))).strip():
            continue  # real content survived on this line
        p_el.getparent().remove(p_el)

    # ---- 5b. repair the lines an inline switch was cut out of ----
    # Removing runs from the middle of a sentence leaves the seams: the spaces
    # that separated the deleted runs collapse together, and a trailing
    # instruction takes the sentence's full stop with it when it goes.
    for p_idx in sorted(touched_paragraphs):
        if p_idx >= len(paragraphs):
            continue
        p_el = paragraphs[p_idx]
        if p_el.getparent() is None:
            continue
        t_els = [t for t in p_el.iter(_q("t"))]
        prev_ends_space = False
        for t_el in t_els:
            if not t_el.text:
                continue
            text = re.sub(r"[ \t]{2,}", " ", t_el.text)
            if prev_ends_space and text.startswith(" "):
                text = text.lstrip(" ")
            if text:
                prev_ends_space = text.endswith(" ")
                t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            t_el.text = text

        original = original_paragraph_text.get(p_idx, "")
        current = "".join(t.text or "" for t in t_els)
        if original.rstrip().endswith(".") and current.strip() and not current.rstrip().endswith("."):
            for t_el in reversed(t_els):
                if t_el.text and t_el.text.strip():
                    t_el.text = t_el.text.rstrip() + "."
                    break

    # ---- 5c. highlights are markup, not formatting ----
    # Green/yellow annotation tells the assembler what to do with a run; it is
    # not something the recipient of the letter should ever see. Only touched
    # when the template actually used the highlight convention, so a red/blue
    # template keeps whatever highlighting its author intended.
    if uses_highlight_markup:
        for hl in list(body_el.iter(_q("highlight"))):
            if (hl.get(_q("val")) or "").lower() in HIGHLIGHT_ROLES and hl.getparent() is not None:
                hl.getparent().remove(hl)

    # ---- 5e. repeating table rows: clone the prototype once per record ----
    # Dead last on purpose: every coordinate-driven pass above is finished, so
    # the clones never shift an index anything still needs. See step 0.
    if repeat_rows:
        _apply_row_repeats(repeat_rows, source_record, locale, result, drop_paragraph_indices)

    # ---- 5f. Korean postposition agreement ----
    # Runs after filling AND after the repeat pass, because the syllable that
    # decides the form is the value that was just inserted -- and for a
    # particle beside a line-item token, that value only exists in the
    # clones. Resolving earlier read the raw token's '>' and fixed the vowel
    # form into every row. The decision is made on the whole paragraph's text
    # -- the deciding syllable and the alternation routinely sit in different runs
    # -- while the edit is applied to the single run that carries the alternation,
    # so no run's formatting is disturbed.
    if manifest.get("particles"):
        from app.generation.korean import ALTERNATION_RE, choose, has_batchim  # noqa: F401
        from app.generation.korean import _final_consonant

        # body_el.iter yields only attached paragraphs -- the same set the
        # old paragraphs-list walk visited, plus the clones.
        for p_el in body_el.iter(_q("p")):
            t_els = [t for t in p_el.iter(_q("t"))]
            joined = "".join(t.text or "" for t in t_els)
            if not ALTERNATION_RE.search(joined):
                continue

            # Map each match onto the run that contains it, walking offsets once.
            bounds, offset = [], 0
            for t_el in t_els:
                length = len(t_el.text or "")
                bounds.append((offset, offset + length, t_el))
                offset += length

            for m in reversed(list(ALTERNATION_RE.finditer(joined))):
                preceding = joined[: m.start()].rstrip()
                if not preceding:
                    continue
                target = next(
                    (t for start, end, t in bounds if start <= m.start() and m.end() <= end), None
                )
                if target is None:
                    result.qa_notes.append(
                        "A Korean particle alternation is split across runs and was left as written; "
                        "resave the template in Word to merge the runs."
                    )
                    continue
                for start, end, t in bounds:
                    if t is target:
                        local = m.start() - start, m.end() - start
                        break
                replacement = choose(preceding, m.group(1), m.group(2))
                text = target.text or ""
                target.text = text[: local[0]] + replacement + text[local[1]:]
                if _final_consonant(preceding) is None:
                    result.qa_notes.append(
                        f"Korean particle after {preceding[-12:]!r} could not be decided from the "
                        f"value's script; the vowel form was used. Confirm with a native reader."
                    )

    document.save(output_path)

    # ---- 5d. archive normalisation ----
    # §18 asks for 100% reproducibility "after archive timestamps are
    # normalised", and normalising them is the *writer's* job: a caller who
    # forgets makes every digest in the audit trail a record of when the file
    # was written rather than of what it contains. Done before the QA re-read,
    # so the gates below inspect the file that is actually delivered.
    normalise_docx(output_path)

    # ---- 6. QA gates ----
    # Each gate is a named check in `app/qa` (§15), and the manifest's
    # `qa_policy` (§6) decides what each one costs. The renderer runs them and
    # records what they said; it no longer holds an opinion about severity,
    # which is what let "turn this check off to unblock the batch" mean editing
    # the renderer for every tenant at once.
    check_doc = docx.Document(output_path)
    check_body = check_doc.element.body

    policy = resolve_policy(manifest.get("qa_policy"))
    qa_findings = placeholder_findings(check_body, policy)
    qa_findings += branch_findings(manifest, source_record, result.condition_lineage, drop_block_ids, policy)
    qa_findings += overflow_findings(
        check_body, scalar_fields, {fid: value for fid, (value, _source) in field_values.items()}, policy
    )
    qa_findings += layout_findings(template_path, output_path, policy)
    # Every check above this line asks "what is left over?" -- a placeholder, a
    # mergefield, an instruction, a changed region. None of them can see a
    # section that is simply gone, which is how a letter shipped missing its
    # offer sentence, its date and its hours of work while reporting qa_passed.
    # This one asks the opposite question: the user supplied a value, so where
    # is it?
    qa_findings += content_loss_findings(
        check_body, {**manifest, "fields": scalar_fields}, field_values, drop_block_ids, policy
    )
    # The template's body is passed so the doubled-word gate can tell the
    # engine's fault from the author's: this estate's own transfer letter
    # contains "offer to to the position" as a typo, and blaming the renderer
    # for it is how a gate earns a reputation for crying wolf.
    qa_findings += value_format_findings(check_body, policy, docx.Document(template_path).element.body)

    outcome = route(qa_findings, policy)
    result.qa_notes.extend(outcome.notes)
    # Warnings are kept alongside the blocking findings: a warning a person
    # approved over is precisely the "human overrode it" case §22 asks for, and
    # it cannot be recorded later if it was never recorded at all.
    result.qa_findings.extend(
        {"check": f.check, "object_id": None, "detail": f.detail, "severity": f.severity}
        for f in outcome.findings
    )
    # `and` rather than `=`. These checks are the last gate, not the only one: a
    # straight assignment here discarded every block already recorded during the
    # fill -- an unresolved required field, an undecidable condition -- because a
    # document can have all of those and still contain no leftover bracket for
    # this pass to find.
    result.qa_passed = result.qa_passed and outcome.passed

    return result
