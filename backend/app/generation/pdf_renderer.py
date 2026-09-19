"""§12's PDF path: overlay approved values onto an immutable PDF, touch nothing else.

Two unrelated jobs live here, and the docstring says so up front because
confusing them is how somebody ships a preview as a deliverable.

The first is the renderer. §12 gives the PDF path six steps -- "Approved
immutable PDF -> Stored dynamic regions / anchors / bounding boxes -> Mask
placeholder/instruction region when approved -> Overlay exact text at
deterministic coordinates -> Validate page dimensions + static-region integrity
-> Save generated PDF version" -- and one constraint that decides the whole
design: "the PDF renderer should modify only approved dynamic regions. Static
text, page geometry, logos and unaffected elements remain untouched."

That constraint rules out the obvious implementation. Rebuilding the page and
redrawing it would reproduce the layout approximately, and approximately is the
wrong answer for a document somebody signs. So this renderer never rebuilds a
page. It parses the page's existing content stream, removes exactly the
show-text operations whose origin falls inside an approved region, appends an
opaque rectangle over that region and one text operation carrying the value, and
writes the file back. Every other operator in the stream -- the logo XObject, the
rules, the static prose, the page geometry -- is passed through untouched,
because it is literally the same object.

Removing the placeholder's operators rather than painting over them is the part
worth defending. A white rectangle hides a placeholder from a reader and leaves
`<Employee Name>` sitting in the file for anyone who selects the text, copies the
page or runs `pdftotext` over it. On a letter that carries legal weight that is
not a cosmetic difference; it is the redaction failure that reaches a newspaper.
So the operators go, and the check that proves they went is text extraction on
the file that was actually written.

What "deterministic coordinates" means here is a stated rule rather than a
layout algorithm. A value inherits the baseline and the font size of the text it
replaces -- the position and the size the template's author approved -- and both
are properties of a template pinned by hash, so the rule is reproducible rather
than merely repeatable. A region that was blank in the approved PDF has nothing
to inherit, so the value's feet go on the bottom edge of the approved box and
its size comes from the caller.

The second job in this file is the review-queue preview, unchanged and still
LibreOffice's: the `.docx` is the deliverable and Word is the authority on how it
paginates, so a preview that shows five pages where Word shows four is a renderer
difference and not a defect.

Known limits, stated rather than discovered later:

  * The overlay draws in Helvetica, a standard-14 font every reader has, and
    encodes to WinAnsi. A value carrying a character WinAnsi cannot represent --
    any CJK text, for instance -- is refused loudly rather than drawn as
    mojibake. A region needing one of those needs an embedded font, which is a
    larger piece of work than this module.
  * `scan_page_regions` measures glyph widths exactly for Helvetica and Arial
    and estimates them for every other font. Its output is onboarding evidence
    for a reviewer to confirm, never a production anchor on its own; the
    manifest's stored bounding box is the anchor.
  * A page whose text is drawn as vector outlines or as an image has no
    show-text operations, so there is nothing to mask and nothing to check. The
    renderer reports that it masked nothing rather than pretending it did.
"""

from __future__ import annotations

import logging
import math
import shutil
import subprocess
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from pypdf import PdfReader, PdfWriter
from pypdf.generic import ContentStream, DecodedStreamObject, DictionaryObject, NameObject

from app.generation.reproducibility import content_digest
from app.qa.layout_integrity import PageGeometry, TextChunk, pdf_findings
from app.qa.policy import PLACEHOLDER_REMAINS, findings_for
from app.qa.policy import resolve_policy, route
from app.templates.semantic_model import (
    ANCHOR_BBOX,
    APPROVED,
    Anchor,
    PageRegion,
    TemplateInventory,
    resolve,
)

log = logging.getLogger(__name__)

#: Recorded in every document's lineage (§17). §19 lists "Renderer version
#: change alters output" as a risk whose whole mitigation is that the version was
#: written down, so this string changes whenever the drawn output would.
PDF_RENDERER_VERSION = "pdf_overlay/1.0.1"

#: Used when a region was blank in the approved PDF, so there is no masked text
#: whose size the value could inherit.
DEFAULT_FONT_SIZE = 11.0

#: Inset from the left and right edges of an approved box, in points. Text that
#: starts exactly on the box edge reads as touching whatever is beside it.
DEFAULT_PADDING = 1.0

ALIGN_LEFT, ALIGN_CENTRE, ALIGN_RIGHT = "left", "centre", "right"
ALIGNMENTS = (ALIGN_LEFT, ALIGN_CENTRE, ALIGN_RIGHT)

#: Adobe's Helvetica metrics: ascender, descender and glyph widths in 1/1000 em.
#: Carried here rather than taken from `pypdf._text_extraction`, whose copy of
#: the table is both private and visibly mis-transcribed (it gives capital U the
#: width of a full stop). A wrong width silently mis-measures every overflow
#: check, so the numbers are the ones from Helvetica.afm.
HELVETICA_ASCENDER = 0.718
HELVETICA_DESCENDER = -0.207

HELVETICA_WIDTHS = {
    " ": 278, "!": 278, '"': 355, "#": 556, "$": 556, "%": 889, "&": 667, "'": 191,
    "(": 333, ")": 333, "*": 389, "+": 584, ",": 278, "-": 333, ".": 278, "/": 278,
    "0": 556, "1": 556, "2": 556, "3": 556, "4": 556, "5": 556, "6": 556, "7": 556,
    "8": 556, "9": 556, ":": 278, ";": 278, "<": 584, "=": 584, ">": 584, "?": 556,
    "@": 1015, "A": 667, "B": 667, "C": 722, "D": 722, "E": 667, "F": 611, "G": 778,
    "H": 722, "I": 278, "J": 500, "K": 667, "L": 556, "M": 833, "N": 722, "O": 778,
    "P": 667, "Q": 778, "R": 722, "S": 667, "T": 611, "U": 722, "V": 667, "W": 944,
    "X": 667, "Y": 667, "Z": 611, "[": 278, "\\": 278, "]": 278, "^": 469, "_": 556,
    "`": 333, "a": 556, "b": 556, "c": 500, "d": 556, "e": 556, "f": 278, "g": 556,
    "h": 556, "i": 222, "j": 222, "k": 500, "l": 222, "m": 833, "n": 556, "o": 556,
    "p": 556, "q": 556, "r": 333, "s": 500, "t": 278, "u": 556, "v": 500, "w": 722,
    "x": 500, "y": 500, "z": 500, "{": 334, "|": 260, "}": 334, "~": 584,
    # The WinAnsi slots a real letter actually uses: curly quotes in a surname,
    # an en dash in a date range, a euro sign in a salary.
    "€": 556, "‘": 222, "’": 222, "“": 333, "”": 333,
    "–": 556, "—": 1000, "•": 350, "…": 1000, "‹": 333,
    "›": 333, "™": 1000, "Š": 667, "š": 500, "Ž": 611,
    "ž": 500, "Œ": 1000, "œ": 944, "Ÿ": 667, "ƒ": 556,
    "†": 556, "‡": 556, "‰": 1000, "ˆ": 333, "˜": 333,
    "‚": 222, "„": 333,
}

#: The characters WinAnsi encodes in 0x80-0x9F, where it and Latin-1 disagree.
#: Without this table a curly apostrophe in "O'Brien" either raises or draws as
#: a control character.
_WINANSI_HIGH = {
    "€": 0x80, "‚": 0x82, "ƒ": 0x83, "„": 0x84, "…": 0x85,
    "†": 0x86, "‡": 0x87, "ˆ": 0x88, "‰": 0x89, "Š": 0x8A,
    "‹": 0x8B, "Œ": 0x8C, "Ž": 0x8E, "‘": 0x91, "’": 0x92,
    "“": 0x93, "”": 0x94, "•": 0x95, "–": 0x96, "—": 0x97,
    "˜": 0x98, "™": 0x99, "š": 0x9A, "›": 0x9B, "œ": 0x9C,
    "ž": 0x9E, "Ÿ": 0x9F,
}

#: Base fonts whose metrics are Helvetica's. Arial is metrically identical by
#: design, which is why substituting one for the other never reflows a document.
_HELVETICA_ALIASES = ("helvetica", "arial")

#: What a glyph is assumed to cost when the font is not one whose metrics are
#: known. Only ever used to advance the pen between two show operations that
#: share a text object; a chunk that sets its own text matrix is positioned
#: exactly whatever font it is in.
FALLBACK_ADVANCE = 0.5

_IDENTITY = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)

#: The resource name the overlay font is added under. A generic name, because
#: resource names are visible to anyone who opens the file in a PDF inspector
#: and a product-specific one says what drew the text. If it collides with a
#: font the template already declares, a digit is appended, deterministically.
OVERLAY_FONT_NAME = "/F9"


class PdfOverlayError(RuntimeError):
    """The overlay cannot be performed as specified, so nothing is written.

    Every use of this is a refusal rather than a failure: an unapproved region,
    a value that cannot be encoded, two fills claiming the same box. Rendering
    something approximate instead would produce a document that looks finished.
    """


# ---------------------------------------------------------------- geometry --

def _multiply(m, n) -> tuple:
    """m x n, both 2x3 affine matrices in PDF's (a b c d e f) order."""
    a, b, c, d, e, f = m
    A, B, C, D, E, F = n
    return (a * A + b * C, a * B + b * D,
            c * A + d * C, c * B + d * D,
            e * A + f * C + E, e * B + f * D + F)


def _translation(tx: float, ty: float) -> tuple:
    return (1.0, 0.0, 0.0, 1.0, tx, ty)


def _number(value) -> float:
    """A content-stream operand as a float, or 0 when it is not a number.

    Operands are whatever the producer wrote. A malformed one must not take the
    whole render down: the walker's job is to locate text, and a stream that
    cannot be walked is reported by the integrity gate, which compares the file
    that was written rather than trusting this parse.
    """
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


# ------------------------------------------------------------------ fonts --

def glyph_width(character: str, base_font: str | None) -> float:
    """Advance width of one glyph, in em units.

    Accented letters are decomposed before lookup, because Helvetica gives
    "E-acute" exactly the width of "E" -- that is what makes an accent free to
    add. Anything still unknown falls back, which the docstring of this module
    names as an estimate rather than a measurement.
    """
    if base_font is not None and not any(a in base_font.lower() for a in _HELVETICA_ALIASES):
        return FALLBACK_ADVANCE
    width = HELVETICA_WIDTHS.get(character)
    if width is None:
        stripped = unicodedata.normalize("NFD", character)
        width = HELVETICA_WIDTHS.get(stripped[0]) if stripped else None
    return (width if width is not None else FALLBACK_ADVANCE * 1000) / 1000.0


def measure_text(text: str, font_size: float, base_font: str | None = "Helvetica",
                 *, char_spacing: float = 0.0, word_spacing: float = 0.0,
                 horizontal_scale: float = 1.0) -> float:
    """How wide `text` will actually be, in points.

    Exact for the font this renderer draws in, which is the case the overflow
    gate depends on: a value refused for overflow must genuinely not fit.
    """
    width = sum(glyph_width(c, base_font) for c in text) * font_size
    width += char_spacing * len(text)
    width += word_spacing * text.count(" ")
    return width * horizontal_scale


def _winansi(text: str) -> bytes:
    """`text` in WinAnsi, or a refusal naming the character that stopped it."""
    out = bytearray()
    for character in text:
        code = _WINANSI_HIGH.get(character)
        if code is None and ord(character) < 256:
            code = ord(character)
        if code is None:
            raise PdfOverlayError(
                f"{character!r} (U+{ord(character):04X}) is not in WinAnsi, so Helvetica cannot "
                f"draw it and the value {text!r} would reach the reader as mojibake. This region "
                "needs a template with an embedded font covering that script."
            )
        out.append(code)
    return bytes(out)


def _pdf_string(text: str) -> bytes:
    escaped = (_winansi(text)
               .replace(b"\\", b"\\\\")
               .replace(b"(", b"\\(")
               .replace(b")", b"\\)")
               .replace(b"\r", b"\\r")
               .replace(b"\n", b"\\n"))
    return b"(" + escaped + b")"


def _fmt(value: float) -> str:
    """Three decimals, always. A fixed format is a reproducible one."""
    return f"{value:.3f}"


def _resources(page) -> DictionaryObject:
    """The page's resource dictionary, creating one only if it truly has none.

    A page may inherit its resources from the page tree. Writing a fresh empty
    dictionary onto such a page would shadow the inherited one and unmap every
    font the static text is drawn in -- the whole page would render blank while
    every structural check still passed.
    """
    own = page.get(NameObject("/Resources"))
    if own is not None:
        return own.get_object()
    inherited = page.get_inherited("/Resources")
    if inherited is not None:
        return inherited.get_object()
    created = DictionaryObject()
    page[NameObject("/Resources")] = created
    return created


def _font_map(page) -> dict:
    """Resource name -> BaseFont, for locating text drawn by the template."""
    fonts = _resources(page).get(NameObject("/Font"))
    if fonts is None:
        return {}
    fonts = fonts.get_object()
    mapping = {}
    for name, font in fonts.items():
        try:
            base = font.get_object().get(NameObject("/BaseFont"))
        except Exception:  # a broken font entry is not a reason to stop reading text
            base = None
        mapping[str(name)] = str(base) if base is not None else None
    return mapping


def _overlay_font(page) -> str:
    """Add Helvetica to the page's resources under a name nothing else uses."""
    resources = _resources(page)
    fonts = resources.get(NameObject("/Font"))
    fonts = fonts.get_object() if fonts is not None else None
    if fonts is None:
        fonts = DictionaryObject()
        resources[NameObject("/Font")] = fonts

    name = OVERLAY_FONT_NAME
    suffix = 0
    while NameObject(name) in fonts and str(
        fonts[NameObject(name)].get_object().get(NameObject("/BaseFont"), "")
    ) != "/Helvetica":
        suffix += 1
        name = f"{OVERLAY_FONT_NAME}{suffix}"

    if NameObject(name) not in fonts:
        helvetica = DictionaryObject()
        helvetica[NameObject("/Type")] = NameObject("/Font")
        helvetica[NameObject("/Subtype")] = NameObject("/Type1")
        helvetica[NameObject("/BaseFont")] = NameObject("/Helvetica")
        helvetica[NameObject("/Encoding")] = NameObject("/WinAnsiEncoding")
        fonts[NameObject(name)] = helvetica
    return name


# --------------------------------------------------- reading what is there --

_SHOW_OPERATORS = (b"Tj", b"TJ", b"'", b'"')


def _shown_text(operands, operator: bytes) -> str:
    """The characters a show operator puts on the page.

    The numbers inside a `TJ` array are kerning adjustments, not content, so
    they are dropped from the text and applied to the pen instead.
    """
    if operator == b"TJ":
        parts = []
        for item in (operands[0] if operands else ()):
            if isinstance(item, str):
                parts.append(item)
            elif isinstance(item, bytes):
                parts.append(item.decode("latin-1"))
        return "".join(parts)
    index = 2 if operator == b'"' else 0
    if len(operands) <= index:
        return ""
    item = operands[index]
    return item if isinstance(item, str) else bytes(item).decode("latin-1")


def _tj_kerning(operands) -> float:
    return sum(_number(i) for i in (operands[0] if operands else ()) if not isinstance(i, (str, bytes)))


@dataclass(frozen=True)
class LocatedShow:
    """One show-text operation: where it is, what it says, and where it sits in
    the stream.

    `index` is what makes masking possible -- the caller removes exactly those
    operations and leaves every other operator in the stream as the template
    wrote it. `base_font` travels alongside because measuring the text needs the
    font it was drawn in, and the chunk is the QA gate's type rather than this
    module's.
    """

    index: int
    chunk: TextChunk
    base_font: str | None


def located_show_operations(operations, page_number: int, fonts: dict):
    """Every show-text operation on one page, with where it starts.

    The walker tracks the graphics and text matrices rather than trusting
    `Tm`-per-line, because a producer is free to emit one text object holding a
    whole paragraph and move the pen with `Td` and `TJ` alone.
    """
    ctm = _IDENTITY
    stack: list = []
    tm = tlm = _IDENTITY
    font_size, base_font, leading = 0.0, None, 0.0
    char_spacing = word_spacing = 0.0
    horizontal_scale = 1.0

    for index, (operands, operator) in enumerate(operations):
        op = bytes(operator)
        if op == b"q":
            stack.append(ctm)
        elif op == b"Q":
            ctm = stack.pop() if stack else _IDENTITY
        elif op == b"cm" and len(operands) >= 6:
            ctm = _multiply(tuple(_number(v) for v in operands[:6]), ctm)
        elif op == b"BT":
            tm = tlm = _IDENTITY
        elif op == b"Tf" and len(operands) >= 2:
            base_font = fonts.get(str(operands[0]))
            font_size = _number(operands[1])
        elif op == b"Tm" and len(operands) >= 6:
            tm = tlm = tuple(_number(v) for v in operands[:6])
        elif op in (b"Td", b"TD") and len(operands) >= 2:
            if op == b"TD":
                leading = -_number(operands[1])
            tlm = _multiply(_translation(_number(operands[0]), _number(operands[1])), tlm)
            tm = tlm
        elif op == b"T*":
            tlm = _multiply(_translation(0.0, -leading), tlm)
            tm = tlm
        elif op == b"TL" and operands:
            leading = _number(operands[0])
        elif op == b"Tc" and operands:
            char_spacing = _number(operands[0])
        elif op == b"Tw" and operands:
            word_spacing = _number(operands[0])
        elif op == b"Tz" and operands:
            horizontal_scale = _number(operands[0]) / 100.0
        elif op in _SHOW_OPERATORS:
            if op in (b"'", b'"'):
                if op == b'"' and len(operands) >= 2:
                    word_spacing, char_spacing = _number(operands[0]), _number(operands[1])
                tlm = _multiply(_translation(0.0, -leading), tlm)
                tm = tlm
            text = _shown_text(operands, op)
            device = _multiply(tm, ctm)
            # The size a reader sees is the declared size scaled by whatever the
            # matrices do to it; a producer that sets 1pt text and scales it by
            # ten has drawn 10pt text.
            scale = math.hypot(device[2], device[3]) or 1.0
            yield LocatedShow(
                index=index,
                chunk=TextChunk(page=page_number, x=device[4], y=device[5],
                                text=text, font_size=font_size * scale),
                base_font=base_font,
            )
            advance = measure_text(text, font_size, base_font, char_spacing=char_spacing,
                                   word_spacing=word_spacing, horizontal_scale=horizontal_scale)
            if op == b"TJ":
                advance -= _tj_kerning(operands) / 1000.0 * font_size * horizontal_scale
            tm = _multiply(_translation(advance, 0.0), tm)


def page_chunks(page, page_number: int, container) -> tuple:
    """Every located show operation on `page`, as a tuple.

    A page whose content stream cannot be parsed yields nothing rather than
    raising: it is a page with no *readable* text, which the integrity gate then
    compares against a rendered page with no readable text, and a real change
    still shows up as a geometry difference.
    """
    contents = page.get_contents()
    if contents is None:
        return ()
    try:
        stream = ContentStream(contents, container)
    except Exception:
        return ()
    return tuple(found.chunk for found in located_show_operations(
        stream.operations, page_number, _font_map(page)))


def page_geometry(reader) -> tuple:
    """One `PageGeometry` per page, 1-based, in page order."""
    geometry = []
    for number, page in enumerate(reader.pages, start=1):
        box = page.mediabox
        geometry.append(PageGeometry(
            page=number,
            width=float(box.width),
            height=float(box.height),
            rotation=int(page.get(NameObject("/Rotate")) or 0),
        ))
    return tuple(geometry)


def scan_page_regions(pdf_path: str) -> tuple:
    """Derive a region inventory from a PDF, for a reviewer to confirm.

    §12 wants "stored dynamic regions / anchors / bounding boxes", and they have
    to come from somewhere the first time a template is onboarded. This is that
    somewhere -- one `PageRegion` per show operation, boxed from the text's
    origin, its measured width and the font's ascender and descender.

    It is deliberately *not* the production anchor. The doc's anchor table puts
    the bounding box "fixed to the page" and the manifest stores it; a box
    re-derived at render time would move whenever the extraction heuristics did,
    which is exactly the silent repointing §6 refuses for style-id anchors.
    """
    reader = PdfReader(pdf_path)
    regions = []
    for number, page in enumerate(reader.pages, start=1):
        contents = page.get_contents()
        if contents is None:
            continue
        stream = ContentStream(contents, reader)
        for found in located_show_operations(stream.operations, number, _font_map(page)):
            chunk = found.chunk
            if not chunk.text.strip():
                continue
            width = measure_text(chunk.text, chunk.font_size, found.base_font)
            regions.append(PageRegion(
                page=number,
                x0=chunk.x,
                y0=chunk.y + HELVETICA_DESCENDER * chunk.font_size,
                x1=chunk.x + width,
                y1=chunk.y + HELVETICA_ASCENDER * chunk.font_size,
                text=chunk.text,
            ))
    return tuple(regions)


# ------------------------------------------------------------- the renderer --

@dataclass(frozen=True)
class RegionFill:
    """One approved dynamic region and the exact text to put in it.

    `status` is not decoration. §12's third step is "mask placeholder/instruction
    region *when approved*", and §6's first lock rule is that every non-STATIC
    object carries APPROVED. So a fill that is still PROPOSED is refused here
    rather than rendered and flagged afterwards: a proposed mapping that reached
    a customer's letterhead is not a QA finding, it is a document nobody agreed
    to send.
    """

    object_id: str
    anchor: Anchor
    text: str
    status: str = APPROVED
    #: None means "inherit the size of the text this region replaces", which is
    #: the size the template's author approved.
    font_size: float | None = None
    align: str = ALIGN_LEFT
    padding: float = DEFAULT_PADDING
    #: False draws the value without removing what is underneath. Only correct
    #: for a region that is genuinely blank in the approved PDF -- a ruled line
    #: waiting for a name. Over a placeholder it leaves the placeholder text
    #: extractable, which is the redaction failure this renderer exists to avoid.
    mask: bool = True

    def __post_init__(self):
        if not isinstance(self.anchor, Anchor):
            raise TypeError(
                f"fill {self.object_id!r} needs a semantic_model.Anchor, got "
                f"{type(self.anchor).__name__}")
        if self.anchor.kind != ANCHOR_BBOX:
            raise PdfOverlayError(
                f"fill {self.object_id!r} carries a {self.anchor.kind!r} anchor. The doc's anchor "
                f"table gives PDF exactly one kind -- a region bounding box, 'immutable PDF "
                f"templates only' -- because there are no runs or fields on a page to address.")
        if not isinstance(self.text, str):
            raise TypeError(
                f"fill {self.object_id!r} must carry formatted text, got "
                f"{type(self.text).__name__}; formatting is `value_format`'s decision and the "
                "renderer must not make it a second time")
        if self.align not in ALIGNMENTS:
            raise ValueError(
                f"fill {self.object_id!r} declares align={self.align!r}; expected {list(ALIGNMENTS)}")
        if self.font_size is not None and self.font_size <= 0:
            raise ValueError(
                f"fill {self.object_id!r} declares font_size={self.font_size!r}; a size must be "
                "positive, and None is how a fill says 'inherit the approved one'")

    @property
    def box(self) -> tuple:
        return (self.anchor.x0, self.anchor.y0, self.anchor.x1, self.anchor.y1)

    @property
    def page(self) -> int:
        return self.anchor.page

    def overlaps(self, other: "RegionFill") -> bool:
        if self.page != other.page:
            return False
        ax0, ay0, ax1, ay1 = self.box
        bx0, by0, bx1, by1 = other.box
        return ax0 < bx1 and bx0 < ax1 and ay0 < by1 and by0 < ay1


@dataclass
class OverlayResult:
    """What §12's `RenderResult` carries, for the PDF path.

    Named to match `docx_renderer.FillResult` where they mean the same thing, so
    the batch runner and the audit writer read one shape rather than two.
    """

    output_path: str
    renderer: str = PDF_RENDERER_VERSION
    #: One entry per fill: which region, what went in, what came out from under
    #: it. This is the §17 "resolved field -> source lineage" row for a PDF.
    region_lineage: list = field(default_factory=list)
    qa_passed: bool = True
    qa_notes: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    #: `sha256:...` of the file written. Meaningful because the write is
    #: reproducible: the same template, fills and renderer produce it again.
    digest: str = ""


def _validate_fills(fills, page_count: int) -> list:
    """Everything that must be true before a single byte is written."""
    fills = list(fills)
    for fill in fills:
        if not isinstance(fill, RegionFill):
            raise TypeError(f"every fill must be a RegionFill, got {type(fill).__name__}")
        if fill.status != APPROVED:
            raise PdfOverlayError(
                f"fill {fill.object_id!r} has status {fill.status!r}. §12 masks and overlays a "
                "region only when it is approved, and §6 requires APPROVED on every non-STATIC "
                "object before a manifest can lock.")
        if not 1 <= fill.page <= page_count:
            raise PdfOverlayError(
                f"fill {fill.object_id!r} addresses page {fill.page} of a {page_count}-page PDF. "
                "The approved template is immutable, so this anchor was compiled against a "
                "different document.")

    seen: dict = {}
    for fill in fills:
        if fill.object_id in seen:
            raise PdfOverlayError(
                f"two fills both claim object_id {fill.object_id!r}; lineage would record one "
                "value for a region that received another")
        seen[fill.object_id] = fill

    for i, one in enumerate(fills):
        for other in fills[i + 1:]:
            if one.overlaps(other):
                raise PdfOverlayError(
                    f"fills {one.object_id!r} and {other.object_id!r} claim overlapping regions on "
                    f"page {one.page}. §6: no two objects may claim overlapping anchor ranges -- "
                    "whichever is drawn second destroys the other.")

    # Page first, then object id: the content stream this produces must not
    # depend on the order a caller happened to build its list in, or two renders
    # of the same manifest would differ in bytes while agreeing in every pixel.
    return sorted(fills, key=lambda f: (f.page, f.object_id))


def _baseline(fill: RegionFill, font_size: float, inherited: float | None) -> float:
    """Where the value's baseline goes.

    The placeholder's own baseline when there was a placeholder, because that is
    the position the template's author approved and the box around it is
    deliberately drawn a little larger so the anchor still resolves after a
    re-save. Bottom-aligning to the box instead would drop every value a few
    points below the line it belongs on -- correct content, visibly broken
    letter.

    A region that was blank in the approved PDF has no baseline to inherit, so
    the rule falls back to the box's bottom edge plus the font's descender,
    which puts the glyphs' feet on the bottom of the approved region.
    """
    if inherited is not None:
        return inherited
    return fill.anchor.y0 - HELVETICA_DESCENDER * font_size


def _pen_x(fill: RegionFill, width: float) -> float:
    """Where the pen starts, horizontally, inside the approved box.

    Taken from the box rather than inherited from the masked text, unlike the
    baseline. The two axes of a bounding box mean different things: vertically it
    is a band from ascender to descender and the baseline is a line inside it,
    so there is an approved line to inherit; horizontally it is the field's
    approved extent, which is exactly what alignment chooses a position within
    and what the overflow gate measures against.
    """
    x0, _y0, x1, _y1 = fill.box
    if fill.align == ALIGN_CENTRE:
        return (x0 + x1) / 2 - width / 2
    if fill.align == ALIGN_RIGHT:
        return x1 - fill.padding - width
    return x0 + fill.padding


def _overlay_stream(page_fills, font_name: str, sizes: dict, baselines: dict) -> bytes:
    """The operators appended to a page: a mask rectangle, then the value.

    One `q ... Q` pair per drawing, so the graphics state the template left
    behind is never inherited by the overlay and the overlay never leaks into
    whatever a later stream does.
    """
    out = bytearray()
    for fill in page_fills:
        x0, y0, x1, y1 = fill.box
        if fill.mask:
            out += (f"q 1.000 1.000 1.000 rg {_fmt(x0)} {_fmt(y0)} "
                    f"{_fmt(x1 - x0)} {_fmt(y1 - y0)} re f Q\n").encode("ascii")
    for fill in page_fills:
        if not fill.text:
            # A region whose approved value is empty is masked and left empty.
            # Drawing an empty string would add an operator that says nothing.
            continue
        size = sizes[fill.object_id]
        width = measure_text(fill.text, size)
        baseline = _baseline(fill, size, baselines[fill.object_id])
        out += (f"q BT {font_name} {_fmt(size)} Tf 0.000 0.000 0.000 rg "
                f"1 0 0 1 {_fmt(_pen_x(fill, width))} {_fmt(baseline)} Tm ").encode("ascii")
        out += _pdf_string(fill.text)
        out += b" Tj ET Q\n"
    return bytes(out)


def render_overlay(
    template_pdf: str,
    output_path: str,
    fills,
    *,
    page_regions,
    qa_policy=None,
    default_font_size: float = DEFAULT_FONT_SIZE,
) -> OverlayResult:
    """§12's PDF path, end to end.

    `page_regions` is the stored inventory the manifest pins -- the second step
    of the doc's flow -- and it has no default. Deriving it from the template at
    render time would mean the coordinates a value lands on were recomputed by
    the renderer rather than approved by a human, which is the whole distinction
    between an anchor and a guess. `scan_page_regions` produces a candidate
    inventory at onboarding time; a reviewer confirms it; the manifest stores it.

    Raises rather than returning a failed result for anything that makes the
    request itself wrong -- an unapproved fill, an anchor that no longer
    resolves, a value Helvetica cannot draw. A QA finding describes a document
    that was produced; these describe one that must not be.
    """
    reader = PdfReader(template_pdf)
    ordered = _validate_fills(fills, len(reader.pages))

    inventory = TemplateInventory(paragraph_texts=(), page_regions=tuple(page_regions))
    result = OverlayResult(output_path=str(output_path))

    # Step 2 of the flow, before anything is written: every anchor resolves
    # exactly once against the pinned inventory, or the render stops. `resolve`
    # raises AnchorNotFoundError for a page that was re-laid out and
    # AnchorAmbiguousError for a box that addresses two regions.
    resolved = {fill.object_id: resolve(fill.anchor, inventory) for fill in ordered}

    before_geometry = page_geometry(reader)
    before_chunks = tuple(
        chunk
        for number, page in enumerate(reader.pages, start=1)
        for chunk in page_chunks(page, number, reader)
    )

    writer = PdfWriter(clone_from=reader)
    by_page: dict = {}
    for fill in ordered:
        by_page.setdefault(fill.page, []).append(fill)

    overflows: list = []
    masked_placeholders: list = []
    for number, page_fills in sorted(by_page.items()):
        page = writer.pages[number - 1]
        contents = page.get_contents()
        if contents is None:
            raise PdfOverlayError(
                f"page {number} of {template_pdf} has no content stream, so there is nothing to "
                "overlay onto and nothing to mask; the approved template is not the document "
                "these anchors were compiled against")
        stream = ContentStream(contents, writer)
        found = list(located_show_operations(stream.operations, number, _font_map(page)))

        # Step 3: mask. The placeholder's operators are *removed*, not covered.
        masked_indices: set = set()
        masked_text: dict = {}
        for fill in page_fills:
            if not fill.mask:
                continue
            for item in found:
                if item.chunk.inside(fill.box):
                    masked_indices.add(item.index)
                    masked_text.setdefault(fill.object_id, []).append(item.chunk)

        if masked_indices:
            stream.operations = [
                operation for index, operation in enumerate(stream.operations)
                if index not in masked_indices
            ]
            page.replace_contents(stream)

        # Step 4: overlay, at the coordinates the manifest approved.
        sizes, baselines = {}, {}
        for fill in page_fills:
            covered = masked_text.get(fill.object_id, [])
            sizes[fill.object_id] = (
                fill.font_size or (covered[0].font_size if covered else None) or default_font_size)
            # Stream order rather than lowest-on-the-page: a region holding two
            # show operations holds one line split by a kern, and the first is
            # where the line starts.
            baselines[fill.object_id] = covered[0].y if covered else None

        font_name = _overlay_font(page)
        appended = DecodedStreamObject()
        appended.set_data(_overlay_stream(page_fills, font_name, sizes, baselines))
        merged = ContentStream(page.get_contents(), writer)
        merged.operations = merged.operations + ContentStream(appended, writer).operations
        page.replace_contents(merged)

        for fill in page_fills:
            size = sizes[fill.object_id]
            width = measure_text(fill.text, size)
            available = (fill.anchor.x1 - fill.anchor.x0) - 2 * fill.padding
            if width > available:
                overflows.append(
                    f"{fill.object_id}: {fill.text!r} needs {width:.1f} pt at {size:.1f} pt but the "
                    f"approved region on page {number} offers {available:.1f} pt")
            covered = masked_text.get(fill.object_id, [])
            stored = resolved[fill.object_id].matched_text
            if fill.mask and stored:
                # Only the placeholders a fill was meant to remove. The stored
                # inventory also holds labels ("Name:"), headings and
                # placeholders no fill targets in this render -- all of which are
                # supposed to survive, so checking the whole inventory flags a
                # correct document.
                masked_placeholders.append(stored)
            if fill.mask and stored and not any(stored in c.text or c.text in stored for c in covered):
                # A warning, not a block: the value still lands where the
                # manifest said, so the letter is correct even though the stored
                # inventory no longer matches the template. What must block is
                # narrower and is checked against the OUTPUT below -- whether the
                # placeholder is still readable on the finished page.
                result.warnings.append(
                    f"{fill.object_id}: the stored region says {stored!r} but the page at those "
                    f"coordinates carries {[c.text for c in covered] or 'nothing'}. The inventory "
                    "and the template have drifted apart; re-scan before the next batch.")
            result.region_lineage.append({
                "object_id": fill.object_id,
                "page": number,
                "box": list(fill.box),
                "text": fill.text,
                "font_size": size,
                "font_size_source": ("declared" if fill.font_size
                                     else "inherited" if covered else "default"),
                "baseline": _baseline(fill, size, baselines[fill.object_id]),
                "baseline_source": "inherited" if covered else "region",
                "masked_text": [c.text for c in covered],
                "stored_region_text": stored,
            })

    writer.write(str(output_path))

    # Step 5: validate against the file that was actually written, not against
    # the object graph that was meant to produce it.
    written = PdfReader(str(output_path))
    after_chunks = tuple(
        chunk
        for number, page in enumerate(written.pages, start=1)
        for chunk in page_chunks(page, number, written)
    )
    approved_regions = tuple(
        PageRegion(page=f.page, x0=f.anchor.x0, y0=f.anchor.y0, x1=f.anchor.x1, y1=f.anchor.y1,
                   text=f.text)
        for f in ordered
    )
    policy = resolve_policy(qa_policy)
    outcome = route(
        pdf_findings(
            before_geometry=before_geometry,
            after_geometry=page_geometry(written),
            before_chunks=before_chunks,
            after_chunks=after_chunks,
            approved_regions=approved_regions,
            overflows=overflows,
            policy=policy,
        ),
        policy,
    )
    result.qa_notes.extend(outcome.notes)
    result.qa_passed = outcome.passed

    # §17: "Placeholder remains | Block document". The DOCX path has blocked on
    # this from the beginning; the overlay path had no equivalent, so a mask that
    # matched nothing shipped a letter with `<Employee Name>` still printed on it
    # and reported qa_passed True. Checked against the written output rather than
    # against whether masking found anything, because those are different facts:
    # a stale region over a template that no longer has the placeholder is a
    # drifted inventory (a warning, above), while a placeholder legible on the
    # finished page is a defective document.
    survived = sorted({
        text for text in masked_placeholders
        if any(text in chunk.text for chunk in after_chunks)
    })
    if survived:
        placeholder_outcome = route(
            findings_for(
                PLACEHOLDER_REMAINS,
                [f"{text!r} is still readable on the finished page" for text in survived],
                policy,
            ),
            policy,
        )
        result.qa_notes.extend(placeholder_outcome.notes)
        result.qa_passed = result.qa_passed and placeholder_outcome.passed
    result.digest = content_digest(str(output_path))
    return result


# --------------------------------------------------------------------------
# The review-queue preview: a picture of the letter, not the letter.
#
# The `.docx` is the deliverable and Word is the authority on how it paginates.
# This renders a preview so a reviewer can see the letter without opening Word,
# and the distinction matters: LibreOffice and Word break pages differently, so
# a preview that shows five pages where Word shows four is a renderer difference
# and not a defect. The document's layout properties are never touched by the
# engine, which is why that difference is cosmetic.
#
# CJK is the operational trap. Without the Noto CJK fonts installed, LibreOffice
# renders every Chinese, Japanese and Korean glyph as a tofu box -- a preview
# that looks catastrophically broken while the `.docx` is perfectly correct. A
# reviewer seeing that will reject good letters, so the missing font is reported
# as a preview problem rather than left to look like a document problem.
#
# Nothing above this line is involved: an overlay render never converts anything.
# --------------------------------------------------------------------------

SOFFICE_CANDIDATES = (
    "soffice",
    "libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
    "/usr/lib/libreoffice/program/soffice",
)

# Scripts whose glyphs come from a font pack that is not installed by default.
CJK_SCRIPT_LANGUAGES = ("zh", "ja", "ko")
# Linux, Windows *and* macOS. The list used to name only the first two, so on a
# Mac -- which ships PingFang, Heiti, Songti and Hiragino -- `cjk_fonts_installed`
# reported no CJK fonts at all while `fc-list` was reporting forty-eight of them.
CJK_FONT_HINTS = (
    "noto sans cjk", "noto serif cjk", "source han", "simsun", "msmincho", "malgun",
    "pingfang", "heiti", "songti", "stsong", "stheiti", "hiragino", "yu gothic", "meiryo",
)


def _cjk_chars(text: str) -> set:
    """The CJK ideographs in a string. Kana and Hangul come along for the ride."""
    return {
        ch for ch in text
        if "一" <= ch <= "鿿"      # CJK unified ideographs
        or "぀" <= ch <= "ヿ"      # hiragana + katakana
        or "가" <= ch <= "힯"      # hangul syllables
    }

RENDER_TIMEOUT_SECONDS = 120


class PreviewUnavailable(RuntimeError):
    """No renderer on this host. A preview is a convenience, never a gate."""


@dataclass
class PreviewResult:
    pdf_path: str
    renderer: str
    notes: list = field(default_factory=list)


def find_soffice() -> str | None:
    for candidate in SOFFICE_CANDIDATES:
        resolved = shutil.which(candidate) or (candidate if Path(candidate).exists() else None)
        if resolved:
            return resolved
    return None


def cjk_fonts_installed() -> bool:
    """Best-effort: `fc-list` is present wherever LibreOffice renders on Linux."""
    fc = shutil.which("fc-list")
    if not fc:
        return True  # cannot tell; do not cry wolf
    try:
        out = subprocess.run([fc], capture_output=True, text=True, timeout=20).stdout.lower()
    except (OSError, subprocess.SubprocessError):
        return True
    return any(hint in out for hint in CJK_FONT_HINTS)


def render_pdf(docx_path: str | Path, out_dir: str | Path, language: str = "en") -> PreviewResult:
    """Render `docx_path` to a PDF beside it. Raises `PreviewUnavailable` if it cannot."""
    soffice = find_soffice()
    if not soffice:
        raise PreviewUnavailable(
            "LibreOffice is not installed on this host, so no PDF preview can be produced. "
            "The .docx is unaffected."
        )

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    notes: list[str] = []

    if language.split("-")[0].lower() in CJK_SCRIPT_LANGUAGES and not cjk_fonts_installed():
        notes.append(
            "No CJK font pack found on this host. The preview will show empty boxes where the "
            "letter has Chinese, Japanese or Korean text; the .docx itself is correct. "
            "Install fonts-noto-cjk on the rendering host."
        )

    # A private profile per run: concurrent conversions sharing the default one
    # deadlock, and a worker pool is exactly the shape that trips it.
    with tempfile.TemporaryDirectory() as profile:
        proc = subprocess.run(
            [
                soffice, "--headless", "--norestore", "--nolockcheck",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path),
            ],
            capture_output=True, text=True, timeout=RENDER_TIMEOUT_SECONDS,
        )

    pdf = out_dir / (Path(docx_path).stem + ".pdf")
    if not pdf.exists():
        # The converter's output names paths on this host; it is for the log.
        log.error("PDF conversion produced no file (exit %s): %s",
                  proc.returncode, (proc.stderr or proc.stdout)[:2000])
        raise PreviewUnavailable(
            "The PDF conversion did not produce a file. The .docx is unaffected."
        )

    _refuse_if_script_was_lost(docx_path, pdf)
    scrub_pdf_metadata(pdf, docx_path)

    notes.append(
        "Page breaks are LibreOffice's. Word is the authority on pagination for the .docx, "
        "which this preview does not modify."
    )
    return PreviewResult(pdf_path=str(pdf), renderer=soffice, notes=notes)


def _docx_identity(docx_path) -> dict:
    """The title and author the source document itself declares, if any."""
    try:
        from docx import Document

        core = Document(str(docx_path)).core_properties
    except Exception:  # noqa: BLE001 - an unreadable source just has no identity
        return {}
    identity = {}
    if (core.title or "").strip():
        identity["/Title"] = core.title.strip()
    if (core.author or "").strip():
        identity["/Author"] = core.author.strip()
    return identity


def scrub_pdf_metadata(pdf_path, docx_path=None) -> None:
    """Rewrite `pdf_path` so its metadata describes the document, not the toolchain.

    The converter stamps every PDF with its own name and build ("LibreOffice
    26.8.0.3 (AARCH64)", Creator "Writer"), a creation date in the server's
    time zone, and an XMP packet repeating all of it. None of that is about the
    letter, and all of it tells the recipient how the letter was produced. So
    the document information dictionary is replaced with the title and author
    the source .docx declares (nothing else, and no dates), the XMP metadata
    stream is removed from the catalogue, and the objects nothing refers to any
    more are dropped so the old packet is not still sitting in the file.

    Pages, the structure tree and the outline are untouched: this rewrites the
    envelope, never what is drawn.
    """
    reader = PdfReader(str(pdf_path))
    writer = PdfWriter(clone_from=reader)
    writer._root_object.pop(NameObject("/Metadata"), None)
    writer.metadata = _docx_identity(docx_path) if docx_path else {}
    writer.compress_identical_objects(remove_identicals=False, remove_orphans=True)
    staged = Path(str(pdf_path) + ".scrubbing")
    try:
        writer.write(str(staged))
        staged.replace(pdf_path)
    finally:
        staged.unlink(missing_ok=True)


def _refuse_if_script_was_lost(docx_path: str | Path, pdf: Path) -> None:
    """Read the PDF back and refuse it if the conversion ate the letter.

    A missing font does not make LibreOffice fail. It makes it substitute, and
    the substitute has no Chinese glyphs, so every ideograph in the document
    comes out as the same wrong character:

        隐隐隐 Gao Yan隐
        隐隐隐隐隐隐隐隐隐隐隐隐隐隐 Promotion 隐隐隐隐隐隐隐隐隐隐隐隐隐

    Exit code zero, a plausible file size, a PDF that opens. The Latin text is
    perfect, which is what makes it convincing -- the name, the salary and the
    dates are all correct and every word around them is destroyed.

    The guard that was supposed to catch this could not, three times over. It
    keyed on the document's declared `language`, and this letter is Chinese
    content generated under a project tagged `en`. Its font list knew only Linux
    and Windows names, so a Mac with forty-eight CJK fonts reported none. And it
    only ever appended a *note*, which the download endpoint discards.

    So this asks the question directly instead of predicting the answer: the
    source has CJK, does the output? It costs one PDF text extraction, it cannot
    false-positive on a host where the conversion works, and it catches every
    cause rather than the one cause somebody thought of -- missing font, broken
    fallback, bad substitution map.

    Raises `PreviewUnavailable`, which the API already turns into a 503 naming
    what is missing. Refusing is the right answer: a letter whose entire body is
    one repeated glyph is not a preview with a caveat, it is a wrong document
    that looks like a right one, and handing it to somebody to send is worse
    than handing them nothing.
    """
    try:
        from docx import Document
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - both are hard dependencies
        return

    try:
        source = Document(str(docx_path))
        source_text = "\n".join(p.text for p in source.paragraphs)
        for table in source.tables:
            for row in table.rows:
                source_text += "\n" + "\n".join(c.text for c in row.cells)
    except Exception:  # noqa: BLE001 - an unreadable source is not this check's business
        return

    wanted = _cjk_chars(source_text)
    if not wanted:
        return

    try:
        rendered = "\n".join((page.extract_text() or "") for page in PdfReader(str(pdf)).pages)
    except Exception:  # noqa: BLE001 - a PDF we cannot read back is not proof of anything
        return

    got = _cjk_chars(rendered)
    # Compared as *distinct* characters, not as a count. Tofu substitution
    # preserves the character count exactly -- one wrong glyph per right one --
    # so a length check sees nothing wrong. What collapses is the variety.
    kept = len(wanted & got) / len(wanted)
    if kept >= 0.5:
        return

    missing_fonts = _requested_fonts_not_installed(docx_path)
    detail = (
        f" The template asks for {', '.join(sorted(missing_fonts))}, which "
        f"{'is' if len(missing_fonts) == 1 else 'are'} not installed on this host."
        if missing_fonts else ""
    )
    raise PreviewUnavailable(
        "This document's Chinese, Japanese or Korean text did not survive the PDF conversion — "
        f"only {kept:.0%} of its characters came through, and the rest were replaced with a "
        f"placeholder glyph.{detail} The .docx is correct and unaffected; download that instead, "
        "or install the font on the machine running the API and try again."
    )


def _requested_fonts_not_installed(docx_path: str | Path) -> set:
    """Which fonts the .docx asks for that this host does not have.

    Best effort, and only ever used to make the refusal above more specific --
    naming `MingLiU` is the difference between a message somebody can act on and
    one they can only forward.
    """
    import re
    import zipfile

    fc = shutil.which("fc-list")
    if not fc:
        return set()
    try:
        installed = subprocess.run(
            [fc, "--format", "%{family}\n"], capture_output=True, text=True, timeout=20).stdout.lower()
        with zipfile.ZipFile(str(docx_path)) as z:
            xml = "".join(
                z.read(name).decode("utf8", "replace")
                for name in ("word/document.xml", "word/styles.xml") if name in z.namelist())
    except (OSError, subprocess.SubprocessError, zipfile.BadZipFile, KeyError):
        return set()

    requested = set()
    for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
        requested |= set(re.findall(rf'w:{attr}="([^"]+)"', xml))
    # `w:eastAsia` also carries language tags like "zh-CN" in some producers,
    # and a language tag is not a missing font.
    return {
        f for f in requested
        if f.lower() not in installed and not re.fullmatch(r"[a-z]{2}-[A-Z]{2}", f)
    }
