"""Source-file extraction + chunking (spec §9). Simplified vs. full spec: no OCR
fallback for scanned PDFs, no PPTX support in this MVP pass -- docx/pdf/xlsx/csv/txt
covers everything the frontend's upload dialogs advertise except pptx, which is a
documented gap (raises a clear ingest_error instead of silently no-op'ing).
"""

import csv
import hashlib
import io
from dataclasses import dataclass
from datetime import date, datetime


@dataclass
class RawChunk:
    text: str
    element_type: str
    heading_path: str | None = None


def _split_paragraph(text: str, max_chars: int = 700) -> list[str]:
    text = text.strip()
    if not text:
        return []
    if len(text) <= max_chars:
        return [text]
    parts, current = [], []
    length = 0
    for sentence in text.replace("\n", " ").split(". "):
        sentence = sentence.strip()
        if not sentence:
            continue
        if length + len(sentence) > max_chars and current:
            parts.append(". ".join(current) + ".")
            current, length = [], 0
        current.append(sentence)
        length += len(sentence)
    if current:
        parts.append(". ".join(current))
    return parts


def extract_docx(path: str) -> list[RawChunk]:
    import docx

    document = docx.Document(path)
    chunks: list[RawChunk] = []
    heading_path: list[str] = []
    for p in document.paragraphs:
        style = (p.style.name if p.style else "") or ""
        if style.lower().startswith("heading"):
            level = int(style[-1]) if style[-1].isdigit() else 1
            heading_path = heading_path[: level - 1] + [p.text.strip()]
            continue
        for part in _split_paragraph(p.text):
            chunks.append(RawChunk(text=part, element_type="paragraph", heading_path=" > ".join(heading_path) or None))
    for t_idx, table in enumerate(document.tables):
        rows = [[c.text.strip() for c in row.cells] for row in table.rows]
        text = "\n".join(" | ".join(r) for r in rows)
        if text.strip():
            chunks.append(RawChunk(text=text, element_type="table", heading_path=" > ".join(heading_path) or None))
    return chunks


def extract_pdf(path: str) -> list[RawChunk]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    chunks: list[RawChunk] = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for part in _split_paragraph(text):
            chunks.append(RawChunk(text=part, element_type="paragraph", heading_path=f"page {page_no}"))
    if not chunks:
        raise ValueError("No extractable text found (the PDF may be a scanned image without OCR support in this build).")
    return chunks


def extract_xlsx(path: str) -> list[RawChunk]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    chunks: list[RawChunk] = []
    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header, *body = rows
        header = [str(h) if h is not None else "" for h in header]
        for row in body:
            record = {header[i]: row[i] for i in range(min(len(header), len(row)))}
            text = "; ".join(f"{k}: {v}" for k, v in record.items() if v not in (None, ""))
            if text:
                chunks.append(RawChunk(text=text, element_type="sheet_rows", heading_path=sheet.title))
    return chunks


def extract_csv(raw_bytes: bytes) -> list[RawChunk]:
    text = raw_bytes.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []
    header, *body = rows
    chunks: list[RawChunk] = []
    for row in body:
        record = {header[i]: row[i] for i in range(min(len(header), len(row)))}
        chunk_text = "; ".join(f"{k}: {v}" for k, v in record.items() if v)
        if chunk_text:
            chunks.append(RawChunk(text=chunk_text, element_type="sheet_rows"))
    return chunks


def extract_txt(raw_bytes: bytes) -> list[RawChunk]:
    text = raw_bytes.decode("utf-8", errors="ignore")
    return [RawChunk(text=part, element_type="paragraph") for para in text.split("\n\n") for part in _split_paragraph(para)]


def extract(path: str, file_type: str) -> list[RawChunk]:
    if file_type == "docx":
        return extract_docx(path)
    if file_type == "pdf":
        return extract_pdf(path)
    if file_type == "xlsx":
        return extract_xlsx(path)
    if file_type == "csv":
        return extract_csv(open(path, "rb").read())
    if file_type in ("txt", "html", "json"):
        return extract_txt(open(path, "rb").read())
    raise ValueError(f"Unsupported source file type for ingestion: {file_type}")


def content_sha256(text: str) -> str:
    return hashlib.sha256(text.encode()).hexdigest()


# --------------------------------------------------------------------------- records
# The chunk path above exists for retrieval: it flattens a spreadsheet into
# "col: val; col2: val2" strings that a TF-IDF index can rank. That shape is
# lossy on purpose and is the wrong input for template filling, where one row
# is one document and every cell has to be addressable by its column name.
# `extract_records` is the structured counterpart -- same files, different
# contract -- and it is the seam an external extraction service (Azure Document
# Intelligence, AWS Bedrock Data Automation) would slot in behind if scanned
# sources ever arrive.

RECORD_FILE_TYPES = {"xlsx", "csv"}


def _clean_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # openpyxl reads whole numbers as 38.0
    # A date-formatted cell arrives from openpyxl as a real `datetime`, and
    # `str()` of one renders the midnight it never had:
    #
    #     变动将于2026-09-01 00:00:00生效
    #
    # in a letter that asked for a date. This is the layer where it has to be
    # fixed: everything downstream -- binding, conditions, `format_value` -- is
    # handed the string this function returns, so by the time a formatter could
    # notice the value was a date, it is already text with a time stuck to it.
    #
    # A datetime at midnight is a date; every date-formatted spreadsheet cell is
    # one. ISO, because this is the normalised form the rest of the pipeline
    # parses -- `value_format.parse_date` reads it, so a field the compiler typed
    # as a date still renders in the reader's locale from here.
    #
    # A datetime carrying an actual time keeps it. Dropping that would be
    # inventing a fact rather than normalising one; what goes away is only the
    # midnight nobody meant.
    if isinstance(value, datetime):
        return (value.date().isoformat() if not (value.hour or value.minute or value.second)
                else value.isoformat(sep=" ", timespec="minutes"))
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _rows_to_records(rows: list[list], sheet: str | None = None) -> tuple[list[str], list[dict]]:
    if not rows:
        return [], []
    header_row, *body = rows
    columns, seen = [], {}
    for index, raw in enumerate(header_row):
        name = _clean_cell(raw) or f"column_{index + 1}"
        # Duplicate headers are common in exported spreadsheets and would
        # otherwise silently overwrite each other in the record dict.
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        columns.append(name)

    records = []
    for row_index, row in enumerate(body):
        values = {columns[i]: _clean_cell(row[i]) for i in range(min(len(columns), len(row)))}
        if not any(values.values()):
            continue  # trailing blank rows
        records.append({"_row_index": row_index, "_sheet": sheet, **values})
    return columns, records


def extract_records(path: str, file_type: str, sheet: str | None = None) -> tuple[list[str], list[dict]]:
    """Read a tabular source into (column_names, records) where each record is
    one row keyed by column name. One record produces one document."""
    if file_type == "xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
        worksheet = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in worksheet.iter_rows(values_only=True)]
        return _rows_to_records(rows, worksheet.title)

    if file_type == "csv":
        text = open(path, "rb").read().decode("utf-8", errors="ignore")
        return _rows_to_records(list(csv.reader(io.StringIO(text))))

    raise ValueError(
        f"Cannot read tabular records from a {file_type} file. "
        f"Supported: {', '.join(sorted(RECORD_FILE_TYPES))}."
    )


def sheet_names(path: str, file_type: str) -> list[str]:
    if file_type != "xlsx":
        return []
    import openpyxl

    return openpyxl.load_workbook(path, read_only=True).sheetnames
