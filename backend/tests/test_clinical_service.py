"""The clinical service, as a caller meets it: a study book, per-type
numbering, and one endpoint that turns a published clinical template plus
typed values into a numbered, stored, downloadable study document.

The flow under test is the whole vertical: kit -> blueprint -> publish ->
/clinical-documents:generate. Everything downstream of the manifest is the
machinery the rest of the product already trusts (fill engine, QA gates,
document trail); what these tests pin is the part that makes a clinical
document a clinical document -- the study snapshot that keeps saying what it
said, the per-type sequence that cannot repeat a number, and the derived
totals that either sum every row or honestly stay blank.
"""

import pytest

from app.clinical.derivations import derive_row_totals


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def org_a(two_orgs):
    token_a, project_a, _token_b, _project_b = two_orgs
    return token_a, project_a


@pytest.fixture
def published_csr_manifest(app_client, org_a):
    """A published clinical-kit blueprint and its manifest id."""
    token, project_id = org_a
    made = app_client.post(
        "/api/v1/template-blueprints", headers=_auth(token),
        json={"name": "Studio CSR", "kit": "clinical", "project_id": project_id})
    assert made.status_code == 201, made.text
    blueprint_id = made.json()["id"]

    published = app_client.post(
        f"/api/v1/template-blueprints/{blueprint_id}:publish", headers=_auth(token),
        json={"recompile": False})
    assert published.status_code == 201, published.text
    return token, project_id, published.json()["manifest_id"]


DISPOSITION_ROWS = [
    {"site_name": "Pune General", "subjects_enrolled": 80,
     "subjects_completed": 70, "subjects_withdrawn": 10},
    {"site_name": "Mumbai Central", "subjects_enrolled": 40,
     "subjects_completed": 34, "subjects_withdrawn": 6},
]

STUDY = {
    "protocol_number": "ONC-2026-014", "title": "A Phase 2 study of drug X",
    "sponsor": "Acme Pharma", "phase": "2", "indication": "NSCLC",
    "principal_investigator": "Dr. A. Rao",
}

FIELDS = {
    "efficacy_summary": "The primary endpoint was met.",
    "safety_summary": "No new safety signals were observed.",
    "conclusions": "The results support continued development.",
}

COLUMNS = [
    {"source_key": "site_name", "type": "string"},
    {"source_key": "subjects_enrolled", "type": "number"},
    {"source_key": "subjects_completed", "type": "number"},
    {"source_key": "subjects_withdrawn", "type": "number"},
]


# ---------------------------------------------------------------- derivations

def test_totals_sum_every_numeric_column():
    out = derive_row_totals(DISPOSITION_ROWS, COLUMNS)
    assert out["row_count"] == 2
    assert out["total_subjects_enrolled"] == "120"
    assert out["total_subjects_completed"] == "104"
    assert out["total_subjects_withdrawn"] == "16"
    assert "total_site_name" not in out  # string columns never sum


def test_a_partial_column_stays_honestly_blank():
    """One blank cell and the column's total does not fill at all -- a partial
    sum on a study report is a wrong number that looks deliberate."""
    rows = [dict(DISPOSITION_ROWS[0]), {"site_name": "Delhi", "subjects_enrolled": ""}]
    out = derive_row_totals(rows, COLUMNS)
    assert "total_subjects_enrolled" not in out
    assert out["row_count"] == 2


def test_totals_are_plain_decimal_notation_never_scientific():
    rows = [{"n": 100}, {"n": 20}]
    out = derive_row_totals(rows, [{"source_key": "n", "type": "number"}])
    assert out["total_n"] == "120"  # Decimal("120").normalize() prints 1.2E+2


def test_no_rows_means_a_count_of_zero_and_no_totals():
    out = derive_row_totals([], COLUMNS)
    assert out == {"row_count": 0}


# ---------------------------------------------------------------- study book

def test_study_book_crud(app_client, org_a):
    token, _ = org_a
    made = app_client.post("/api/v1/studies", headers=_auth(token), json=STUDY)
    assert made.status_code == 201, made.text
    study_id = made.json()["id"]

    got = app_client.get(f"/api/v1/studies/{study_id}", headers=_auth(token))
    assert got.json()["sponsor"] == "Acme Pharma"

    patched = app_client.patch(f"/api/v1/studies/{study_id}", headers=_auth(token),
                               json={"phase": "3"})
    assert patched.json()["phase"] == "3"
    assert patched.json()["protocol_number"] == "ONC-2026-014"  # untouched fields stay

    listed = app_client.get("/api/v1/studies", headers=_auth(token),
                            params={"q": "onc-2026"})
    assert any(s["id"] == study_id for s in listed.json()["items"])

    gone = app_client.delete(f"/api/v1/studies/{study_id}", headers=_auth(token))
    assert gone.json()["deleted"] is True
    assert app_client.get(f"/api/v1/studies/{study_id}",
                          headers=_auth(token)).status_code == 404


def test_study_needs_a_protocol_number(app_client, org_a):
    token, _ = org_a
    res = app_client.post("/api/v1/studies", headers=_auth(token),
                          json={"protocol_number": "  "})
    assert res.status_code == 422
    assert res.json()["detail"]["error"]["code"] == "STUDY_NEEDS_PROTOCOL"


# ---------------------------------------------------------------- the flow

def test_generate_csr_end_to_end(app_client, published_csr_manifest):
    token, project_id, manifest_id = published_csr_manifest

    study = app_client.post("/api/v1/studies", headers=_auth(token), json=STUDY).json()

    res = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study_id": study["id"],
        "rows": DISPOSITION_ROWS, "fields": FIELDS,
        "document_date": "2026-09-08", "version_label": "1.0",
    })
    assert res.status_code == 201, res.text
    body = res.json()

    # The number came from the org's csr sequence, the totals from Decimal.
    assert body["number"] == "CSR-0001"
    assert body["document_type"] == "csr"
    assert body["qa_passed"] is True, body["qa_notes"]
    assert body["status"] == "final"  # the admin fixture can approve
    assert body["study"]["sponsor"] == "Acme Pharma"

    # The stored document is real and carries the rendered rows and totals.
    import docx as docx_lib

    from app.db import SessionLocal
    from app.models import DocumentVersion
    from app.storage import abs_path

    db = SessionLocal()
    try:
        blob = db.get(DocumentVersion, body["document_version_id"]).blob_path
    finally:
        db.close()
    document = docx_lib.Document(str(abs_path(blob)))
    table = document.tables[0]
    assert len(table.rows) == 3  # header + 2 sites
    assert table.rows[1].cells[0].text == "Pune General"
    text = "\n".join(p.text for p in document.paragraphs)
    assert "CSR-0001" in text
    assert "Acme Pharma" in text
    assert "120" in text  # the derived enrolment total, printed after the table
    assert "The primary endpoint was met." in text

    # The registry row can be read back, snapshot, record and all.
    detail = app_client.get(f"/api/v1/clinical-documents/{body['id']}",
                            headers=_auth(token)).json()
    assert detail["study"]["protocol_number"] == "ONC-2026-014"
    assert detail["source_record"]["document_number"] == "CSR-0001"
    assert detail["source_record"]["total_subjects_enrolled"] == "120"


def test_each_document_type_numbers_its_own_sequence(app_client, org_a):
    """CSR-0001 and PA-0001 come from different counters -- generating one
    must not advance the other."""
    token, project_id = org_a

    def manifest_for(kit):
        made = app_client.post(
            "/api/v1/template-blueprints", headers=_auth(token),
            json={"name": f"{kit} template", "kit": kit, "project_id": project_id})
        published = app_client.post(
            f"/api/v1/template-blueprints/{made.json()['id']}:publish",
            headers=_auth(token), json={"recompile": False})
        return published.json()["manifest_id"]

    csr_manifest = manifest_for("clinical")

    def generate_csr():
        res = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
            "manifest_id": csr_manifest, "project_id": project_id,
            "document_type": "csr", "study": STUDY,
            "rows": DISPOSITION_ROWS, "fields": FIELDS, "version_label": "1.0"})
        assert res.status_code == 201, res.text
        return res.json()["number"]

    first_csr = generate_csr()
    assert first_csr.startswith("CSR-")

    amendment = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_for("clinical_protocol"), "project_id": project_id,
        "document_type": "protocol_amendment", "study": STUDY,
        "rows": [{"section_reference": "5.1", "previous_text": "8 visits",
                  "revised_text": "6 visits"}],
        "fields": {"amendment_rationale": "Fewer visits, same endpoints."},
        "version_label": "2.0"})
    assert amendment.status_code == 201, amendment.text
    assert amendment.json()["number"].startswith("PA-")

    # The amendment advanced its own counter only: the next CSR is exactly one
    # step on, however many documents this session has already numbered.
    second_csr = generate_csr()
    assert (int(second_csr.split("-")[1])
            == int(first_csr.split("-")[1]) + 1), (first_csr, second_csr)


def test_a_fill_crash_after_allocation_rolls_the_number_back(
        app_client, published_csr_manifest, monkeypatch):
    token, project_id, manifest_id = published_csr_manifest
    payload = {
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study": STUDY,
        "rows": DISPOSITION_ROWS, "fields": FIELDS, "version_label": "1.0"}

    first = app_client.post("/api/v1/clinical-documents:generate",
                            headers=_auth(token), json=payload)
    assert first.status_code == 201, first.text

    import app.generation.single as single_module

    def explode(*_args, **_kwargs):
        raise RuntimeError("renderer crashed mid-fill")

    monkeypatch.setattr(single_module, "fill_template", explode)
    crashed = app_client.post("/api/v1/clinical-documents:generate",
                              headers=_auth(token), json=payload)
    assert crashed.status_code == 422
    assert crashed.json()["detail"]["error"]["code"] == "FILL_FAILED"
    monkeypatch.undo()

    after = app_client.post("/api/v1/clinical-documents:generate",
                            headers=_auth(token), json=payload)
    n_first = int(first.json()["number"].split("-")[1])
    n_after = int(after.json()["number"].split("-")[1])
    assert n_after == n_first + 1, "the crashed generation burned a number"


def test_a_required_table_with_no_rows_is_refused(app_client, published_csr_manifest):
    token, project_id, manifest_id = published_csr_manifest
    res = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study": STUDY, "rows": [], "fields": FIELDS})
    assert res.status_code == 422
    assert res.json()["detail"]["error"]["code"] == "CLINICAL_NEEDS_ROWS"


def test_a_document_about_no_study_is_refused(app_client, published_csr_manifest):
    token, project_id, manifest_id = published_csr_manifest
    res = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "rows": DISPOSITION_ROWS, "fields": FIELDS})
    assert res.status_code == 422
    assert res.json()["detail"]["error"]["code"] == "CLINICAL_NEEDS_STUDY"


def test_an_unknown_document_type_is_refused(app_client, published_csr_manifest):
    token, project_id, manifest_id = published_csr_manifest
    res = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "lab_report", "study": STUDY,
        "rows": DISPOSITION_ROWS, "fields": FIELDS})
    assert res.status_code == 422
    assert res.json()["detail"]["error"]["code"] == "CLINICAL_BAD_DOCUMENT_TYPE"


def test_an_optional_visit_table_may_stay_empty(app_client, org_a):
    """The ICF kit's visit schedule is optional (REMOVE_TABLE): a consent form
    with no scheduled visits generates cleanly, table and all gone."""
    token, project_id = org_a
    made = app_client.post(
        "/api/v1/template-blueprints", headers=_auth(token),
        json={"name": "ICF", "kit": "clinical_icf", "project_id": project_id})
    published = app_client.post(
        f"/api/v1/template-blueprints/{made.json()['id']}:publish",
        headers=_auth(token), json={"recompile": False})
    manifest_id = published.json()["manifest_id"]

    res = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "icf", "study": STUDY, "rows": [],
        "fields": {"purpose_description": "To learn whether drug X helps.",
                   "procedures_description": "You will take one tablet daily.",
                   "risks_description": "Headache and nausea are possible.",
                   "benefits_description": "Your condition may improve.",
                   "participant_name": "____________________"},
        "version_label": "1.0"})
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["number"] == "ICF-0001"
    assert body["qa_passed"] is True, body["qa_notes"]


def test_void_keeps_the_number(app_client, published_csr_manifest):
    token, project_id, manifest_id = published_csr_manifest
    made = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study": STUDY,
        "rows": DISPOSITION_ROWS, "fields": FIELDS, "version_label": "1.0"}).json()

    voided = app_client.post(f"/api/v1/clinical-documents/{made['id']}:void",
                             headers=_auth(token))
    assert voided.status_code == 200
    assert voided.json()["status"] == "void"
    assert voided.json()["number"] == made["number"]  # the number is not reused

    again = app_client.post(f"/api/v1/clinical-documents/{made['id']}:void",
                            headers=_auth(token))
    assert again.status_code == 409

    listed = app_client.get("/api/v1/clinical-documents", headers=_auth(token),
                            params={"status": "void"})
    assert any(d["id"] == made["id"] for d in listed.json()["items"])


def test_void_needs_sign_off_authority(app_client, published_csr_manifest):
    """A role that cannot approve a study document has no business cancelling
    one."""
    token, project_id, manifest_id = published_csr_manifest
    made = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study": STUDY,
        "rows": DISPOSITION_ROWS, "fields": FIELDS, "version_label": "1.0"}).json()

    from app.db import SessionLocal
    from app.models import User
    from app.security import create_access_token, hash_password

    db = SessionLocal()
    try:
        admin = db.query(User).filter(User.email == "user-a@tenant.test").one()
        clerk = db.query(User).filter(User.email == "coordinator-a@tenant.test").one_or_none()
        if clerk is None:
            clerk = User(org_id=admin.org_id, email="coordinator-a@tenant.test",
                         full_name="Coordinator A", password_hash=hash_password("pw"),
                         role_key="generator")
            db.add(clerk)
            db.commit()
        clerk_token = create_access_token(clerk.id, clerk.org_id)
    finally:
        db.close()

    refused = app_client.post(f"/api/v1/clinical-documents/{made['id']}:void",
                              headers=_auth(clerk_token))
    assert refused.status_code == 403

    allowed = app_client.post(f"/api/v1/clinical-documents/{made['id']}:void",
                              headers=_auth(token))
    assert allowed.status_code == 200


def test_the_other_tenant_cannot_see_clinical_documents(
        app_client, published_csr_manifest, two_orgs):
    token, project_id, manifest_id = published_csr_manifest
    _token_a, _pa, token_b, _pb = two_orgs

    made = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study": STUDY,
        "rows": DISPOSITION_ROWS, "fields": FIELDS, "version_label": "1.0"}).json()

    assert app_client.get(f"/api/v1/clinical-documents/{made['id']}",
                          headers=_auth(token_b)).status_code == 404
    listed_b = app_client.get("/api/v1/clinical-documents",
                              headers=_auth(token_b)).json()["items"]
    assert not any(d["id"] == made["id"] for d in listed_b)


def test_the_snapshot_survives_a_later_study_edit(app_client, published_csr_manifest):
    """The book row can change; the document must keep saying what it said."""
    token, project_id, manifest_id = published_csr_manifest
    study = app_client.post("/api/v1/studies", headers=_auth(token), json=STUDY).json()

    made = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study_id": study["id"],
        "rows": DISPOSITION_ROWS, "fields": FIELDS, "version_label": "1.0"}).json()

    app_client.patch(f"/api/v1/studies/{study['id']}", headers=_auth(token),
                     json={"sponsor": "Renamed Pharma Ltd"})

    detail = app_client.get(f"/api/v1/clinical-documents/{made['id']}",
                            headers=_auth(token)).json()
    assert detail["study"]["sponsor"] == "Acme Pharma"


def test_a_sparse_book_study_does_not_erase_supplied_fields(
        app_client, published_csr_manifest):
    """A book study with no recorded sponsor must not delete the sponsor_name
    the caller typed into fields -- absences never win."""
    token, project_id, manifest_id = published_csr_manifest
    study = app_client.post("/api/v1/studies", headers=_auth(token), json={
        "protocol_number": "SPARSE-01", "title": "A sparse study"}).json()

    res = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study_id": study["id"],
        "rows": DISPOSITION_ROWS,
        "fields": {**FIELDS, "sponsor_name": "Typed Sponsor GmbH",
                   "investigator_name": "Dr. Typed"},
        "version_label": "1.0"})
    assert res.status_code == 201, res.text
    record = app_client.get(f"/api/v1/clinical-documents/{res.json()['id']}",
                            headers=_auth(token)).json()["source_record"]
    assert record["sponsor_name"] == "Typed Sponsor GmbH"
    assert record["investigator_name"] == "Dr. Typed"
    assert record["protocol_number"] == "SPARSE-01"  # real values still win


def test_document_list_filters(app_client, published_csr_manifest):
    token, project_id, manifest_id = published_csr_manifest
    study = app_client.post("/api/v1/studies", headers=_auth(token), json=STUDY).json()
    made = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study_id": study["id"],
        "rows": DISPOSITION_ROWS, "fields": FIELDS, "version_label": "1.0"}).json()

    by_study = app_client.get("/api/v1/clinical-documents", headers=_auth(token),
                              params={"study_id": study["id"]}).json()["items"]
    assert any(d["id"] == made["id"] for d in by_study)
    by_type = app_client.get("/api/v1/clinical-documents", headers=_auth(token),
                             params={"document_type": "icf"}).json()["items"]
    assert not any(d["id"] == made["id"] for d in by_type)
    by_project = app_client.get("/api/v1/clinical-documents", headers=_auth(token),
                                params={"project_id": "no-such-project"}).json()["items"]
    assert not by_project


# ------------------------------------------------- review-hardening cases

def test_nan_and_comma_cells_leave_the_total_honestly_blank():
    """Decimal('NaN') constructs without raising, and '1,5' is 1.5 in half
    the world -- both must mean 'no total', never a printed guess."""
    cols = [{"source_key": "n", "type": "number"}]
    assert "total_n" not in derive_row_totals([{"n": "5"}, {"n": "NaN"}], cols)
    assert "total_n" not in derive_row_totals([{"n": "inf"}], cols)
    assert "total_n" not in derive_row_totals([{"n": "1,5"}], cols)
    assert derive_row_totals([{"n": "1.5"}, {"n": 2}], cols)["total_n"] == "3.5"


def test_a_caller_cannot_fabricate_a_derived_total(app_client, published_csr_manifest):
    """When a column cannot sum, its honest blank must not be fillable by a
    caller-typed fields['total_...'] wearing the server's clothes."""
    token, project_id, manifest_id = published_csr_manifest
    rows = [dict(DISPOSITION_ROWS[0]),
            {"site_name": "Delhi", "subjects_enrolled": "n/a",
             "subjects_completed": 1, "subjects_withdrawn": 0}]
    res = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study": STUDY, "rows": rows,
        "fields": {**FIELDS, "total_subjects_enrolled": "999"},
        "version_label": "1.0"})
    assert res.status_code == 201, res.text
    record = app_client.get(f"/api/v1/clinical-documents/{res.json()['id']}",
                            headers=_auth(token)).json()["source_record"]
    assert "total_subjects_enrolled" not in record  # not 999, not anything
    # Columns that DID sum still carry the server's own figure.
    assert record["total_subjects_completed"] == "71"


def test_rows_cannot_be_smuggled_through_fields(app_client, org_a):
    """An optional table with no rows must stay empty: a row list hidden in
    fields[<collection>] would bypass the row filter and the derivations."""
    token, project_id = org_a
    made = app_client.post(
        "/api/v1/template-blueprints", headers=_auth(token),
        json={"name": "ICF smuggle", "kit": "clinical_icf", "project_id": project_id})
    published = app_client.post(
        f"/api/v1/template-blueprints/{made.json()['id']}:publish",
        headers=_auth(token), json={"recompile": False})
    manifest_id = published.json()["manifest_id"]

    res = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "icf", "study": STUDY, "rows": [],
        "fields": {"purpose_description": "p", "procedures_description": "p",
                   "risks_description": "r", "benefits_description": "b",
                   "participant_name": "____",
                   "visits": [{"visit_name": "Ghost", "visit_week": 1,
                               "visit_procedures": "smuggled"}]},
        "version_label": "1.0"})
    assert res.status_code == 201, res.text
    record = app_client.get(f"/api/v1/clinical-documents/{res.json()['id']}",
                            headers=_auth(token)).json()["source_record"]
    assert "visits" not in record


def test_an_inline_study_with_only_a_title_is_refused(app_client, published_csr_manifest):
    """The error text says 'at least a protocol number', and the gate agrees."""
    token, project_id, manifest_id = published_csr_manifest
    res = app_client.post("/api/v1/clinical-documents:generate", headers=_auth(token), json={
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study": {"title": "A study with no protocol"},
        "rows": DISPOSITION_ROWS, "fields": FIELDS})
    assert res.status_code == 422
    assert res.json()["detail"]["error"]["code"] == "CLINICAL_NEEDS_STUDY"


def test_a_template_whose_collection_clashes_or_is_unnamed_is_refused(
        app_client, published_csr_manifest):
    """iterate_over = 'document_number' would have the rows list clobber the
    allocated number; an empty iterate_over would file rows under a name the
    renderer never looks up. Both are template defects, both refuse loudly."""
    token, project_id, manifest_id = published_csr_manifest

    from app.db import SessionLocal
    from app.models import TemplateManifest

    def set_iterate_over(value):
        db = SessionLocal()
        try:
            m = db.get(TemplateManifest, manifest_id)
            blocks = [dict(b) for b in (m.blocks or [])]
            for b in blocks:
                if str(b.get("object_type") or "").upper() == "TABLE_ROW":
                    b["iterate_over"] = value
            m.blocks = blocks
            db.commit()
        finally:
            db.close()

    original = "disposition_rows"
    payload = {
        "manifest_id": manifest_id, "project_id": project_id,
        "document_type": "csr", "study": STUDY,
        "rows": DISPOSITION_ROWS, "fields": FIELDS, "version_label": "1.0"}
    try:
        set_iterate_over("document_number")
        clash = app_client.post("/api/v1/clinical-documents:generate",
                                headers=_auth(token), json=payload)
        assert clash.status_code == 422
        assert clash.json()["detail"]["error"]["code"] == "CLINICAL_BAD_TEMPLATE"

        set_iterate_over("")
        unnamed = app_client.post("/api/v1/clinical-documents:generate",
                                  headers=_auth(token), json=payload)
        assert unnamed.status_code == 422
        assert unnamed.json()["detail"]["error"]["code"] == "CLINICAL_BAD_TEMPLATE"
    finally:
        set_iterate_over(original)
