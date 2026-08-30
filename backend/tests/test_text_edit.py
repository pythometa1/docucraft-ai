"""Editing a finished letter's words without rebuilding the letter.

The guarantee is one sentence and it is the reason this module exists rather
than a rich-text editor: **a text edit changes text and nothing else.** Every
test here is a way of trying to break that.

The alternative has already shipped here as a §20 CRITICAL defect -- an HTML
round-trip that rebuilt the .docx from `<p>` tags and replaced finished letters,
letterhead and all, at the same blob path. `is_html_editable` refuses that path
for template-filled documents; this one exists so the refusal does not also mean
"no editing".
"""

from __future__ import annotations

import docx
import pytest

from app.generation.text_edit import (
    EditRejected,
    apply_edits,
    read_document,
    structural_diff,
)


@pytest.fixture
def letter(tmp_path):
    """A letter with the shapes that break naive editors: a table, a header,
    an empty paragraph and a run split mid-sentence."""
    d = docx.Document()
    d.add_paragraph("PRIVATE & CONFIDENTIAL")
    d.add_paragraph("")
    p = d.add_paragraph("Dear ")
    p.add_run("Amelia").bold = True
    p.add_run(", welcome aboard.")
    d.add_paragraph("Your salary is $82,000.00 per annum.")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Base salary"
    table.cell(0, 1).text = "$82,000.00"
    d.sections[0].header.paragraphs[0].text = "Hospira Australia"
    path = tmp_path / "letter.docx"
    d.save(str(path))
    return str(path)


def _para(paragraphs, needle):
    return next(p for p in paragraphs if needle in p.text)


def test_the_letter_reads_back_as_addressable_runs(letter):
    paragraphs = read_document(letter)
    assert _para(paragraphs, "PRIVATE").text == "PRIVATE & CONFIDENTIAL"
    assert "Amelia" in _para(paragraphs, "Amelia").text

    # Every span carries the coordinate the rest of the system uses, so an edit,
    # a manifest slot and a QA finding all point at the same run.
    for p in paragraphs:
        for s in p.spans:
            assert s.paragraph_index == p.paragraph_index
            assert isinstance(s.span_index, int)


def test_empty_paragraphs_are_kept_so_coordinates_do_not_shift(letter):
    paragraphs = read_document(letter)
    indices = [p.paragraph_index for p in paragraphs]
    assert indices == sorted(indices)
    assert indices == list(range(len(indices))), (
        "renumbering around blank lines would make every coordinate below the "
        "first one point at the wrong run"
    )
    assert any(not p.text.strip() for p in paragraphs)


def test_an_edit_changes_the_words_and_nothing_else(letter, tmp_path):
    paragraphs = read_document(letter)
    target = _para(paragraphs, "salary is")
    out = str(tmp_path / "edited.docx")

    outcome = apply_edits(letter, out, [{
        "paragraph_index": target.paragraph_index,
        "span_index": target.spans[0].span_index,
        "text": "Your salary is $91,020.00 per annum.",
    }])

    assert "$91,020.00" in _para(read_document(out), "salary is").text
    assert outcome["applied"][0]["before"] != outcome["applied"][0]["after"]
    # The guarantee, measured rather than asserted in prose.
    assert structural_diff(letter, out) == []


def test_a_table_cell_survives_an_edit_elsewhere(letter, tmp_path):
    """HTML round-trips lose tables. This must not."""
    paragraphs = read_document(letter)
    target = _para(paragraphs, "PRIVATE")
    out = str(tmp_path / "edited.docx")
    apply_edits(letter, out, [{
        "paragraph_index": target.paragraph_index,
        "span_index": target.spans[0].span_index,
        "text": "PRIVATE AND CONFIDENTIAL",
    }])

    after = docx.Document(out)
    assert len(after.tables) == 1
    assert after.tables[0].cell(0, 1).text == "$82,000.00"


def test_the_header_survives_an_edit(letter, tmp_path):
    """A header lives in its own package part. An edit must not touch it --
    `structural_diff` would catch it, and so does this."""
    paragraphs = read_document(letter)
    target = _para(paragraphs, "PRIVATE")
    out = str(tmp_path / "edited.docx")
    apply_edits(letter, out, [{
        "paragraph_index": target.paragraph_index,
        "span_index": target.spans[0].span_index,
        "text": "STRICTLY CONFIDENTIAL",
    }])

    assert docx.Document(out).sections[0].header.paragraphs[0].text == "Hospira Australia"


def test_several_edits_apply_in_one_save(letter, tmp_path):
    paragraphs = read_document(letter)
    a, b = _para(paragraphs, "PRIVATE"), _para(paragraphs, "salary is")
    out = str(tmp_path / "edited.docx")

    outcome = apply_edits(letter, out, [
        {"paragraph_index": a.paragraph_index, "span_index": a.spans[0].span_index, "text": "CONFIDENTIAL"},
        {"paragraph_index": b.paragraph_index, "span_index": b.spans[0].span_index,
         "text": "Your salary is $95,000.00 per annum."},
    ])
    assert len(outcome["applied"]) == 2
    assert "2 run(s)" in outcome["summary"]
    text = "\n".join(p.text for p in read_document(out))
    assert "CONFIDENTIAL" in text and "$95,000.00" in text


def test_an_edit_naming_a_run_that_does_not_exist_is_refused(letter, tmp_path):
    """A save that reports success and changed nothing is the failure this
    module is careful about."""
    with pytest.raises(EditRejected, match="No editable run"):
        apply_edits(letter, str(tmp_path / "x.docx"),
                    [{"paragraph_index": 999, "span_index": 0, "text": "x"}])


def test_an_edit_that_changes_nothing_is_refused(letter, tmp_path):
    paragraphs = read_document(letter)
    target = _para(paragraphs, "PRIVATE")
    with pytest.raises(EditRejected, match="nothing to save"):
        apply_edits(letter, str(tmp_path / "x.docx"), [{
            "paragraph_index": target.paragraph_index,
            "span_index": target.spans[0].span_index,
            "text": target.spans[0].text,
        }])


def test_no_edits_at_all_is_refused(letter, tmp_path):
    with pytest.raises(EditRejected, match="No edits"):
        apply_edits(letter, str(tmp_path / "x.docx"), [])


@pytest.mark.parametrize("bad", ["a\x00b", "a\x07b", "line\x0bbreak"])
def test_control_characters_are_refused(letter, tmp_path, bad):
    """A control byte in a `<w:t>` produces a file Word declines to open."""
    paragraphs = read_document(letter)
    target = _para(paragraphs, "PRIVATE")
    with pytest.raises(EditRejected, match="control character"):
        apply_edits(letter, str(tmp_path / "x.docx"), [{
            "paragraph_index": target.paragraph_index,
            "span_index": target.spans[0].span_index, "text": bad,
        }])


def test_an_absurdly_long_run_is_refused(letter, tmp_path):
    paragraphs = read_document(letter)
    target = _para(paragraphs, "PRIVATE")
    with pytest.raises(EditRejected, match="limit for one run"):
        apply_edits(letter, str(tmp_path / "x.docx"), [{
            "paragraph_index": target.paragraph_index,
            "span_index": target.spans[0].span_index, "text": "x" * 20001,
        }])


def test_the_edited_file_is_a_document_word_can_open(letter, tmp_path):
    paragraphs = read_document(letter)
    target = _para(paragraphs, "salary is")
    out = str(tmp_path / "edited.docx")
    apply_edits(letter, out, [{
        "paragraph_index": target.paragraph_index,
        "span_index": target.spans[0].span_index,
        "text": "Your salary is $100,000.00 per annum.",
    }])

    import zipfile
    assert zipfile.ZipFile(out).testzip() is None
    reopened = docx.Document(out)
    assert len(reopened.paragraphs) == len(docx.Document(letter).paragraphs)


def test_merge_field_runs_are_not_offered_as_editable(tmp_path):
    """A MERGEFIELD's runs belong to the field. Editing one corrupts the field
    structure, and the fill engine resolves them by replacing the whole run
    sequence -- so they are not text to type into."""
    from tests.conftest import __file__ as _  # noqa: F401  (anchor for fixtures dir)
    from pathlib import Path

    fixture = Path(__file__).resolve().parent / "fixtures" / "templates" / "hospira_offer.docx"
    if not fixture.exists():
        pytest.skip("hospira_offer.docx fixture is missing")

    from app.templates.parsers.docx_prescan import prescan

    scan = prescan(str(fixture))
    assert scan.mergefields, "fixture no longer carries merge fields"
    mf_paragraphs = {m.paragraph_index for m in scan.mergefields}

    paragraphs = read_document(str(fixture))
    for p in paragraphs:
        if p.paragraph_index in mf_paragraphs:
            for span in p.spans:
                assert "MERGEFIELD" not in span.text


@pytest.mark.parametrize("bad", ["Dear Ms Nguyen,\nEmployee ID: A1", "a\tb", "a\r\nb"])
def test_line_breaks_and_tabs_are_refused(letter, tmp_path, bad):
    """Word does not store these as text -- a line break is `<w:br/>` and a tab
    is `<w:tab/>`. One written into a `<w:t>` renders as nothing, so the edit
    appears to save and the words silently disappear.

    Found by asking the model to rewrite a salutation: it returned a three-line
    block and the validator waved it through.
    """
    paragraphs = read_document(letter)
    target = _para(paragraphs, "PRIVATE")
    with pytest.raises(EditRejected, match="tab or line break"):
        apply_edits(letter, str(tmp_path / "x.docx"), [{
            "paragraph_index": target.paragraph_index,
            "span_index": target.spans[0].span_index, "text": bad,
        }])
