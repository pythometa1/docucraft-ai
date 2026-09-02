"""What approval is allowed to wave through.

Approving a manifest used to be a status assignment. Nothing checked that the
thing being approved could run, and nothing retired the manifest it replaced --
so a template could carry three simultaneously-approved manifests and
"which one does production execute?" had no answer.
"""

import pytest
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Project, TemplateFile, TemplateManifest, TemplateVersion, User
from app.manifests.validator import validate_manifest


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


def _ok_manifest():
    """The smallest manifest that passes every rule."""
    return {
        "fields": [{
            "id": "full_name", "type": "text",
            "slots": [{"paragraph_index": 0, "span_index": 0, "text": "<full_name>", "field_id": "full_name"}],
        }],
        "conditions": [{"id": "c1", "expression": "region == 'EU'", "keeps_blocks": ["b1"]}],
        "blocks": [{"id": "b1", "start_paragraph": 1, "end_paragraph": 2}],
    }


def test_a_coherent_manifest_has_no_failures():
    assert validate_manifest(_ok_manifest()) == []


@pytest.mark.parametrize("mutate,rule", [
    # A field with nowhere to go promises a value the letter never shows.
    (lambda m: m["fields"][0].pop("slots"), "field_without_slot"),
    # A typo in the policy must not silently become "leave it empty".
    (lambda m: m["fields"][0].update(on_missing="BLNAK"), "unknown_on_missing_policy"),
    (lambda m: m["fields"][0].update(on_missing="DEFAULT"), "default_policy_without_default"),
    (lambda m: m["fields"].append(dict(m["fields"][0])), "duplicate_id"),
    # A condition governing a block nobody defined silently governs nothing.
    (lambda m: m["conditions"][0].update(keeps_blocks=["b_missing"]), "condition_targets_unknown_block"),
    (lambda m: m["conditions"][0].update(expression="region == = 'EU'"), "condition_does_not_parse"),
    (lambda m: m["conditions"][0].update(expression="1 == 1"), "condition_reads_nothing"),
    (lambda m: m["conditions"][0].update(expression=""), "condition_without_expression"),
    (lambda m: m["blocks"][0].update(start_paragraph=9, end_paragraph=2), "block_range_inverted"),
    (lambda m: m["blocks"][0].pop("end_paragraph"), "block_without_range"),
])
def test_each_incoherence_is_caught(mutate, rule):
    manifest = _ok_manifest()
    mutate(manifest)
    assert rule in {f.rule for f in validate_manifest(manifest)}


def test_an_unresolved_compiler_warning_blocks_approval():
    warnings = [{"code": "W-UNSLOTTED-FIELD", "paragraph_index": -1, "message": "field has no slot"}]
    failures = validate_manifest(_ok_manifest(), warnings=warnings)
    assert [f.rule for f in failures] == ["unresolved_compiler_warning"]


def test_a_disposed_warning_stops_blocking():
    warnings = [{"code": "W-UNSLOTTED-FIELD", "paragraph_index": -1, "message": "field has no slot"}]
    dispositions = {"W-UNSLOTTED-FIELD": {"resolved_by": "u1", "note": "intentional"}}
    assert validate_manifest(_ok_manifest(), warnings=warnings, dispositions=dispositions) == []


# ----------------------------------------------------------------- endpoints

@pytest.fixture()
def template_with_manifests(app_client, two_orgs):
    """A template file plus a factory for manifests on it."""
    _token_a, project_a, *_ = two_orgs
    db = SessionLocal()
    try:
        project = db.get(Project, project_a)
        user = db.scalar(select(User).where(User.org_id == project.org_id))
        template_file = TemplateFile(
            org_id=project.org_id, project_id=project.id, name="v.docx",
            status="parsed", created_by=user.id,
        )
        db.add(template_file)
        db.flush()
        version = TemplateVersion(
            template_file_id=template_file.id, org_id=project.org_id, version_no=1,
            blob_path=f"templates/{project.id}/v.docx", created_by=user.id,
        )
        db.add(version)
        db.flush()
        db.commit()
        ctx = (project.org_id, template_file.id, version.id, user.id)
    finally:
        db.close()

    org_id, template_file_id, version_id, user_id = ctx

    def _make(*, status="draft", manifest=None, warnings=(), version_no=1) -> str:
        payload = manifest if manifest is not None else _ok_manifest()
        db = SessionLocal()
        try:
            row = TemplateManifest(
                org_id=org_id, template_file_id=template_file_id, template_version_id=version_id,
                version_no=version_no, status=status, fields=payload["fields"],
                conditions=payload["conditions"], blocks=payload["blocks"], delete_always=[],
                warnings=list(warnings), created_by=user_id,
            )
            db.add(row)
            db.commit()
            return row.id
        finally:
            db.close()

    return _make


def _status(manifest_id: str) -> str:
    db = SessionLocal()
    try:
        return db.get(TemplateManifest, manifest_id).status
    finally:
        db.close()


def test_approve_refuses_an_incoherent_manifest(app_client, two_orgs, template_with_manifests):
    token_a, *_ = two_orgs
    broken = _ok_manifest()
    broken["conditions"][0]["keeps_blocks"] = ["b_missing"]
    manifest_id = template_with_manifests(manifest=broken)

    res = app_client.post(f"/api/v1/template-manifests/{manifest_id}:approve", headers=_auth(token_a))

    assert res.status_code == 409, res.text
    body = res.json()["detail"]["error"]
    assert body["code"] == "MANIFEST_INVALID"
    assert body["details"]["failures"][0]["rule"] == "condition_targets_unknown_block"
    assert _status(manifest_id) == "draft", "a refused approval must not change the status"


def test_approve_accepts_a_coherent_manifest(app_client, two_orgs, template_with_manifests):
    token_a, *_ = two_orgs
    manifest_id = template_with_manifests()

    res = app_client.post(f"/api/v1/template-manifests/{manifest_id}:approve", headers=_auth(token_a))

    assert res.status_code == 200, res.text
    assert _status(manifest_id) == "approved"


def test_approving_a_new_version_supersedes_the_old_one(app_client, two_orgs, template_with_manifests):
    """Exactly one approved manifest per template, or production has a choice
    to make that nobody made."""
    token_a, *_ = two_orgs
    first = template_with_manifests(version_no=1)
    second = template_with_manifests(version_no=2)

    assert app_client.post(f"/api/v1/template-manifests/{first}:approve", headers=_auth(token_a)).status_code == 200
    res = app_client.post(f"/api/v1/template-manifests/{second}:approve", headers=_auth(token_a))

    assert res.status_code == 200, res.text
    assert first in res.json()["superseded"]
    assert _status(first) == "superseded"
    assert _status(second) == "approved"


def test_a_superseded_manifest_can_no_longer_generate(app_client, two_orgs, template_with_manifests):
    token_a, project_a, *_ = two_orgs
    first = template_with_manifests(version_no=1)
    second = template_with_manifests(version_no=2)
    app_client.post(f"/api/v1/template-manifests/{first}:approve", headers=_auth(token_a))
    app_client.post(f"/api/v1/template-manifests/{second}:approve", headers=_auth(token_a))

    res = app_client.post(
        f"/api/v1/template-manifests/{first}/generate",
        json={"source_record": {"full_name": "Dana"}, "project_id": project_a},
        headers=_auth(token_a),
    )
    assert res.status_code == 409
    # `MANIFEST_RETIRED`, not `MANIFEST_NOT_APPROVED`. Generating no longer asks
    # whether somebody signed the reading -- it asks whether the reading is the
    # current one, which is the question this test was always really about: v1
    # was replaced by v2, so filling from it produces letters built from a
    # version of the template the project has moved off.
    assert res.json()["detail"]["error"]["code"] == "MANIFEST_RETIRED"


def test_the_validation_endpoint_reports_without_approving(app_client, two_orgs, template_with_manifests):
    token_a, *_ = two_orgs
    broken = _ok_manifest()
    broken["fields"][0].pop("slots")
    manifest_id = template_with_manifests(manifest=broken)

    res = app_client.get(f"/api/v1/template-manifests/{manifest_id}/validation", headers=_auth(token_a))

    assert res.status_code == 200, res.text
    body = res.json()
    assert body["can_approve"] is False
    assert body["failures"][0]["rule"] == "field_without_slot"
    assert _status(manifest_id) == "draft"


def test_a_warning_must_be_answered_before_approval(app_client, two_orgs, template_with_manifests):
    token_a, *_ = two_orgs
    warnings = [{"code": "W-DUP-STATIC", "paragraph_index": 4, "message": "static text duplicates block b1"}]
    manifest_id = template_with_manifests(warnings=warnings)

    blocked = app_client.post(f"/api/v1/template-manifests/{manifest_id}:approve", headers=_auth(token_a))
    assert blocked.status_code == 409
    assert blocked.json()["detail"]["error"]["details"]["failures"][0]["rule"] == "unresolved_compiler_warning"

    resolved = app_client.post(
        f"/api/v1/template-manifests/{manifest_id}/warnings:resolve",
        json={"code": "W-DUP-STATIC", "note": "duplication is intended"},
        headers=_auth(token_a),
    )
    assert resolved.status_code == 200, resolved.text
    assert resolved.json()["dispositions"]["W-DUP-STATIC"]["note"] == "duplication is intended"

    assert app_client.post(
        f"/api/v1/template-manifests/{manifest_id}:approve", headers=_auth(token_a)
    ).status_code == 200


def test_an_unknown_warning_code_cannot_be_resolved(app_client, two_orgs, template_with_manifests):
    """Otherwise "resolve everything" is a single call away from meaningless."""
    token_a, *_ = two_orgs
    manifest_id = template_with_manifests(warnings=[{"code": "W-DUP-STATIC", "message": "x"}])

    res = app_client.post(
        f"/api/v1/template-manifests/{manifest_id}/warnings:resolve",
        json={"code": "W-INVENTED", "note": ""},
        headers=_auth(token_a),
    )
    assert res.status_code == 404
    assert res.json()["detail"]["error"]["code"] == "WARNING_NOT_FOUND"


def test_a_blocked_mapping_stops_the_manifest_being_approved(app_client, two_orgs, template_with_manifests):
    """§13's fourth band: "Block | < 0.50, or any veto | The manifest cannot be
    locked until a human resolves it."

    confidence.py computed the band and stamped it on every suggestion from the
    day it landed, and nothing read it back -- so the band that exists purely to
    stop a lock stopped nothing.
    """
    from app.compiler.confidence import Band
    from app.models import SuggestionLog

    token_a, *_ = two_orgs
    manifest_id = template_with_manifests()

    db = SessionLocal()
    try:
        manifest = db.get(TemplateManifest, manifest_id)
        db.add(SuggestionLog(
            org_id=manifest.org_id, manifest_id=manifest_id, source_version_id="sv-1",
            object_id="annual_salary", suggested_column="Amt", method="fuzzy",
            score=0.31, band=Band.BLOCK.value, vetoes=["money_field_needs_two_signals"],
            evidence=[], reviewer_decision="pending",
        ))
        db.commit()
    finally:
        db.close()

    refused = app_client.post(f"/api/v1/template-manifests/{manifest_id}:approve", headers=_auth(token_a))
    assert refused.status_code == 409, refused.text
    body = refused.json()["detail"]["error"]
    assert body["code"] == "MANIFEST_HAS_BLOCKED_MAPPINGS"
    assert "annual_salary" in body["details"]["blocked_objects"]
    assert _status(manifest_id) == "draft", "a refused approval must not lock the manifest"

    # A reviewer deciding the mapping is what clears it -- not re-running the
    # compiler, which would just restate the same score.
    db = SessionLocal()
    try:
        row = db.scalars(select(SuggestionLog).where(SuggestionLog.manifest_id == manifest_id)).one()
        row.reviewer_decision = "edited"
        row.final_column = "Annual Salary"
        db.commit()
    finally:
        db.close()

    assert app_client.post(
        f"/api/v1/template-manifests/{manifest_id}:approve", headers=_auth(token_a)
    ).status_code == 200


def test_a_confirm_band_suggestion_does_not_block_approval(app_client, two_orgs, template_with_manifests):
    """Only the BLOCK band stops a lock. Treating CONFIRM the same way would
    make the four bands one band."""
    from app.compiler.confidence import Band
    from app.models import SuggestionLog

    token_a, *_ = two_orgs
    manifest_id = template_with_manifests()

    db = SessionLocal()
    try:
        manifest = db.get(TemplateManifest, manifest_id)
        db.add(SuggestionLog(
            org_id=manifest.org_id, manifest_id=manifest_id, source_version_id="sv-1",
            object_id="first_name", suggested_column="First Name", method="exact_slug",
            score=0.88, band=Band.CONFIRM.value, vetoes=[], evidence=[],
            reviewer_decision="pending",
        ))
        db.commit()
    finally:
        db.close()

    assert app_client.post(
        f"/api/v1/template-manifests/{manifest_id}:approve", headers=_auth(token_a)
    ).status_code == 200


def test_a_failed_compile_cannot_be_approved(app_client, two_orgs, template_with_manifests):
    """A compile that could not read the template is the absence of a manifest,
    recorded -- not a manifest with problems. Approving it would let documents be
    generated from a reading that does not exist.

    This slipped through once in the running app: `validate_manifest` had the
    rule, but both call sites built its input as `{"fields", "conditions",
    "blocks"}` and dropped `status`, so the check had nothing to read.
    """
    token_a, *_ = two_orgs
    manifest_id = template_with_manifests(status="failed")

    validation = app_client.get(
        f"/api/v1/template-manifests/{manifest_id}/validation", headers=_auth(token_a)
    )
    assert validation.status_code == 200, validation.text
    body = validation.json()
    assert body["can_approve"] is False
    assert "compile_failed" in [f["rule"] for f in body["failures"]]

    res = app_client.post(f"/api/v1/template-manifests/{manifest_id}:approve", headers=_auth(token_a))
    assert res.status_code == 409, res.text
    assert res.json()["detail"]["error"]["code"] == "MANIFEST_INVALID"
    assert _status(manifest_id) == "failed", "a refused approval must not change the status"


def test_the_validator_is_handed_delete_always_so_orphans_can_be_seen():
    """`orphaned_fields` reads `delete_always`. The approval call sites omitted
    the key, so the one check whose whole job is "this field sits only in
    paragraphs the compile deletes" saw nothing deleted and could never fire --
    at the exact moment it exists for."""
    from app.manifests.validator import validate_manifest
    from app.routers.manifests import _validatable

    class _Row:
        fields = [{"id": "x", "slots": [{"paragraph_index": 3, "span_index": 0, "text": "<X>"}]}]
        conditions: list = []
        blocks: list = []
        delete_always = [{"paragraph_index": 3, "span_index": None}]
        status = "draft"

    shape = _validatable(_Row())
    assert shape["delete_always"] == _Row.delete_always
    assert shape["status"] == "draft"
    assert "orphaned_field" in [f.rule for f in validate_manifest(shape)]


# ------------------------------- telling the reviewer before the batch, not after
#
# The China letter reached an approved manifest, a bound source and a running
# batch before anyone learned that `<Pay Rate Monthly>` was never claimed. The
# information existed at compile time; nothing put it in front of a person until
# every row came back "Leftover placeholder brackets".

def _template_with_a_split_placeholder(tmp_path):
    """A .docx whose placeholder Word has split across runs, as real ones are."""
    import docx

    d = docx.Document()
    d.add_paragraph("Dear <Colleague Name>,")
    p = d.add_paragraph("Your new base salary will be CNY")
    p.add_run("<")
    p.add_run("Pay Rate Monthly")
    p.add_run(">")
    p.add_run(" per month.")
    path = tmp_path / "china.docx"
    d.save(str(path))
    return str(path)


def test_an_uncovered_placeholder_is_reported_by_validation(app_client, tmp_path):
    """It is what the reviewer needed to see while deciding, and the message says
    what happens if they do not act on it."""
    from app.compiler.assertions import uncovered_placeholders
    from app.templates.parsers.docx_prescan import prescan

    scan = prescan(_template_with_a_split_placeholder(tmp_path))
    manifest = {
        "fields": [{"id": "colleague_name",
                    "slots": [{"text": "<Colleague Name>", "paragraph_index": 0}]}],
        "conditions": [], "blocks": [], "delete_always": [],
    }
    found, _split = uncovered_placeholders(scan, manifest)
    assert [a.detail.split("'")[1] for a in found] == ["<Pay Rate Monthly>"]


def test_an_undispositioned_warning_stops_approval_and_a_disposition_releases_it():
    """The mechanism the coverage warning rides on. A warning that blocks
    approval must be answerable, or it is a gate with no handle."""
    from app.manifests.validator import validate_manifest

    manifest = {"fields": [], "conditions": [], "blocks": [], "delete_always": [],
                "status": "draft"}
    warning = {"code": "uncovered_placeholder", "paragraph_index": 19,
               "message": "Paragraph 19 contains the placeholder '<Pay Rate Monthly>'."}

    blocked = validate_manifest(manifest, warnings=[warning], dispositions={})
    assert [f.rule for f in blocked] == ["unresolved_compiler_warning"]
    assert "<Pay Rate Monthly>" in blocked[0].detail

    released = validate_manifest(
        manifest, warnings=[warning],
        dispositions={"uncovered_placeholder": {"note": "deleted by the pay-change block"}})
    assert released == []


def test_the_validation_screen_and_approval_see_the_same_warnings(app_client, two_orgs):
    """A warning the screen shows and approval ignores is noise; one approval
    enforces and the screen never showed is a 409 out of nowhere. Both call
    `_all_warnings`, and this pins that they still do.

    The checks approval makes now live in `_approval_blockers`, because a compile
    approves its own output where it may and the two paths must not be able to
    disagree about what "approvable" means. So the assertion moved with them --
    and the third one below is what stops `approve_manifest` growing its own copy
    of a check again."""
    import inspect

    from app.routers import manifests as mod

    source = inspect.getsource(mod)
    validation = source[source.index("def manifest_validation"):]
    validation = validation[:validation.index("\n@router")]
    blockers = source[source.index("def _approval_blockers"):]
    blockers = blockers[:blockers.index("\ndef _record_approval")]
    approve = source[source.index("def approve_manifest"):]
    approve = approve[:approve.index("\n@router")]

    assert "_all_warnings(db, m)" in validation, "the validation screen must merge coverage warnings"
    assert "_all_warnings(db, m)" in blockers, "approval must apply the same set the screen showed"
    assert "_approval_blockers(db, user, m)" in approve, (
        "approval must go through the shared checks rather than re-implementing them")


def test_the_compile_path_and_approve_run_the_same_checks(app_client, two_orgs):
    """`:approve` and the automatic approval on compile are one set of rules.

    Written twice they drift, and the half that drifts is the automatic one --
    nobody clicks it, so nobody notices it stopped enforcing something. Both go
    through `_approval_blockers`, and neither re-implements a check."""
    import inspect

    from app.routers import manifests as mod

    source = inspect.getsource(mod)
    auto = source[source.index("def try_auto_approve"):]
    auto = auto[:auto.index("\n@router")]

    assert "_approval_blockers(db, user, m)" in auto
    assert "_record_approval(db, user, m)" in auto
    # The two the automatic path must refuse outright rather than evaluate.
    assert "has_capability(user, APPROVE_MANIFEST)" in auto, (
        "compiling must not confer an approval right the caller's role withholds")
    assert "legally_binding" in auto, (
        "a four-eyes template can never be approved by whoever compiled it")


def test_a_template_whose_file_is_gone_does_not_break_the_validation_screen(app_client, two_orgs):
    """Best effort. An unreadable template is its own problem, and refusing to
    render the screen would hide every other blocker on it."""
    from app.db import SessionLocal
    from app.models import Project, TemplateFile, TemplateManifest, TemplateVersion, User
    from app.routers.manifests import _coverage_warnings

    _token_a, project_a, *_ = two_orgs
    db = SessionLocal()
    try:
        project = db.get(Project, project_a)
        user = db.query(User).filter(User.org_id == project.org_id).first()
        tf = TemplateFile(org_id=project.org_id, project_id=project.id, name="gone.docx",
                          status="ready", created_by=user.id)
        db.add(tf)
        db.flush()
        tv = TemplateVersion(template_file_id=tf.id, org_id=project.org_id, version_no=1,
                             blob_path="templates/does-not-exist.docx", created_by=user.id)
        db.add(tv)
        db.flush()
        m = TemplateManifest(
            org_id=project.org_id, template_file_id=tf.id, template_version_id=tv.id,
            version_no=1, status="draft", fields=[], conditions=[], blocks=[],
            delete_always=[], confidence=1.0, compiled_by="rule_based",
            prescan_summary={}, created_by=user.id)
        db.add(m)
        db.flush()

        assert _coverage_warnings(db, m) == []
    finally:
        db.rollback()
        db.close()


# --------------------------------------- the placeholders a template will not fill

def test_the_template_row_names_the_placeholders_nothing_will_fill(app_client, two_orgs):
    """A QA block, four steps earlier and on the thing that has to change.

    An unclaimed placeholder is not a risk, it is a decided outcome: the fill
    engine writes values into the run it found the slot in, there is no slot for
    this one, so the literal `<Pay Rate Monthly>` survives into the letter and QA
    refuses it. Every row of the batch fails the same way.

    Before this it was only discoverable by running a batch -- three canary rows
    block, the run stops, and the reader is on the Documents stage reading an
    error about a template they uploaded four steps and one spreadsheet ago. The
    check runs at compile time, so the answer exists the moment the template is
    read; it was simply computed on demand and stored nowhere, which meant the
    only screens that could see it were the ones that asked the validation
    endpoint.
    """
    from app.db import SessionLocal
    from app.models import TemplateFile, TemplateManifest, TemplateVersion
    from app.routers.templates import _manifest_summary

    _token_a, project_a, *_ = two_orgs
    db = SessionLocal()
    try:
        project = db.get(Project, project_a)
        user = db.scalar(select(User).where(User.org_id == project.org_id))
        tf = TemplateFile(org_id=project.org_id, project_id=project.id, name="unfillable.docx",
                          status="parsed", created_by=user.id)
        db.add(tf)
        db.flush()
        tv = TemplateVersion(template_file_id=tf.id, org_id=project.org_id, version_no=1,
                             blob_path=f"templates/{project.id}/unfillable.docx", created_by=user.id)
        db.add(tv)
        db.flush()
        payload = _ok_manifest()
        db.add(TemplateManifest(
            org_id=project.org_id, template_file_id=tf.id, template_version_id=tv.id,
            version_no=1, status="approved", fields=payload["fields"],
            conditions=payload["conditions"], blocks=payload["blocks"], delete_always=[],
            # Stored on the manifest, which is the point: reading a project's
            # templates must not re-parse every .docx to answer this.
            warnings=[
                {"code": "W-SPLIT-PLACEHOLDER", "paragraph_index": 19,
                 "evidence": "<Pay Rate\nMonthly>",
                 "message": "split across runs, so nothing can fill it"},
                {"code": "uncovered_placeholder", "paragraph_index": 15,
                 "evidence": "<Transaction Action Reason>",
                 "message": "no field claims it"},
                # Not a coverage warning, and must not be counted as one.
                {"code": "W-SOMETHING-ELSE", "paragraph_index": 2, "message": "unrelated"},
            ],
            created_by=user.id))
        db.commit()
        summary = _manifest_summary(db, tf.id)
    finally:
        db.close()

    assert summary["unfillable_count"] == 2, "only the coverage warnings count"
    codes = {w["code"] for w in summary["unfillable"]}
    assert codes == {"W-SPLIT-PLACEHOLDER", "uncovered_placeholder"}

    by_code = {w["code"]: w for w in summary["unfillable"]}
    # Newlines flattened: Word splits these across lines, and the raw token
    # renders as a broken three-line chip on the row.
    assert by_code["W-SPLIT-PLACEHOLDER"]["placeholder"] == "<Pay Rate Monthly>"
    assert by_code["W-SPLIT-PLACEHOLDER"]["paragraph_index"] == 19
    # The two codes are kept apart because they need different advice: one is a
    # slot a field could still claim, the other Word has broken across runs and
    # only a change to the document fixes it.
    assert by_code["uncovered_placeholder"]["placeholder"] == "<Transaction Action Reason>"


def test_a_template_with_nothing_outstanding_says_so(app_client, two_orgs, template_with_manifests):
    """The empty case is an explicit zero, not a missing key -- the row reads it
    on every template and an absent field would render as `undefined`."""
    from app.db import SessionLocal
    from app.models import TemplateManifest
    from app.routers.templates import _manifest_summary

    manifest_id = template_with_manifests()
    db = SessionLocal()
    try:
        summary = _manifest_summary(db, db.get(TemplateManifest, manifest_id).template_file_id)
    finally:
        db.close()

    assert summary["unfillable_count"] == 0
    assert summary["unfillable"] == []
