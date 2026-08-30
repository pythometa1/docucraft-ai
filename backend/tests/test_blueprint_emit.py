"""What the emitter has to get right for the pre-scanner to read it back.

The emitter is the inverse of `docx_prescan`, and the pre-scanner is small but
not obvious. Each test here pins one of its mechanics, and each of those
mechanics is a way for an emitted template to be quietly wrong: it compiles, it
approves, and the letters come out with a placeholder printed in them or a value
written into the wrong sentence.
"""

import hashlib
import os

import docx
import pytest

from app.templates import blueprint as bp
from app.templates.emit_docx import EmitError, RoundTripError, emit, emit_from_base
from app.templates.parsers.docx_prescan import prescan
from app.templates.parsers.docx_safety import inspect_package


def _body(*blocks):
    return bp.normalise_body({"blocks": list(blocks), "sect_pr_from": None})


def _p(*segments, style=None):
    return bp.paragraph(segments, style=style)


def _emit(tmp_path, body, name="t.docx"):
    path = str(tmp_path / name)
    return emit(body, path), path


# ---- colours: the whole contract with the compiler ----

def test_a_placeholder_reads_back_blue_and_an_instruction_red(tmp_path):
    """This is the contract. Blue is a slot to fill, red is a note to the author
    that the fill engine deletes; get either wrong and the letter ships with the
    wrong half of the template in it."""
    _r, path = _emit(tmp_path, _body(_p(
        bp.segment("static", "Dear "),
        bp.segment("placeholder", "<Name>"),
        bp.segment("instruction", "delete this line"))))
    spans = prescan(path).spans
    assert [(s.color, s.text) for s in spans] == [
        ("black", "Dear "), ("blue", "<Name>"), ("red", "delete this line")]


def test_a_hyperlink_is_its_own_span_and_is_not_read_as_a_placeholder(tmp_path):
    _r, path = _emit(tmp_path, _body(_p(
        bp.segment("static", "See the "),
        bp.segment("hyperlink", "policy", target="https://example.com/p"),
        bp.segment("static", " for detail."))))
    scan = prescan(path)
    link = [s for s in scan.spans if s.in_hyperlink]
    assert len(link) == 1
    assert link[0].text == "policy"
    assert link[0].color == "black", "a blue link would be compiled as a field to fill"
    assert [(h.text, h.target) for h in scan.hyperlinks] == [
        ("policy", "https://example.com/p")]


# ---- the span-merge trap ----

def test_span_indices_are_contiguous_and_match_what_the_body_declared(tmp_path):
    body = _body(_p(bp.segment("static", "a "), bp.segment("placeholder", "<b>"),
                    bp.segment("static", " c "), bp.segment("instruction", "d")))
    result, path = _emit(tmp_path, body)
    spans = prescan(path).spans
    assert [s.span_index for s in spans] == [0, 1, 2, 3]
    assert result.span_count == 4


def test_a_body_that_was_never_normalised_is_refused(tmp_path):
    """Emitting it would produce a document with fewer spans than the body
    claims, and every slot after the merge would address the span to its left.
    Refusing is the only answer that cannot be silently wrong."""
    unmerged = {"blocks": [_p(bp.segment("static", "a"), bp.segment("static", "b"))],
                "sect_pr_from": None}
    with pytest.raises(EmitError, match="normalised"):
        emit(unmerged, str(tmp_path / "x.docx"))


def test_leading_and_trailing_spaces_survive(tmp_path):
    """Without `xml:space="preserve"` Word collapses "Dear " to "Dear", the
    paragraph reads back a character short, and every anchor context hash taken
    over it stops matching."""
    _r, path = _emit(tmp_path, _body(_p(
        bp.segment("static", "Dear "), bp.segment("placeholder", "<Name>"))))
    assert prescan(path).spans[0].text == "Dear "


# ---- merge fields ----

def test_a_merge_field_reads_back_with_its_code_and_is_not_a_span(tmp_path):
    """Emitted as anything less than the full begin/instrText/separate/end
    sequence the code is unfindable, and the compiler reports a template with no
    salary field in it."""
    body = _body(_p(bp.segment("static", "$"),
                    bp.segment("mergefield", code="LAB__FT_SALARY__38_HR_")))
    _r, path = _emit(tmp_path, body)
    scan = prescan(path)
    assert [(m.paragraph_index, m.code) for m in scan.mergefields] == [
        (0, "LAB__FT_SALARY__38_HR_")]
    assert [s.text for s in scan.spans] == ["$"], "the field's result run must not become a span"


def test_a_merge_field_breaks_the_span_either_side_of_it(tmp_path):
    body = _body(_p(bp.segment("static", "$"), bp.segment("mergefield", code="SALARY"),
                    bp.segment("static", " gross")))
    _r, path = _emit(tmp_path, body)
    spans = prescan(path).spans
    assert [(s.span_index, s.text) for s in spans] == [(0, "$"), (1, " gross")]


def test_a_merge_field_with_no_code_is_refused(tmp_path):
    body = _body(_p(bp.segment("static", "$"), bp.segment("mergefield", code="  ")))
    with pytest.raises(EmitError, match="no code"):
        emit(body, str(tmp_path / "x.docx"))


def test_a_hyperlink_with_no_target_is_refused(tmp_path):
    """Word stores the destination in a relationship, not in the text, so a link
    with no target is a run that merely looks like one."""
    body = _body(_p(bp.segment("hyperlink", "policy")))
    with pytest.raises(EmitError, match="no target"):
        emit(body, str(tmp_path / "x.docx"))


# ---- tables ----

def test_table_paragraphs_take_their_place_in_the_document_order(tmp_path):
    body = _body(
        _p(bp.segment("static", "before")),
        bp.table([[[_p(bp.segment("static", "cell"))]]]),
        _p(bp.segment("static", "after")))
    _r, path = _emit(tmp_path, body)
    scan = prescan(path)
    assert [s.text for s in scan.spans] == ["before", "cell", "after"]
    assert scan.table_paragraph_indices == {1}


def test_a_new_cell_does_not_leave_an_empty_paragraph_behind(tmp_path):
    """python-docx gives every new cell one empty paragraph. Left in place it is
    a paragraph the walk counts and the body does not, so every index after the
    table is one too high."""
    body = _body(bp.table([[[_p(bp.segment("static", "only"))]]]),
                 _p(bp.segment("static", "after")))
    result, path = _emit(tmp_path, body)
    assert result.paragraph_count == 2
    assert len(prescan(path).paragraphs) == 2


def test_a_merge_field_inside_a_table_cell_is_flagged_as_in_table(tmp_path):
    """The fill engine removes a whole table when a doomed paragraph sits in
    one, so whether a field is in a table is load-bearing downstream."""
    body = _body(bp.table([[[_p(bp.segment("static", "Salary"))],
                            [_p(bp.segment("static", "$"),
                                bp.segment("mergefield", code="SALARY"))]]]))
    _r, path = _emit(tmp_path, body)
    assert prescan(path).mergefields[0].in_table is True


def test_a_ragged_table_is_refused(tmp_path):
    body = _body(bp.table([[[_p(bp.segment("static", "a"))], [_p(bp.segment("static", "b"))]],
                           [[_p(bp.segment("static", "c"))]]]))
    with pytest.raises(EmitError, match="same number of cells"):
        emit(body, str(tmp_path / "x.docx"))


def test_a_table_with_no_rows_is_refused(tmp_path):
    with pytest.raises(EmitError, match="no rows"):
        emit(_body(bp.table([])), str(tmp_path / "x.docx"))


# ---- the file itself ----

def test_emitting_the_same_body_twice_produces_the_same_bytes(tmp_path):
    """Otherwise `manifests.versioning.template_hash` records when a file was
    written rather than what is in it, and the (template_hash, manifest_hash)
    pin §6 rests on is a value nobody can reproduce."""
    body = _body(_p(bp.segment("static", "Dear "), bp.segment("placeholder", "<Name>")))
    _r, first = _emit(tmp_path, body, "first.docx")
    _r, second = _emit(tmp_path, body, "second.docx")
    digest = lambda p: hashlib.sha256(open(p, "rb").read()).hexdigest()
    assert digest(first) == digest(second)


def test_the_emitted_package_passes_the_upload_safety_check(tmp_path):
    """It goes back through `POST /projects/{id}/templates` like any other
    template, and that route runs `inspect_package` before python-docx opens
    the file."""
    _r, path = _emit(tmp_path, _body(_p(bp.segment("static", "hello"))))
    inspect_package(path)  # raises if unsafe or malformed


def test_a_paragraph_style_is_written(tmp_path):
    _r, path = _emit(tmp_path, _body(_p(bp.segment("static", "Offer"), style="Heading 1")))
    assert docx.Document(path).paragraphs[0].style.name == "Heading 1"


def test_emit_reports_what_it_wrote(tmp_path):
    body = _body(_p(bp.segment("static", "a"), bp.segment("placeholder", "<b>")),
                 _p(bp.segment("static", "$"), bp.segment("mergefield", code="S")))
    result, path = _emit(tmp_path, body)
    assert (result.paragraph_count, result.span_count, result.mergefield_count) == (2, 3, 1)
    assert os.path.exists(path)


def test_the_round_trip_is_checked_on_every_emit_not_only_in_a_test(monkeypatch, tmp_path):
    """A span-numbering slip discovered during a batch is discovered in a
    letter, after approval. `emit` re-scans its own output so the failure
    happens at the only moment it can still be cheap."""
    from app.templates import emit_docx

    monkeypatch.setattr(emit_docx, "span_plan", lambda segments: list(span_plan_wrong(segments)))

    def span_plan_wrong(segments):
        # Claim one more span than the paragraph can possibly hold.
        yield from bp.span_plan(segments)
        yield (0, 99)

    with pytest.raises(RoundTripError):
        emit(_body(_p(bp.segment("static", "a"))), str(tmp_path / "x.docx"))


# ---- the legacy path: the customer's own file, edited ----

REAL_TEMPLATE = "tests/fixtures/templates/compensation_letter.docx"


def test_publishing_a_legacy_template_preserves_every_part_but_the_text(tmp_path):
    """Not a python-docx round trip, and the difference is measurable. Opening
    this very file with python-docx and saving it drops four relationship parts
    -- comments, endnotes, fontTable and footnotes -- plus the package's
    directory entries. Word opens the result anyway, which is what makes it
    dangerous: nothing is visibly wrong until something needs a relationship
    that is no longer declared."""
    import zipfile

    from app.templates.read_docx import read_body

    body, _notes = read_body(REAL_TEMPLATE)
    out = str(tmp_path / "published.docx")
    emit_from_base(body, REAL_TEMPLATE, out)

    before = zipfile.ZipFile(REAL_TEMPLATE)
    after = zipfile.ZipFile(out)
    assert before.namelist() == after.namelist()
    for name in before.namelist():
        if name.endswith("/") or name == "word/document.xml":
            continue
        assert before.read(name) == after.read(name), f"{name} was rewritten"


def test_a_python_docx_round_trip_would_have_lost_those_parts(tmp_path):
    """The measurement behind the test above, kept so the reason for doing zip
    surgery is not folklore."""
    import shutil
    import zipfile

    out = str(tmp_path / "roundtrip.docx")
    shutil.copyfile(REAL_TEMPLATE, out)
    docx.Document(out).save(out)

    lost = set(zipfile.ZipFile(REAL_TEMPLATE).namelist()) - set(zipfile.ZipFile(out).namelist())
    assert "word/_rels/footnotes.xml.rels" in lost


def test_an_instruction_marked_not_to_emit_is_gone_from_the_published_file(tmp_path):
    """"The AI re-maps it properly and hands it back" -- the author's notes to
    whoever assembles the letter are not in the template a colleague reuses."""
    from app.templates.read_docx import read_body

    source = str(tmp_path / "source.docx")
    emit(_body(_p(bp.segment("static", "Signed "),
                  bp.segment("instruction", "(remove this line)"),
                  bp.segment("static", " by hand"))), source)

    body, _notes = read_body(source)
    body["blocks"][0]["segments"][1]["emit"] = False
    out = str(tmp_path / "clean.docx")
    emit_from_base(body, source, out)

    scan = prescan(out)
    assert [s.text for s in scan.spans] == ["Signed  by hand"], (
        "dropping the instruction leaves its neighbours adjacent, so Word stores them as one run")
    assert not [s for s in scan.spans if s.color == "red"]


def test_dropping_a_segment_can_cost_two_spans_and_the_prediction_knows_it(tmp_path):
    """The subtle half of the merge rule. Removing one instruction between two
    static runs removes *two* spans, not one -- and the empty run left behind
    does not separate them, because `prescan` skips a run with no text without
    breaking the span it sits in."""
    segments = [bp.segment("static", "a"), bp.segment("instruction", "x"),
                bp.segment("static", "b")]
    assert len(bp.emitted_spans(segments)) == 3

    segments[1]["emit"] = False
    emitted = bp.emitted_spans(segments)
    assert [s["text"] for s in emitted] == ["ab"]


def test_a_body_whose_shape_diverged_from_its_base_is_refused(tmp_path):
    """Writing only the text changes of a body that has gained a paragraph would
    produce a file that disagrees with it about how many paragraphs it has, and
    every anchor below the insertion would address the wrong one."""
    from app.templates.read_docx import read_body

    source = str(tmp_path / "source.docx")
    emit(_body(_p(bp.segment("static", "one"))), source)
    body, _notes = read_body(source)
    body["blocks"].append(_p(bp.segment("static", "two")))

    with pytest.raises(EmitError, match="paragraph"):
        emit_from_base(body, source, str(tmp_path / "out.docx"))


def test_editing_the_text_of_a_legacy_template_changes_only_that(tmp_path):
    from app.templates.read_docx import read_body

    source = str(tmp_path / "source.docx")
    emit(_body(_p(bp.segment("static", "Dear "), bp.segment("placeholder", "<Name>"))), source)

    body, _notes = read_body(source)
    body["blocks"][0]["segments"][1]["text"] = "<Colleague First Name>"
    out = str(tmp_path / "out.docx")
    emit_from_base(body, source, out)

    scan = prescan(out)
    assert [(s.color, s.text) for s in scan.spans] == [
        ("black", "Dear "), ("blue", "<Colleague First Name>")]
