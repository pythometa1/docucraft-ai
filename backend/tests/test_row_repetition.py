"""The §6 TABLE_ROW primitive: one template row, rendered once per record.

The design under test is template-row-as-prototype, processed dead last: the
repeat pass runs after every coordinate-driven pass in `fill_template`, so a
cloned row never enters the (paragraph_index, span_index) space anything else
addresses. The coordinate-integrity test is the one that pins that property --
fields *after* the table must land correctly however many rows were rendered.
"""

import hashlib

import docx
import pytest

from app.generation.docx_renderer import fill_template
from app.manifests.validator import validate_manifest
from app.templates.blueprint_lint import legacy_manifest
from app.templates.blueprint_ops import apply_operations
from app.templates.kits import kit_objects, load_kit


@pytest.fixture(scope="module")
def invoice_template(tmp_path_factory):
    """The invoice kit, emitted to a real .docx, plus its manifest dict."""
    from app.templates.emit_docx import emit

    path = tmp_path_factory.mktemp("row_repeat") / "invoice.docx"
    kit = load_kit("invoice")
    objects = kit_objects(kit)
    emit(kit["body"], str(path))
    manifest = legacy_manifest(objects, delete_always=[])
    return str(path), manifest


RECORD = {
    "business_name": "Sundar Decorations", "business_address": "14 MG Road, Pune",
    "business_email": "hello@sundardecor.in", "invoice_number": "INV-0042",
    "invoice_date": "2026-09-05", "due_date": "2026-09-19",
    "customer_name": "Hotel Blue Orchid", "customer_address": "Baner Road, Pune",
    "currency": "INR", "subtotal": 34500, "tax_rate": 18, "tax_amount": 6210,
    "grand_total": 40710, "payment_terms": "14 days, bank transfer.",
}

ITEMS = [
    {"item_description": "Stage decoration", "quantity": 1, "unit_price": 15000, "amount": 15000},
    {"item_description": "Table centrepieces", "quantity": 30, "unit_price": 450, "amount": 13500},
    {"item_description": "Entrance lighting", "quantity": 4, "unit_price": 1500, "amount": 6000},
]


def test_three_items_render_three_rows(invoice_template, tmp_path):
    template, manifest = invoice_template
    out = str(tmp_path / "out.docx")
    fill = fill_template(template, out, manifest, {**RECORD, "line_items": ITEMS})

    assert fill.qa_passed, fill.qa_notes
    table = docx.Document(out).tables[0]
    assert len(table.rows) == 4  # header + one per item, prototype gone
    assert [c.text for c in table.rows[0].cells] == ["Description", "Qty", "Unit price", "Amount"]
    assert table.rows[1].cells[0].text == "Stage decoration"
    assert table.rows[2].cells[1].text == "30"
    assert table.rows[3].cells[3].text == "6,000.00"
    assert fill.repeat_lineage == [{
        "object_id": "line_items", "iterate_over": "line_items",
        "rows_rendered": 3, "missing": [],
    }]


def test_fields_after_the_table_still_land(invoice_template, tmp_path):
    """The coordinate-integrity property: cloning rows must not shift anything.

    Subtotal, tax and total sit in paragraphs *after* the prototype row. If the
    repeat pass disturbed the span space, these values would land one slot to
    the left -- the exact defect the prototype-processed-last design dissolves.
    """
    template, manifest = invoice_template
    out = str(tmp_path / "out.docx")
    fill = fill_template(template, out, manifest, {**RECORD, "line_items": ITEMS})
    assert fill.qa_passed, fill.qa_notes

    text = "\n".join(p.text for p in docx.Document(out).paragraphs)
    assert "Subtotal: INR 34,500.00" in text
    assert "Tax (18%): 6,210.00" in text
    assert "Total due: INR 40,710.00" in text
    assert "<" not in text  # no token survived anywhere


def test_empty_collection_removes_prototype_and_blocks_when_required(invoice_template, tmp_path):
    template, manifest = invoice_template
    out = str(tmp_path / "out.docx")
    fill = fill_template(template, out, manifest, {**RECORD, "line_items": []})

    table = docx.Document(out).tables[0]
    assert len(table.rows) == 1  # header survives, prototype removed
    # The kit declares required: true, so an invoice with no line items blocks.
    assert not fill.qa_passed
    assert any("at least one row" in n for n in fill.qa_notes)


def test_non_list_collection_blocks(invoice_template, tmp_path):
    template, manifest = invoice_template
    fill = fill_template(template, str(tmp_path / "out.docx"), manifest,
                         {**RECORD, "line_items": "not-a-list"})
    assert not fill.qa_passed
    assert any("should be a list" in n for n in fill.qa_notes)


def test_non_dict_row_blocks(invoice_template, tmp_path):
    template, manifest = invoice_template
    fill = fill_template(template, str(tmp_path / "out.docx"), manifest,
                         {**RECORD, "line_items": ["a bare string"]})
    assert not fill.qa_passed
    assert any("not an object of column values" in n for n in fill.qa_notes)


def test_missing_column_value_blank_policy(invoice_template, tmp_path):
    template, manifest = invoice_template
    out = str(tmp_path / "out.docx")
    items = [{"item_description": "Thing", "quantity": 2, "amount": 20}]
    fill = fill_template(template, out, manifest, {**RECORD, "line_items": items})

    assert fill.qa_passed, fill.qa_notes  # BLANK is the declared policy
    row = docx.Document(out).tables[0].rows[1]
    assert [c.text for c in row.cells] == ["Thing", "2", "", "20.00"]
    assert fill.repeat_lineage[0]["missing"] == [
        {"row": 0, "column": "unit_price", "policy": "BLANK"}]


def test_block_policy_on_a_column_blocks_the_document(invoice_template, tmp_path):
    template, manifest = invoice_template
    manifest = {**manifest, "blocks": [
        {**b, "columns": [
            {**c, "on_missing": "BLOCK"} if c["source_key"] == "unit_price" else c
            for c in b["columns"]]}
        if str(b.get("object_type", "")).upper() == "TABLE_ROW" else b
        for b in manifest["blocks"]
    ]}
    fill = fill_template(template, str(tmp_path / "out.docx"), manifest,
                         {**RECORD, "line_items": [{"item_description": "X", "quantity": 1, "amount": 5}]})
    assert not fill.qa_passed
    assert any("has no value for 'unit_price'" in n for n in fill.qa_notes)


def test_tokens_missing_from_document_block(invoice_template, tmp_path):
    """A manifest promising repetition the document cannot honour fails loudly."""
    template, manifest = invoice_template
    manifest = {**manifest, "blocks": [
        {**b, "columns": [{**c, "token": "<No Such Token>"} for c in b["columns"]]}
        if str(b.get("object_type", "")).upper() == "TABLE_ROW" else b
        for b in manifest["blocks"]
    ]}
    fill = fill_template(template, str(tmp_path / "out.docx"), manifest,
                         {**RECORD, "line_items": ITEMS})
    assert not fill.qa_passed
    assert any("no table row in this template carries its column tokens" in n
               for n in fill.qa_notes)


def test_repeated_fill_is_byte_identical(invoice_template, tmp_path):
    """Row cloning must not cost the §18 reproducibility guarantee."""
    template, manifest = invoice_template
    a, b = str(tmp_path / "a.docx"), str(tmp_path / "b.docx")
    fill_template(template, a, manifest, {**RECORD, "line_items": ITEMS})
    fill_template(template, b, manifest, {**RECORD, "line_items": ITEMS})
    digest = lambda p: hashlib.sha256(open(p, "rb").read()).hexdigest()  # noqa: E731
    assert digest(a) == digest(b)


def test_placeholder_blue_is_stripped_from_clones(invoice_template, tmp_path):
    """A rendered line item is the reader's data, not an unfilled slot."""
    from lxml import etree

    template, manifest = invoice_template
    out = str(tmp_path / "out.docx")
    fill_template(template, out, manifest, {**RECORD, "line_items": ITEMS})

    body = docx.Document(out).element.body
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    table = body.find(f"{w}tbl")
    for color in table.iter(f"{w}color"):
        assert (color.get(f"{w}val") or "").upper() != "0000FF"


def test_validator_accepts_table_row_and_refuses_a_broken_one(invoice_template):
    _template, manifest = invoice_template
    assert validate_manifest(manifest) == []

    broken = {**manifest, "blocks": [
        {"id": "bad_rows", "object_type": "TABLE_ROW", "iterate_over": "", "columns": []}]}
    rules = {f.rule for f in validate_manifest(broken)}
    assert "table_row_without_collection" in rules
    assert "table_row_without_columns" in rules
    # And the old failure must NOT fire: a TABLE_ROW legitimately has no range.
    assert "block_without_range" not in rules


def test_set_row_repeat_operation_round_trip():
    """The typed operation the studio and the co-pilot both post."""
    kit = load_kit("invoice")
    objects = [o for o in kit_objects(kit) if o["object_type"] == "FIELD"]

    result = apply_operations(kit["body"], objects, [{
        "op": "set_row_repeat", "iterate_over": "line_items",
        "columns": [{"token": "<Item Description>"}, {"token": "<Quantity>", "type": "number"},
                    {"token": "<Unit Price>", "type": "currency"},
                    {"token": "<Amount>", "type": "currency"}],
        "required": True,
    }])
    assert result.rejected == [], result.rejected
    rows = [o for o in result.objects if o.get("object_type") == "TABLE_ROW"]
    assert len(rows) == 1
    assert rows[0]["object_id"] == "line_items_rows"
    assert rows[0]["columns"][0]["source_key"] == "item_description"

    # Removing it works; removing it twice is refused, not ignored.
    result2 = apply_operations(result.body, result.objects,
                               [{"op": "remove_row_repeat", "id": "line_items_rows"}])
    assert result2.rejected == []
    assert not [o for o in result2.objects if o.get("object_type") == "TABLE_ROW"]
    result3 = apply_operations(result2.body, result2.objects,
                               [{"op": "remove_row_repeat", "id": "line_items_rows"}])
    assert len(result3.rejected) == 1


def test_set_row_repeat_refuses_tokens_outside_tables():
    kit = load_kit("offer")  # no tables at all
    result = apply_operations(kit["body"], [], [{
        "op": "set_row_repeat", "iterate_over": "line_items",
        "columns": [{"token": "<Annual Salary>"}],
    }])
    assert result.rejected, "a token outside any table must be refused"
    assert "cannot be found" in result.rejected[0]["reason"]


def test_gst_kit_renders_cgst_sgst_split(tmp_path):
    from app.templates.emit_docx import emit

    kit = load_kit("invoice_gst")
    template = str(tmp_path / "gst.docx")
    emit(kit["body"], template)
    manifest = legacy_manifest(kit_objects(kit), delete_always=[])

    record = {
        "business_name": "Sundar Decorations", "business_address": "Pune",
        "business_gstin": "27ABCDE1234F1Z5", "invoice_number": "INV-0001",
        "invoice_date": "2026-09-05", "customer_name": "Acme", "customer_address": "Mumbai",
        "customer_gstin": "27FGHIJ5678K2Z6", "subtotal": 10000,
        "cgst_amount": 900, "sgst_amount": 900, "grand_total": 11800,
        "line_items": [{"item_description": "Decor", "hsn_sac": "9985",
                        "quantity": 1, "unit_price": 10000, "amount": 10000}],
    }
    out = str(tmp_path / "gst_out.docx")
    fill = fill_template(template, out, manifest, record, locale="en_IN")
    assert fill.qa_passed, fill.qa_notes
    text = "\n".join(p.text for p in docx.Document(out).paragraphs)
    assert "CGST: ₹ 900.00" in text
    assert "SGST: ₹ 900.00" in text
    table = docx.Document(out).tables[0]
    assert table.rows[1].cells[1].text == "9985"


def _emit_body(body, path):
    from app.templates.emit_docx import emit
    emit(body, str(path))


def test_two_repeating_tables_bind_their_own_rows(tmp_path):
    """Two collections sharing a column name must each render into their own
    table -- first-any-token binding once sent both to the first row."""
    from app.templates import blueprint as bp
    from app.templates.kits import table_row_objects
    from app.templates.blueprint_lint import legacy_manifest

    def row(*texts, roles=None):
        roles = roles or ["placeholder"] * len(texts)
        return [[bp.paragraph([bp.segment(role, text)])]
                for role, text in zip(roles, texts)]

    body = bp.normalise_body({"blocks": [
        bp.paragraph([bp.segment("static", "Services")]),
        bp.table([row("Item", "Amount", roles=["static", "static"]),
                  row("<Service Description>", "<Amount>")]),
        bp.paragraph([bp.segment("static", "Expenses")]),
        bp.table([row("Expense", "Amount", roles=["static", "static"]),
                  row("<Expense Description>", "<Expense Amount>")]),
    ], "sect_pr_from": None})

    specs = [
        {"id": "services", "iterate_over": "services", "columns": [
            {"token": "<Service Description>", "source_key": "description"},
            {"token": "<Amount>", "source_key": "amount", "type": "currency"}]},
        {"id": "expenses", "iterate_over": "expenses", "columns": [
            {"token": "<Expense Description>", "source_key": "description"},
            {"token": "<Expense Amount>", "source_key": "amount", "type": "currency"}]},
    ]
    objects = table_row_objects(body, specs)
    assert len(objects) == 2
    from app.templates.kits import objects_for
    manifest = legacy_manifest(objects_for(body) + objects, delete_always=[])

    template = tmp_path / "two_tables.docx"
    _emit_body(body, template)
    out = tmp_path / "out.docx"
    fill = fill_template(str(template), str(out), manifest, {
        "services": [{"description": "Design", "amount": 100},
                     {"description": "Build", "amount": 200}],
        "expenses": [{"description": "Travel", "amount": 50}],
    })
    assert fill.qa_passed, fill.qa_notes
    tables = docx.Document(str(out)).tables
    assert len(tables[0].rows) == 3  # header + 2 services
    assert len(tables[1].rows) == 2  # header + 1 expense
    assert tables[1].rows[1].cells[0].text == "Travel"


def test_a_column_token_reused_outside_the_row_blocks(tmp_path):
    """<Amount> in the row AND after "Total payable:" has no scalar value; a
    silent blank total is the one outcome this must never produce."""
    from app.templates import blueprint as bp
    from app.templates.kits import objects_for, table_row_objects
    from app.templates.blueprint_lint import legacy_manifest

    body = bp.normalise_body({"blocks": [
        bp.table([
            [[bp.paragraph([bp.segment("static", "Item")])],
             [bp.paragraph([bp.segment("static", "Amount")])]],
            [[bp.paragraph([bp.segment("placeholder", "<Item Description>")])],
             [bp.paragraph([bp.segment("placeholder", "<Amount>")])]],
        ]),
        bp.paragraph([bp.segment("static", "Total payable: "),
                      bp.segment("placeholder", "<Amount>")]),
    ], "sect_pr_from": None})
    specs = [{"id": "line_items", "iterate_over": "line_items", "columns": [
        {"token": "<Item Description>", "source_key": "item_description"},
        {"token": "<Amount>", "source_key": "amount", "type": "currency"}]}]
    manifest = legacy_manifest(objects_for(body) + table_row_objects(body, specs),
                               delete_always=[])
    template = tmp_path / "reused.docx"
    _emit_body(body, template)
    fill = fill_template(str(template), str(tmp_path / "out.docx"), manifest, {
        "line_items": [{"item_description": "X", "amount": 10}]})
    assert not fill.qa_passed
    assert any("outside the row" in n for n in fill.qa_notes)


def test_author_retry_catches_column_token_reuse():
    from app.compiler.blueprint_author import assemble_body

    seg = lambda role, text: {"role": role, "text": text}  # noqa: E731
    reply = {
        "blocks": [
            {"kind": "table_row", "style": "", "repeat": False, "segments": [],
             "cells": [[seg("static", "Item")], [seg("static", "Amount")]]},
            {"kind": "table_row", "style": "", "repeat": True, "segments": [],
             "cells": [[seg("placeholder", "<Item>")], [seg("placeholder", "<Amount>")]]},
            {"kind": "paragraph", "style": "", "repeat": False, "cells": [],
             "segments": [seg("static", "Total: "), seg("placeholder", "<Amount>")]},
        ],
        "line_items_key": "line_items", "field_types": [], "notes": [],
    }
    _body, _specs, problems = assemble_body(reply)
    assert any("its own token" in p for p in problems)


def test_clinical_kit_renders_disposition_rows_and_totals(tmp_path):
    """The clinical CSR kit is a real TABLE_ROW citizen: typed columns, a
    repeat row per site, and total tokens that are their own fields outside
    the table -- the same contract the invoice kits carry."""
    from app.templates.emit_docx import emit

    kit = load_kit("clinical")
    assert kit["table_rows"], "the CSR kit must declare its repeating row"
    objects = kit_objects(kit)
    rows = [o for o in objects if o["object_type"] == "TABLE_ROW"]
    assert len(rows) == 1
    assert rows[0]["iterate_over"] == "disposition_rows"
    assert {c["type"] for c in rows[0]["columns"]} == {"string", "number"}

    template = str(tmp_path / "csr.docx")
    emit(kit["body"], template)
    manifest = legacy_manifest(objects, delete_always=[])

    record = {
        "study_title": "A Phase 2 study of drug X", "protocol_number": "ONC-2026-014",
        "sponsor_name": "Acme Pharma", "phase": "2", "indication": "NSCLC",
        "document_number": "CSR-0007", "report_date": "2026-09-08",
        "version_label": "1.0",
        "total_subjects_enrolled": 120, "total_subjects_completed": 104,
        "total_subjects_withdrawn": 16,
        "efficacy_summary": "The primary endpoint was met.",
        "safety_summary": "No new signals.", "conclusions": "Continue development.",
        "investigator_name": "Dr. A. Rao",
        "disposition_rows": [
            {"site_name": "Pune General", "subjects_enrolled": 80,
             "subjects_completed": 70, "subjects_withdrawn": 10},
            {"site_name": "Mumbai Central", "subjects_enrolled": 40,
             "subjects_completed": 34, "subjects_withdrawn": 6},
        ],
    }
    out = str(tmp_path / "csr_out.docx")
    fill = fill_template(template, out, manifest, record, locale="en_US")
    assert fill.qa_passed, fill.qa_notes
    document = docx.Document(out)
    table = document.tables[0]
    assert len(table.rows) == 3  # header + two sites, prototype gone
    assert table.rows[2].cells[0].text == "Mumbai Central"
    text = "\n".join(p.text for p in document.paragraphs)
    assert "CSR-0007" in text
    assert "Total enrolled: 120" in text
