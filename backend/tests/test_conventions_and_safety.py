"""Convention families, CJK support, package safety, and the structural gate.

These cover the machinery added so the platform can take the client's China and
Japan template sets, and so an uploaded `.docx` is treated as what it is: an
archive from outside the trust boundary.

The first test is the one that makes the convention-family extraction safe.
Moving the marker grammar out of `manifest_compiler` and into
`app/conventions/en_annotated.yaml` had to change no behaviour at all, and the
303 tests that ran before it are the proof for the templates in the fixture set.
This asserts the same thing at the level of the grammar itself, on strings drawn
from the real client masters -- so a future edit to the YAML that quietly breaks
English is caught here rather than in a golden diff nobody reads.
"""

from __future__ import annotations

import zipfile

import pytest

from app.templates.conventions import (
    DEFAULT_FAMILY, all_families, load_family, normalise_text,
)
from app.templates.parsers.docx_safety import (
    MAX_ENTRIES, MalformedPackageError, UnsafePackageError, inspect_package,
)
from app.generation.korean import choose, has_batchim, resolve
from app.compiler.rule_compiler import parse_condition_instruction


# ------------------------------------------------------- the extraction is safe
@pytest.mark.parametrize("text,expected", [
    ("Include the following text only if the Colleague Type is Full time:",
     ("colleague type", "Full time")),
    ("Include the following text only if the Transaction Type is initiated with recruitment "
     "(e.g. Change Job with Recruitment):",
     ("transaction type", "with recruitment")),
    ("For Full time colleagues:", ("colleague_type", "Full time")),
    ("For transactions initiated without recruitment (e.g. Change Job without Recruitment)",
     ("transaction_type", "without recruitment")),
    # The zone hands the parser the previous branch's trailing instruction glued
    # to the front of the next one; the instruction still has to be found.
    ("Refer to Contact Matrix For transactions initiated without recruitment (e.g. x)",
     ("transaction_type", "without recruitment")),
])
def test_the_default_family_reads_the_english_grammar(text, expected):
    parsed = parse_condition_instruction(text)
    assert parsed is not None, f"did not parse: {text!r}"
    _till, field_text, value = parsed
    assert (field_text.lower(), value) == (expected[0].lower(), expected[1])


@pytest.mark.parametrize("text", [
    "Put the LAB__FT_SALARY__38_HR_ from source file",
    "Refer to Contact Matrix",
    "Dates of Effect",
    "",
])
def test_ordinary_instruction_text_is_not_read_as_a_condition(text):
    assert parse_condition_instruction(text) is None


# ------------------------------------------------------------------ family set
def test_every_shipped_family_loads():
    families = all_families()
    assert DEFAULT_FAMILY in families
    for family_id in families:
        family = load_family(family_id)
        assert family.id == family_id
        assert family.delimiters, f"{family_id} declares no placeholder delimiters"
        # Compiles without raising -- a bad pattern in YAML is a startup error,
        # not a mystery at generation time.
        family.bracket_re()


def test_cjk_families_are_present_for_the_client_estate():
    assert {"zh_annotated", "ja_annotated", "ko_annotated"} <= set(all_families())


@pytest.mark.parametrize("family_id,text,expected_field,expected_value", [
    ("zh_annotated", "适用于全职员工：", "colleague_type", "全职"),
    ("zh_annotated", "全职员工适用：", "colleague_type", "全职"),
    ("zh_annotated", "兼职员工适用：", "colleague_type", "兼职"),
    ("ko_annotated", "정규직 직원용:", "colleague_type", "정규직"),
    ("ko_annotated", "계약직 직원용:", "colleague_type", "계약직"),
    ("ja_annotated", "正社員の方へ：", "colleague_type", "正社員"),
])
def test_cjk_markers_parse(family_id, text, expected_field, expected_value):
    parsed = parse_condition_instruction(text, load_family(family_id))
    assert parsed is not None, f"{family_id} did not read {text!r}"
    _till, field_text, value = parsed
    assert (field_text, value) == (expected_field, expected_value)


def test_fullwidth_and_cjk_delimiters_are_found():
    family = load_family("zh_annotated")
    found = family.bracket_re().findall("姓名：＜姓名＞ 部門：《部門》 職位：【職位】 code:<Code>")
    flat = [g for groups in found for g in (groups if isinstance(groups, tuple) else (groups,)) if g]
    assert set(flat) == {"姓名", "部門", "職位", "Code"}


def test_nfkc_folds_fullwidth_onto_ascii():
    # The reason a CJK master's placeholders and headers bind at all.
    assert normalise_text("＜Ｎａｍｅ＞") == "<Name>"
    assert normalise_text("：") == ":"
    assert normalise_text("１２３") == "123"


# --------------------------------------------------------------------- Korean
@pytest.mark.parametrize("word,expected", [
    ("김민준", True),    # ㄴ final
    ("이서아", False),   # open syllable
    ("서울", True),      # ㄹ final
])
def test_batchim_detection(word, expected):
    assert has_batchim(word) is expected


@pytest.mark.parametrize("text,expected", [
    ("김민준이(가) 입사합니다.", "김민준이 입사합니다."),
    ("이서아이(가) 입사합니다.", "이서아가 입사합니다."),
    ("계약을(를) 체결합니다.", "계약을 체결합니다."),
    # ㄹ-final takes 로, every other final consonant takes 으로.
    ("서울으로(로) 이동합니다.", "서울로 이동합니다."),
    ("부산으로(로) 이동합니다.", "부산으로 이동합니다."),
])
def test_particle_agreement(text, expected):
    resolved, _ambiguous = resolve(text)
    assert resolved == expected


def test_a_non_hangul_value_is_flagged_rather_than_guessed_silently():
    resolved, ambiguous = resolve("Alex이(가) 입사합니다.")
    assert resolved == "Alex가 입사합니다."
    assert ambiguous, "a value whose script cannot decide the particle must be reported"


def test_choose_is_pure():
    assert choose("김민준", "이", "가") == "이"
    assert choose("이서아", "이", "가") == "가"


# -------------------------------------------------------------- package safety
def _zip(path, entries, **kw):
    with zipfile.ZipFile(path, "w") as z:
        for name, data in entries.items():
            z.writestr(name, data, **kw)
    return str(path)


def test_a_real_template_passes(fixtures_dir):
    report = inspect_package(fixtures_dir / "templates" / "hospira_offer.docx")
    assert report.entries > 0
    assert report.total_uncompressed > 0


def test_a_symlink_entry_is_rejected(tmp_path):
    path = tmp_path / "evil.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", "<x/>")
        info = zipfile.ZipInfo("word/link")
        info.create_system = 3
        info.external_attr = (0o120777 << 16)
        z.writestr(info, "/etc/passwd")
    with pytest.raises(UnsafePackageError, match="symlink"):
        inspect_package(path)


@pytest.mark.parametrize("name", ["../../etc/cron.d/x", "/etc/passwd", "a/../../b"])
def test_path_traversal_is_rejected(tmp_path, name):
    path = _zip(tmp_path / "t.docx", {"word/document.xml": "<x/>", name: "x"})
    with pytest.raises(UnsafePackageError, match="escapes"):
        inspect_package(path)


def test_a_decompression_bomb_is_rejected(tmp_path):
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<x/>")
        z.writestr("word/bomb.xml", "0" * (50 * 1024 * 1024))
    with pytest.raises(UnsafePackageError, match="compression ratio|inflates"):
        inspect_package(path)


def test_an_entry_flood_is_rejected(tmp_path):
    path = tmp_path / "many.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", "<x/>")
        for i in range(MAX_ENTRIES + 1):
            z.writestr(f"junk/{i}", "")
    with pytest.raises(UnsafePackageError, match="entries"):
        inspect_package(path)


def test_a_corrupt_file_is_malformed_not_hostile(tmp_path):
    """The two are handled oppositely: hostile is refused, broken is stored with
    an error so the person who uploaded it can see what arrived."""
    path = tmp_path / "notazip.docx"
    path.write_bytes(b"%PDF-1.4 this is not a docx")
    with pytest.raises(MalformedPackageError):
        inspect_package(path)


def test_a_zip_without_a_document_part_is_malformed(tmp_path):
    path = _zip(tmp_path / "empty.docx", {"hello.txt": "hi"})
    with pytest.raises(MalformedPackageError, match="Not a Word document"):
        inspect_package(path)


def test_the_upload_endpoint_refuses_a_hostile_package(app_client, two_orgs):
    token, project_id, _tb, _pb = two_orgs
    import io
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<x/>")
        info = zipfile.ZipInfo("word/link")
        info.create_system = 3
        info.external_attr = (0o120777 << 16)
        z.writestr(info, "/etc/passwd")
    res = app_client.post(
        f"/api/v1/projects/{project_id}/templates",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("evil.docx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 422, res.text
    assert res.json()["detail"]["error"]["code"] == "UNSAFE_TEMPLATE_PACKAGE"


# ------------------------------------------------- the particle pass, end to end
def _korean_letter(path, alternation: str = "이(가)"):
    """A minimal Korean master: static run, blue placeholder, alternation run."""
    import docx as _docx
    from docx.shared import RGBColor

    d = _docx.Document()
    p = d.add_paragraph()
    p.add_run("신입사원 ")
    placeholder = p.add_run("<성명>")
    placeholder.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
    p.add_run(f"{alternation} 입사하였습니다.")
    d.save(str(path))
    return str(path)


@pytest.mark.parametrize("name,expected", [
    ("김민준", "김민준이 입사하였습니다."),   # ㄴ final -> 이
    ("이서아", "이서아가 입사하였습니다."),   # open syllable -> 가
])
def test_the_particle_agrees_with_the_value_that_was_inserted(tmp_path, name, expected):
    """The deciding syllable does not exist until the fill happens, and it lands
    in a different run from the alternation -- which is why the pass reads the
    paragraph and edits the run."""
    import dataclasses

    from app.templates.parsers.docx_prescan import W_NS, prescan
    from app.generation.docx_renderer import fill_template
    from app.compiler.rule_compiler import compile_manifest

    template = _korean_letter(tmp_path / "ko.docx")
    manifest = dataclasses.asdict(compile_manifest(prescan(template), load_family("ko_annotated")))
    assert manifest["particles"] is True, "the Korean family must turn the pass on"

    out = tmp_path / "letter.docx"
    result = fill_template(template, str(out), manifest, {"성명": name})

    import docx as _docx
    text = "".join(t.text or "" for t in _docx.Document(str(out)).element.body.iter(f"{{{W_NS}}}t"))
    assert expected in text
    assert "(" not in text, f"the alternation was left unresolved: {text!r}"
    assert result.qa_passed, result.qa_notes


def test_a_latin_value_before_a_particle_is_reported(tmp_path):
    import dataclasses

    from app.templates.parsers.docx_prescan import prescan
    from app.generation.docx_renderer import fill_template
    from app.compiler.rule_compiler import compile_manifest

    template = _korean_letter(tmp_path / "ko.docx")
    manifest = dataclasses.asdict(compile_manifest(prescan(template), load_family("ko_annotated")))
    result = fill_template(template, str(tmp_path / "letter.docx"), manifest, {"성명": "Alex"})
    assert any("particle" in note for note in result.qa_notes), result.qa_notes


def test_a_template_without_particles_is_untouched(tmp_path, fixtures_dir):
    """The pass rewrites text, so it must not run for a family that did not ask."""
    import dataclasses

    from app.templates.parsers.docx_prescan import prescan
    from app.compiler.rule_compiler import compile_manifest

    manifest = dataclasses.asdict(compile_manifest(prescan(str(fixtures_dir / "templates" / "hospira_offer.docx"))))
    assert manifest["particles"] is False


# ------------------------------------------------------------- structural gate
def test_the_structural_gate_notices_a_dropped_part(tmp_path, fixtures_dir):
    """Text gates read what the letter says; this reads what it is."""
    from app.generation.docx_renderer import _structural_failures

    template = fixtures_dir / "templates" / "hospira_offer.docx"
    mangled = tmp_path / "mangled.docx"
    with zipfile.ZipFile(template) as src, zipfile.ZipFile(mangled, "w") as dst:
        for name in src.namelist():
            if name == "word/styles.xml":
                continue  # a letter that lost its styles renders as plain text
            dst.writestr(name, src.read(name))

    failures = _structural_failures(str(template), str(mangled))
    assert any("styles.xml" in f for f in failures), failures


def test_the_structural_gate_notices_an_edited_part(tmp_path, fixtures_dir):
    from app.generation.docx_renderer import _structural_failures

    template = fixtures_dir / "templates" / "hospira_offer.docx"
    mangled = tmp_path / "mangled.docx"
    with zipfile.ZipFile(template) as src, zipfile.ZipFile(mangled, "w") as dst:
        for name in src.namelist():
            data = src.read(name)
            if name == "word/styles.xml":
                data = data.replace(b"</w:styles>", b"<w:docDefaults/></w:styles>")
            dst.writestr(name, data)

    failures = _structural_failures(str(template), str(mangled))
    assert any("must not touch" in f for f in failures), failures


def test_the_structural_gate_rejects_an_unreadable_output(tmp_path, fixtures_dir):
    from app.generation.docx_renderer import _structural_failures

    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip at all")
    failures = _structural_failures(str(fixtures_dir / "templates" / "hospira_offer.docx"), str(broken))
    assert failures


# ------------------------------------------------------------ profiler warnings
def test_the_annotation_gap_in_the_client_master_is_named(fixtures_dir):
    """The client left `initiated` unhighlighted inside two instruction phrases.
    The zone rule absorbs it at render time; onboarding still has to be told."""
    from app.templates.parsers.docx_prescan import prescan
    from app.compiler.rule_compiler import compile_manifest

    manifest = compile_manifest(prescan(str(fixtures_dir / "templates" / "icc_ct036_original.docx")))
    gaps = [w for w in manifest.warnings if w["code"] == "W-HL-GAP"]
    assert gaps, "the unmarked fragment inside the instruction zone was not reported"
    assert any("initiated" in w["detail"] for w in gaps), gaps
    assert all(w["message"] for w in manifest.warnings), "every warning needs a human-readable reason"


def test_cjk_field_names_do_not_collide():
    """Regression: `[^a-z0-9]+` deleted every CJK character, so `<姓名>` and
    `<部門>` both slugged to the fallback `"field"` and merged into one manifest
    entry -- one placeholder would be filled with the other's value, and nothing
    downstream could tell. Found by the first end-to-end Korean fill."""
    from app.compiler.rule_compiler import _slug

    assert _slug("姓名") != _slug("部門")
    assert _slug("성명") != _slug("부서")
    assert _slug("姓名") not in ("", "field")
    # ASCII slugs must be exactly what they always were; the goldens depend on it.
    assert _slug("Colleague First Name") == "colleague_first_name"
    assert _slug("LAB__FT_SALARY__38_HR_") == "lab_ft_salary_38_hr"
    assert _slug("＜Ｎａｍｅ＞") == "name"


def test_a_deleted_trailing_instruction_does_not_take_the_full_stop_with_it(tmp_path):
    """The last instruction run on a line often carries the sentence's period.
    Deleting it leaves the letter reading `...contact the APAC team` with no
    stop -- which reads as sloppiness rather than as a bug, so nothing catches
    it. The repair restores the punctuation the original paragraph ended with."""
    import dataclasses

    import docx as _docx
    from docx.shared import RGBColor

    from app.templates.parsers.docx_prescan import W_NS, prescan
    from app.generation.docx_renderer import fill_template
    from app.compiler.rule_compiler import compile_manifest

    def red(run):
        run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        return run

    def blue(run):
        run.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
        return run

    path = tmp_path / "switch.docx"
    d = _docx.Document()
    p = d.add_paragraph()
    p.add_run("Please contact ")
    red(p.add_run("For transactions initiated with recruitment (e.g. Change Job with Recruitment)"))
    blue(p.add_run("<Team A>"))
    red(p.add_run("For transactions initiated without recruitment (e.g. Change Job without Recruitment)"))
    blue(p.add_run("<Team B>"))
    red(p.add_run(" Refer to Contact Matrix."))   # carries the sentence's full stop
    d.save(str(path))

    manifest = dataclasses.asdict(compile_manifest(prescan(str(path))))
    assert sum(1 for b in manifest["blocks"] if b.get("start_span") is not None) == 2

    out = tmp_path / "letter.docx"
    result = fill_template(str(path), str(out), manifest, {
        "transaction_type": "With Recruitment", "team_a": "APAC Recruitment", "team_b": "PX Group",
    })
    text = "".join(t.text or "" for t in _docx.Document(str(out)).element.body.iter(f"{{{W_NS}}}t"))

    assert result.qa_passed, result.qa_notes
    assert "Please contact APAC Recruitment" in text
    assert "PX Group" not in text
    assert "Contact Matrix" not in text
    assert text.strip().endswith("."), f"the full stop went with the instruction: {text!r}"


def test_a_particle_split_across_runs_is_reported_not_guessed(tmp_path):
    """Word fragments text across runs for its own reasons. When the split lands
    inside the alternation itself the pass cannot rewrite it without merging
    runs, which would move formatting -- so it says so instead of guessing."""
    import dataclasses

    import docx as _docx
    from docx.shared import RGBColor

    from app.templates.parsers.docx_prescan import prescan
    from app.generation.docx_renderer import fill_template
    from app.compiler.rule_compiler import compile_manifest

    path = tmp_path / "split.docx"
    d = _docx.Document()
    p = d.add_paragraph()
    p.add_run("신입사원 ")
    r = p.add_run("<성명>")
    r.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
    # The alternation straddles a run boundary, and the two runs differ in style
    # so nothing upstream merges them.
    p.add_run("이(")
    p.add_run("가) 입사하였습니다.").bold = True
    d.save(str(path))

    manifest = dataclasses.asdict(compile_manifest(prescan(str(path)), load_family("ko_annotated")))
    result = fill_template(str(path), str(tmp_path / "out.docx"), manifest, {"성명": "김민준"})
    assert any("split across runs" in note for note in result.qa_notes), result.qa_notes


def test_the_structural_gate_notices_an_added_part(tmp_path, fixtures_dir):
    from app.generation.docx_renderer import _structural_failures

    template = fixtures_dir / "templates" / "hospira_offer.docx"
    mangled = tmp_path / "mangled.docx"
    with zipfile.ZipFile(template) as src, zipfile.ZipFile(mangled, "w") as dst:
        for name in src.namelist():
            dst.writestr(name, src.read(name))
        dst.writestr("word/smuggled.xml", "<x/>")

    failures = _structural_failures(str(template), str(mangled))
    assert any("added" in f for f in failures), failures


def test_the_structural_gate_notices_a_malformed_body(tmp_path, fixtures_dir):
    from app.generation.docx_renderer import _structural_failures

    template = fixtures_dir / "templates" / "hospira_offer.docx"
    mangled = tmp_path / "mangled.docx"
    with zipfile.ZipFile(template) as src, zipfile.ZipFile(mangled, "w") as dst:
        for name in src.namelist():
            data = b"<w:document><unclosed>" if name == "word/document.xml" else src.read(name)
            dst.writestr(name, data)

    failures = _structural_failures(str(template), str(mangled))
    assert any("well-formed" in f for f in failures), failures
