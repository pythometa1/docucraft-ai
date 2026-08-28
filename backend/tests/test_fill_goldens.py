"""Golden-file tests for the fill engine — the determinism claim, checked.

`fill_engine` is the one module in this codebase where a regression produces a
*wrong legal document that looks correct*. Content assertions ("the salary
appears") catch the failures you thought of; a golden catches the ones you did
not — a dropped table, a surviving instruction line, a stray empty paragraph
between clauses, a number that quietly changed its grouping.

Two template families, chosen because they fail differently:

* **Hospira** — colour-coded, English instruction prose, MERGEFIELDs, tables.
  Its manifest is *recompiled in-test*, because `compile_manifest` is pure
  deterministic code; that way these tests cover the compiler too.
* **Compensation** — `[[IF]]`/`[[ENDIF]]` control tokens, a delete-when-true
  condition, Indian locale, 51 columns. Its manifest is a **checked-in
  snapshot**: compiling it requires a language model, and a test that needs a
  paid API key is a test that gets skipped. Snapshotting draws the line exactly
  where determinism starts.

Regenerate with `pytest --update-goldens`. A golden that changes in a pull
request without an explanation is the signal this file exists to produce.
"""

from __future__ import annotations

import io
import json
from datetime import date

import docx
import pytest

from app.generation import docx_renderer as fill_engine
from app.generation.source_resolver import apply_binding, suggest_bindings
from app.templates.parsers.docx_prescan import prescan
from app.generation.docx_renderer import fill_template
from app.generation.source_ingestion import extract_records
from app.compiler.rule_compiler import compile_manifest

from golden import assert_docx_equal, assert_matches_golden

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# These templates print the generation date, so without pinning it a golden
# records whatever day it was captured on and every later day fails the byte
# comparison. The value is arbitrary; only its fixedness matters.
GOLDEN_DATE = date(2026, 8, 5)


@pytest.fixture(autouse=True)
def _pin_generation_date(monkeypatch):
    monkeypatch.setattr(fill_engine, "_today", lambda: GOLDEN_DATE)


# --------------------------------------------------------------------- setup
@pytest.fixture(scope="module")
def hospira(fixtures_dir):
    """Rule-compiled in-test: the compiler is deterministic and needs no model."""
    template = fixtures_dir / "templates" / "hospira_offer.docx"
    compiled = compile_manifest(prescan(str(template)))
    manifest = {
        "fields": compiled.fields, "conditions": compiled.conditions,
        "blocks": compiled.blocks, "delete_always": compiled.delete_always,
    }
    columns, records = extract_records(str(fixtures_dir / "records" / "colleagues.xlsx"), "xlsx")
    bindings = suggest_bindings(manifest, columns).as_field_bindings()
    return template, manifest, bindings, {r["Colleague Type"]: r for r in records}


@pytest.fixture(scope="module")
def compensation(fixtures_dir):
    """Manifest loaded from a snapshot -- compiling it needs a model."""
    template = fixtures_dir / "templates" / "compensation_letter.docx"
    snapshot = json.loads((fixtures_dir / "manifests" / "compensation.json").read_text())
    manifest = {k: snapshot[k] for k in ("fields", "conditions", "blocks", "delete_always")}
    columns, records = extract_records(str(fixtures_dir / "records" / "compensation.xlsx"), "xlsx")
    bindings = suggest_bindings(manifest, columns).as_field_bindings()
    return template, manifest, bindings, {r["employee_id"]: r for r in records}


def _text(path) -> str:
    """All text, including inside tables -- `Document.paragraphs` skips those.

    Runs are joined *within* a paragraph with no separator and paragraphs with a
    newline. Joining every `w:t` with a newline instead splits a value from the
    static text beside it: the template writes `$` as its own run, so `$82,000.00`
    becomes `$\\n82,000.00` and a perfectly correct document looks broken.
    """
    d = docx.Document(str(path))
    return "\n".join(
        "".join(t.text or "" for t in p.iter(f"{W}t"))
        for p in d.element.body.iter(f"{W}p")
    )


# ------------------------------------------------- Hospira: cases 1-5, 8, 10
@pytest.mark.parametrize("colleague_type", ["Full Time", "Part Time", "Fixed Term"])
def test_hospira_letter_matches_golden(hospira, goldens_dir, update_goldens, tmp_path, colleague_type):
    template, manifest, bindings, records = hospira
    out = tmp_path / "letter.docx"
    fill_template(str(template), str(out), manifest,
                  apply_binding(records[colleague_type], bindings), locale="en_AU")
    assert_matches_golden(out, goldens_dir / f"hospira_{colleague_type.replace(' ', '_').lower()}",
                          update=update_goldens)


def test_hospira_hyperlink_survives_the_fill(hospira, tmp_path):
    """Hyperlink runs are excluded from span filling; a broken FUSE link in an
    offer letter is a support ticket the reader raises, not the engine."""
    template, manifest, bindings, records = hospira
    out = tmp_path / "letter.docx"
    fill_template(str(template), str(out), manifest, apply_binding(records["Full Time"], bindings))

    before = len(docx.Document(str(template)).element.body.findall(f".//{W}hyperlink"))
    after = len(docx.Document(str(out)).element.body.findall(f".//{W}hyperlink"))
    assert before > 0, "fixture no longer contains a hyperlink -- this test is now vacuous"
    assert after == before


def test_mergefield_currency_carries_exactly_one_symbol(hospira, tmp_path):
    """The template writes `$` as static text before the field, so a formatter
    that emits its own symbol renders `$$82,000.00`."""
    template, manifest, bindings, records = hospira
    out = tmp_path / "letter.docx"
    fill_template(str(template), str(out), manifest, apply_binding(records["Full Time"], bindings))

    text = _text(out)
    assert "$82,000.00" in text
    assert "$$" not in text


def test_fixed_term_letter_keeps_no_other_types_table(hospira, tmp_path):
    """Regression: dropped tables were matched by `id()` across two lxml passes.
    Element proxies are created on demand and freed with their last reference,
    so the deletion intermittently missed and a Fixed Term letter shipped with
    the Full Time remuneration table and raw «LAB__FT_SALARY__38_HR_» codes."""
    template, manifest, bindings, records = hospira
    out = tmp_path / "letter.docx"
    result = fill_template(str(template), str(out), manifest,
                           apply_binding(records["Fixed Term"], bindings))

    assert result.qa_passed, result.qa_notes
    assert "«" not in _text(out)
    joined = " ".join(c.text for t in docx.Document(str(out)).tables for r in t.rows for c in r.cells)
    for figure in ("82,000.00", "48,000.00"):
        assert figure not in joined


# ------------------------------------------- Compensation: cases 11, 12, 13
@pytest.mark.parametrize("employee_id,why", [
    ("IN10001", "one bonus, no relocation, Permanent with probation"),
    ("IN10002", "all three bonuses plus relocation"),
    ("IN10009", "no bonuses at all and an Internship -- exercises the "
                "delete-when-true condition and the non-Permanent branch"),
])
def test_compensation_letter_matches_golden(compensation, goldens_dir, update_goldens,
                                            tmp_path, employee_id, why):
    template, manifest, bindings, records = compensation
    out = tmp_path / "letter.docx"
    fill_template(str(template), str(out), manifest,
                  apply_binding(records[employee_id], bindings), locale="en_IN")
    assert_matches_golden(out, goldens_dir / f"compensation_{employee_id.lower()}",
                          update=update_goldens)


def test_no_control_tokens_or_scaffolding_reach_the_reader(compensation, tmp_path):
    template, manifest, bindings, records = compensation
    out = tmp_path / "letter.docx"
    result = fill_template(str(template), str(out), manifest,
                           apply_binding(records["IN10001"], bindings), locale="en_IN")

    text = _text(out)
    assert result.qa_passed, result.qa_notes
    assert "[[" not in text, "control tokens survived"
    assert "TEMPLATE CONTROL PAGE" not in text, "the legend page reached the letter"


def test_indian_amounts_use_lakh_grouping(compensation, tmp_path):
    """`1,200,000.00` is simply the wrong number in an Indian letter."""
    template, manifest, bindings, records = compensation
    out = tmp_path / "letter.docx"
    fill_template(str(template), str(out), manifest,
                  apply_binding(records["IN10001"], bindings), locale="en_IN")

    text = _text(out)
    assert "12,00,000.00" in text
    assert "1,200,000.00" not in text


def test_a_letter_keeps_only_the_clauses_its_data_calls_for(compensation, tmp_path):
    """Regression: `[[IF bonus = 0]] Delete this entire section [[ENDIF]]` was
    compiled keep-when-true, so the bonus section survived only for people with
    *no* bonuses -- fourteen of twenty letters silently lost a section, and
    every one passed QA."""
    template, manifest, bindings, records = compensation

    def clauses(employee_id):
        out = tmp_path / f"{employee_id}.docx"
        fill_template(str(template), str(out), manifest,
                      apply_binding(records[employee_id], bindings), locale="en_IN")
        text = _text(out)
        return {name: name in text for name in ("Joining bonus:", "Retention bonus:", "Stock units:")}

    assert clauses("IN10001") == {"Joining bonus:": True, "Retention bonus:": False, "Stock units:": False}
    assert clauses("IN10002") == {"Joining bonus:": True, "Retention bonus:": True, "Stock units:": True}
    assert clauses("IN10009") == {"Joining bonus:": False, "Retention bonus:": False, "Stock units:": False}


# ------------------------------------------------------- case 9: idempotence
@pytest.mark.parametrize("family", ["hospira", "compensation"])
def test_the_same_input_produces_the_same_document(request, tmp_path, family):
    """The product's central claim, asserted rather than asserted-about."""
    if family == "hospira":
        template, manifest, bindings, records = request.getfixturevalue("hospira")
        record, locale = records["Full Time"], "en_AU"
    else:
        template, manifest, bindings, records = request.getfixturevalue("compensation")
        record, locale = records["IN10002"], "en_IN"

    bound = apply_binding(record, bindings)
    first, second = tmp_path / "a.docx", tmp_path / "b.docx"
    fill_template(str(template), str(first), manifest, bound, locale=locale)
    fill_template(str(template), str(second), manifest, bound, locale=locale)

    assert_docx_equal(first, second)


# ------------------------------------------------- cases 6-7: loud failures
def _one_paragraph_template(text: str, tmp_path) -> str:
    d = docx.Document()
    d.add_paragraph(text)
    path = tmp_path / "broken.docx"
    buf = io.BytesIO()
    d.save(buf)
    path.write_bytes(buf.getvalue())
    return str(path)


def test_a_leftover_placeholder_fails_qa(tmp_path):
    """The document is still produced -- a reviewer needs to see what went
    wrong -- but it must never be reported as clean."""
    template = _one_paragraph_template("Dear <Colleague First Name>,", tmp_path)
    out = tmp_path / "out.docx"
    result = fill_template(template, str(out),
                           {"fields": [], "conditions": [], "blocks": [], "delete_always": []}, {})

    assert result.qa_passed is False
    assert any("bracket" in note.lower() for note in result.qa_notes), result.qa_notes
    assert out.exists()


def test_a_missing_value_is_named_in_the_lineage(hospira, tmp_path):
    """A field the record cannot supply is recorded as `missing`, not blanked
    and forgotten."""
    template, manifest, bindings, records = hospira
    stripped = {k: v for k, v in apply_binding(records["Full Time"], bindings).items()
                if k != "position_title"}
    out = tmp_path / "out.docx"
    result = fill_template(str(template), str(out), manifest, stripped)

    missing = {e["field_id"] for e in result.field_lineage if e["source"] == "missing"}
    assert "position_title" in missing
