"""§11: template families, manifest inheritance, and the cost of getting it wrong.

Onboarding a template with no family match costs ~15-40 model calls; with a
strong match, ~2-6 (§18). The doc draws the commercial conclusion itself --
"family reuse is the primary cost lever, not model selection" -- so the failure
these tests pin is not a crash. It is an estate that compiles every file from
scratch because nothing remembered what the last one looked like, and, on the
other side, an inherited mapping that arrives already approved and quietly fills
a paragraph that no longer exists.
"""

import io
import json

import docx
import pytest

from app.compiler.confidence import (
    Band, Candidate, band_for_score, exact_name_match, family_inheritance, score,
)
from app.manifests.models import ManifestEnvelope, ManifestObject, ManifestStatus
from app.templates.fingerprint import StructuralFingerprint, structural_similarity
from app.templates.inheritance import (
    FamilyRecord, InheritanceBranch, NewTemplateVersion, REUSE_MANIFEST_SIMILARITY,
    TARGETED_REVIEW_SIMILARITY, as_evidence_only, decide_inheritance,
    fingerprint_from_dict, inherit_manifest, inheritance_summary,
    objects_needing_review, rank_candidates, with_family,
)
from app.templates.semantic_model import (
    Anchor, PROPOSED, context_hash, document_text, paragraph_offsets, static_text_hash,
)


# --------------------------------------------------------------- fingerprints
def _fingerprint(**kw) -> StructuralFingerprint:
    base = dict(
        mergefield_codes=frozenset({"LAB__FT_SALARY__38_HR_", "LAB__FT_TP_38_HR"}),
        paragraph_band=3, table_count=1, table_shapes=(12,),
        blue_span_count=17, red_span_count=27, hyperlink_count=1,
        bracket_tokens=frozenset({"colleague first name", "position title"}),
    )
    base.update(kw)
    return StructuralFingerprint(**base)


def test_a_stored_fingerprint_reads_back_as_the_one_that_was_written():
    """`as_dict` turns the frozensets into lists because JSON has no sets. A
    reader that left them as lists would compare list against frozenset and score
    every stored family at zero -- the estate would look like it had no families
    while the table was full of them."""
    original = _fingerprint()
    restored = fingerprint_from_dict(original.as_dict())

    assert restored == original
    assert structural_similarity(original, restored) == 1.0


def test_a_fingerprint_with_a_key_this_build_cannot_read_raises():
    """A newer writer's extra signal, silently ignored, produces a similarity
    score that is quietly wrong rather than absent -- and a wrong score decides
    whether a customer pays the 2-call or the 40-call onboarding."""
    payload = _fingerprint().as_dict()
    payload["footnote_count"] = 3

    with pytest.raises(ValueError, match="footnote_count"):
        fingerprint_from_dict(payload)


# ------------------------------------------------------------- the thresholds
def _inherited_field_score(similarity: float) -> float:
    """§13's own arithmetic for an inherited field that also matches by name."""
    return score(Candidate(
        target="new_reporting_to", source_ref="source.new_manager_name",
        signals=(exact_name_match(True), family_inheritance(similarity)),
    ))


def test_the_reuse_threshold_is_where_an_inherited_mapping_reaches_confirm():
    """0.90 is derived from §13, not chosen for looking decisive: at that
    similarity an inherited field that also matches by name scores 0.813, inside
    the Confirm band, which is what "one click to accept" means. At 0.85 the same
    field scores 0.799 and lands in Review, where §13 requires ranked
    alternatives -- so the branch would have saved the reviewer nothing while
    claiming it had."""
    at_threshold = _inherited_field_score(REUSE_MANIFEST_SIMILARITY)
    just_below = _inherited_field_score(0.85)

    assert band_for_score(at_threshold, vetoed=False, money_or_identifier=False) is Band.CONFIRM
    assert band_for_score(just_below, vetoed=False, money_or_identifier=False) is Band.REVIEW


def test_the_targeted_review_floor_matches_the_clustering_threshold():
    """Two numbers that both mean "same family" and can drift apart is how an
    estate gets clustered one way at upload and matched another way at
    inheritance."""
    from app.templates.family_matcher import SIMILARITY_THRESHOLD

    assert TARGETED_REVIEW_SIMILARITY == SIMILARITY_THRESHOLD


def test_a_medium_match_stays_in_the_review_band():
    value = _inherited_field_score(TARGETED_REVIEW_SIMILARITY)
    assert band_for_score(value, vetoed=False, money_or_identifier=False) is Band.REVIEW


# ------------------------------------------------------------- family matching
def _record(family_id: str, fingerprint, *, approved="mf_approved", name=None) -> FamilyRecord:
    return FamilyRecord(
        family_id=family_id, name=name or family_id, fingerprint=fingerprint,
        representative_template_version_id=f"tv_{family_id}", approved_manifest_id=approved,
    )


def test_candidates_rank_closest_first_and_break_ties_stably():
    """A ranking that reorders itself between two calls makes an onboarding
    decision nobody can explain afterwards."""
    near = _fingerprint()
    far = _fingerprint(bracket_tokens=frozenset({"invoice number"}))

    ranked = rank_candidates(near, [_record("fam_b", far), _record("fam_a", near),
                                    _record("fam_c", near)])

    assert [c.family_id for c in ranked] == ["fam_a", "fam_c", "fam_b"]


def test_a_tenant_with_no_families_compiles_as_a_new_one():
    decision = decide_inheritance(_fingerprint(), [])

    assert decision.branch is InheritanceBranch.NEW_FAMILY
    assert decision.family_id is None
    assert not decision.joins_existing_family
    assert "no template families yet" in decision.reason


def test_a_very_close_match_reuses_the_approved_manifest():
    decision = decide_inheritance(_fingerprint(), [_record("fam_offer", _fingerprint())])

    assert decision.branch is InheritanceBranch.REUSE_MANIFEST
    assert decision.reuses_manifest
    assert decision.parent_manifest_id == "mf_approved"
    assert decision.similarity == 1.0


def test_a_medium_match_reuses_the_evidence_and_asks_for_a_targeted_review():
    """The translated-template case: same MERGEFIELD codes and layout, different
    placeholder vocabulary. Close enough to be evidence, not close enough to
    carry an approval across."""
    german = _fingerprint(bracket_tokens=frozenset({"vorname des kollegen",
                                                    "positionsbezeichnung"}))
    decision = decide_inheritance(_fingerprint(), [_record("fam_offer", german)])

    assert decision.branch is InheritanceBranch.REUSE_EVIDENCE
    assert decision.reuses_evidence
    assert TARGETED_REVIEW_SIMILARITY <= decision.similarity < REUSE_MANIFEST_SIMILARITY


def test_a_distant_family_is_not_inherited_from_at_all():
    """Correcting a wrong inherited mapping costs a reviewer more attention than
    making a fresh one, so below the floor reuse stops paying."""
    invoice = StructuralFingerprint(
        mergefield_codes=frozenset({"INV_TOTAL"}), paragraph_band=0, table_count=1,
        table_shapes=(3,), blue_span_count=2, red_span_count=0, hyperlink_count=0,
        bracket_tokens=frozenset({"invoice number"}),
    )
    decision = decide_inheritance(_fingerprint(), [_record("fam_invoice", invoice)])

    assert decision.branch is InheritanceBranch.NEW_FAMILY
    assert decision.family_id is None
    assert decision.parent_manifest_id is None


def test_a_close_family_with_no_approved_manifest_is_joined_but_not_inherited():
    """§11 says "nearest *approved* family", and the word carries weight: a
    family whose manifest is still a draft has no signed-off knowledge to lend.
    The template still joins it, so the next upload of this document type finds
    one family rather than a row per upload."""
    decision = decide_inheritance(
        _fingerprint(), [_record("fam_offer", _fingerprint(), approved=None)]
    )

    assert decision.branch is InheritanceBranch.NEW_FAMILY
    assert decision.family_id == "fam_offer"
    assert decision.joins_existing_family
    assert decision.parent_manifest_id is None


def test_a_decision_that_reuses_a_manifest_must_name_it():
    from app.templates.inheritance import InheritanceDecision

    with pytest.raises(ValueError, match="must name the manifest"):
        InheritanceDecision(branch=InheritanceBranch.REUSE_MANIFEST, reason="because")


# ----------------------------------------------------------- inheriting objects
PAD = ("This paragraph is deliberately long so that an edit elsewhere in the document stays "
       "well outside the sixty-character window the anchor context hash covers.")

PARENT_PARAGRAPHS = (
    "Dear <Colleague First Name>,",
    PAD,
    "You will report to <New Reporting To>.",
    PAD,
    "Countersigned by <Regional Head>.",
)

# The revision: same document, the countersignature paragraph dropped.
CHILD_PARAGRAPHS = PARENT_PARAGRAPHS[:4]


def _anchor(paragraphs, token: str, paragraph_index: int) -> dict:
    whole = document_text(paragraphs)
    at = paragraph_offsets(paragraphs)[paragraph_index] + paragraphs[paragraph_index].index(token)
    return Anchor(
        kind="run_path", path=f"body/p[{paragraph_index}]/r[1]", ordinal=1,
        token=token, context_hash=context_hash(whole, at),
    ).as_dict()


def _approved_field(object_id: str, token: str, paragraph_index: int) -> ManifestObject:
    return ManifestObject(object_id=object_id, object_type="FIELD", attributes={
        "anchor": _anchor(PARENT_PARAGRAPHS, token, paragraph_index),
        "token": token,
        "source_ref": f"source.{object_id}",
        "format": None,
        "on_missing": "BLOCK",
        "status": "APPROVED",
        "approved_by": "u_1042",
        "approved_at": "2026-08-25T10:58:02Z",
        "evidence": ["exact_name_match"],
    })


def _parent_manifest(*extra) -> ManifestEnvelope:
    return ManifestEnvelope(
        manifest_id="mf_parent", manifest_version=7, status=ManifestStatus.LOCKED,
        organization_id="org_a", template_version_id="tv_parent",
        template_family_id="fam_offer", template_hash="sha256:" + "c" * 64,
        source_schema_ref="workday/hr_letter_extract",
        source_schema_hash="sha256:" + "d" * 64,
        qa_policy={"blocking": ["placeholder_check"]},
        objects=(
            _approved_field("greeting", "<Colleague First Name>", 0),
            _approved_field("reporting_to", "<New Reporting To>", 2),
            _approved_field("regional_head", "<Regional Head>", 4),
        ) + extra,
    )


def _target(**kw) -> NewTemplateVersion:
    payload = dict(template_version_id="tv_child", inventory=CHILD_PARAGRAPHS,
                   template_hash="sha256:" + "e" * 64)
    payload.update(kw)
    return NewTemplateVersion(**payload)


def _inherit(parent=None, **kw) -> ManifestEnvelope:
    return inherit_manifest(parent or _parent_manifest(), _target(),
                            manifest_id="mf_child", similarity=0.98, **kw)


def test_an_inherited_manifest_is_a_draft_for_the_new_template():
    draft = _inherit()

    assert draft.status is ManifestStatus.DRAFT
    assert draft.manifest_version == 1
    assert draft.template_version_id == "tv_child"
    assert draft.template_hash == "sha256:" + "e" * 64
    assert draft.template_family_id == "fam_offer"
    # Not a supersession: the parent belongs to a different template version and
    # stays LOCKED and runnable.
    assert draft.supersedes is None
    assert draft.manifest_hash is None


def test_inherited_objects_are_proposed_and_carry_no_approval_stamp():
    """The approval belonged to a review of a different document. Carrying the
    stamp across would make the audit trail claim a named person signed off a
    letter they have never seen, and §13's bands would never get a chance to
    decide."""
    draft = _inherit()
    greeting = draft.object_by_id("greeting")

    assert greeting.attributes["status"] == PROPOSED
    assert "approved_by" not in greeting.attributes
    assert "approved_at" not in greeting.attributes


def test_inheritance_records_the_measured_similarity_as_evidence():
    """§13's family_inheritance signal is "scaled by measured family similarity",
    and §6's worked FIELD example carries `family_inheritance:0.98`. A reviewer
    told only that something was inherited cannot weigh it."""
    greeting = _inherit().object_by_id("greeting")

    assert "family_inheritance:0.98" in greeting.attributes["evidence"]
    assert "exact_name_match" in greeting.attributes["evidence"]
    assert greeting.attributes["inherited_from"]["manifest_id"] == "mf_parent"
    assert greeting.attributes["inherited_from"]["object_id"] == "greeting"


def test_an_object_whose_anchor_no_longer_resolves_comes_back_for_review():
    """Not dropped. A dropped object leaves the reviewer a shorter, cleaner
    manifest and no indication that a paragraph the parent filled is now
    unaddressed -- §19's silently blank required field, arriving as a legally
    defective letter."""
    draft = _inherit()
    orphan = draft.object_by_id("regional_head")

    assert orphan is not None, "the object was dropped instead of flagged"
    assert orphan.attributes["needs_review"] is True
    assert "no longer resolves" in orphan.attributes["inheritance_note"]
    assert [o.object_id for o in objects_needing_review(draft)] == ["regional_head"]


def test_an_object_whose_anchor_still_resolves_is_not_flagged():
    draft = _inherit()

    assert draft.object_by_id("reporting_to").attributes["needs_review"] is False
    assert "inheritance_note" not in draft.object_by_id("reporting_to").attributes


def test_a_static_region_whose_text_is_gone_comes_back_for_review():
    """§6: a changed STATIC region forces a new template version rather than a
    re-approval, so inheriting one whose text has vanished must not pass."""
    parent = _parent_manifest(
        ManifestObject(object_id="pad", object_type="STATIC",
                       attributes={"text_hash": static_text_hash(PAD)}),
        ManifestObject(object_id="countersignature", object_type="STATIC",
                       attributes={"text_hash": static_text_hash(PARENT_PARAGRAPHS[4])}),
    )
    draft = _inherit(parent)

    assert draft.object_by_id("pad").attributes["needs_review"] is False
    assert draft.object_by_id("countersignature").attributes["needs_review"] is True
    assert "byte-for-byte" in draft.object_by_id("countersignature").attributes["inheritance_note"]


def test_a_legacy_object_addressed_only_by_token_is_checked_for_uniqueness():
    """The manifests this build compiled before anchors existed address a slot by
    token. A bare token is a weaker address, but "does it still appear exactly
    once" is the same question §6 asks, and asking it badly beats not asking."""
    parent = _parent_manifest(
        ManifestObject(object_id="pad_token", object_type="FIELD", attributes={
            "token": PAD, "source_ref": "source.pad", "format": None,
            "on_missing": "BLANK", "status": "APPROVED",
        }),
    )
    flagged = _inherit(parent).object_by_id("pad_token")

    assert flagged.attributes["needs_review"] is True
    assert "appears 2 times" in flagged.attributes["inheritance_note"]


def test_an_anchor_bearing_object_with_no_address_at_all_is_not_waved_through():
    parent = _parent_manifest(
        ManifestObject(object_id="mystery", object_type="SIGNATURE", attributes={
            "signer_source_ref": "source.signer", "image_policy": "NONE", "esign_ref": None,
        }),
    )
    flagged = _inherit(parent).object_by_id("mystery")

    assert flagged.attributes["needs_review"] is True
    assert "§6 requires a SIGNATURE to carry an anchor" in flagged.attributes["inheritance_note"]


def test_a_positionally_addressed_region_cannot_be_revalidated_and_says_so():
    parent = _parent_manifest(
        ManifestObject(object_id="temp_block", object_type="SECTION", attributes={
            "anchor_range": {"start_index": 2, "end_index": 4}, "repeat_over": None,
            "ordering": None, "empty_behaviour": "REMOVE",
        }),
    )
    flagged = _inherit(parent).object_by_id("temp_block")

    assert flagged.attributes["needs_review"] is True
    assert "position rather than a reference" in flagged.attributes["inheritance_note"]


def test_the_source_contract_travels_with_the_mappings():
    """Mappings are only meaningful against the schema they were compiled for.
    An empty pin would let the new template read a different extract without
    anybody deciding to."""
    draft = _inherit()

    assert draft.source_schema_ref == "workday/hr_letter_extract"
    assert draft.source_schema_hash == "sha256:" + "d" * 64
    assert draft.qa_policy == {"blocking": ["placeholder_check"]}


def test_inheriting_from_an_unapproved_manifest_is_refused():
    """A draft manifest's mappings are exactly the ones nobody has checked.
    Spreading them across a family would multiply one unreviewed guess by the
    size of the estate."""
    from dataclasses import replace

    draft_parent = replace(_parent_manifest(), status=ManifestStatus.DRAFT)
    with pytest.raises(ValueError, match="nearest \\*approved\\* family"):
        _inherit(draft_parent)


def test_inheritance_never_crosses_a_tenant():
    """§19: "cross-tenant leakage via family matching -- contract breach;
    potentially existential". This function is where it would happen, because
    structure looks the same across customers."""
    with pytest.raises(ValueError, match="never crosses a tenant"):
        inherit_manifest(_parent_manifest(), _target(organization_id="org_b"),
                         manifest_id="mf_child", similarity=0.98)


def test_similarity_outside_the_zero_to_one_range_is_refused():
    with pytest.raises(ValueError, match="0..1"):
        inherit_manifest(_parent_manifest(), _target(), manifest_id="mf_child", similarity=1.4)


def test_an_inherited_manifest_needs_its_own_id():
    with pytest.raises(ValueError, match="manifest_id"):
        inherit_manifest(_parent_manifest(), _target(), manifest_id="", similarity=0.9)


def test_the_medium_branch_flags_every_inherited_object():
    """"Reuse mapping evidence + targeted review" is not "reuse the manifest with
    a shorter review". At medium similarity an object that re-anchors cleanly has
    only proved a token exists in both documents."""
    draft = as_evidence_only(_inherit())

    assert all(o.attributes["needs_review"] for o in draft.objects)
    assert "medium band" in draft.object_by_id("reporting_to").attributes["inheritance_note"]
    # The object that actually broke keeps the sharper note.
    assert "no longer resolves" in draft.object_by_id("regional_head").attributes["inheritance_note"]


def test_the_summary_counts_what_was_carried_and_what_survived():
    summary = inheritance_summary(_inherit())

    assert summary["objects_total"] == 3
    assert summary["objects_inherited"] == 3
    assert summary["objects_needing_review"] == 1
    assert summary["needs_review"][0]["object_id"] == "regional_head"


def test_family_membership_is_recorded_without_mutating_the_manifest():
    draft = _inherit()
    moved = with_family(draft, "fam_apac")

    assert moved.template_family_id == "fam_apac"
    assert draft.template_family_id == "fam_offer", "the envelope is frozen for a reason"


# ------------------------------------------------------------------ endpoints
def _template_bytes(*, with_countersignature: bool = True) -> bytes:
    d = docx.Document()
    d.add_heading("Offer", level=1)
    d.add_paragraph("Dear <Colleague First Name>,")
    d.add_paragraph("You are offered the role of <Position Title>.")
    d.add_heading("Remuneration", level=1)
    d.add_paragraph("Your salary will be <Annual Salary> per year.")
    d.add_paragraph("You will report to <New Reporting To>.")
    if with_countersignature:
        d.add_paragraph("Countersigned by <Regional Head>.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _upload(app_client, token, project_id, name, payload) -> dict:
    res = app_client.post(
        f"/api/v1/projects/{project_id}/templates", headers=_auth(token),
        files={"file": (name, payload, DOCX_TYPE)},
    )
    assert res.status_code == 201, res.text
    return res.json()


def _approve_a_manifest_for(template_version_id: str, org_id: str, user_id: str,
                            template_file_id: str) -> str:
    """Insert a locked manifest directly.

    Driving the approval endpoint would make this test depend on manifest
    validation rather than on inheritance, which is what it is about.
    """
    from app.db import SessionLocal
    from app.models import TemplateManifest

    db = SessionLocal()
    try:
        manifest = TemplateManifest(
            org_id=org_id, template_file_id=template_file_id,
            template_version_id=template_version_id, version_no=1, status="approved",
            fields=[
                {"id": "reporting_to", "token": "<New Reporting To>",
                 "source_ref": "source.new_manager_name", "format": None,
                 "on_missing": "BLOCK", "status": "APPROVED", "approved_by": user_id},
                {"id": "regional_head", "token": "<Regional Head>",
                 "source_ref": "source.regional_head", "format": None,
                 "on_missing": "BLOCK", "status": "APPROVED", "approved_by": user_id},
            ],
            conditions=[], blocks=[], delete_always=[], confidence=1.0,
            compiled_by="rule_based", prescan_summary={}, created_by=user_id,
        )
        db.add(manifest)
        db.commit()
        return manifest.id
    finally:
        db.close()


@pytest.fixture(scope="module")
def org_a(two_orgs):
    from app.db import SessionLocal
    from app.models import User

    token_a, project_a, _token_b, _project_b = two_orgs
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.email == "user-a@tenant.test").one()
        return token_a, project_a, user.org_id, user.id
    finally:
        db.close()


def test_the_first_template_of_its_kind_mints_a_family(app_client, org_a):
    """"Approved family knowledge grows" only if the fingerprint is kept. A
    cluster row records who arrived together; without a family row the next
    upload has nothing to compare against and pays the no-match price again."""
    token, project_id, _org_id, _user_id = org_a
    template = _upload(app_client, token, project_id, "fam_first.docx", _template_bytes())

    res = app_client.post(
        f"/api/v1/templates/{template['id']}/inherit-manifest", headers=_auth(token), json={},
    )

    assert res.status_code == 201, res.text
    body = res.json()
    # The public outcome is a plain word; the branch enum names the method.
    assert body["decision"]["outcome"] == "new"
    assert "branch" not in body["decision"]
    assert body["family_created"] is True
    assert body["family_id"]
    assert body["manifest"] is None


def test_a_revision_inherits_the_approved_manifest_and_reports_the_delta(app_client, org_a):
    token, project_id, org_id, user_id = org_a
    parent_template = _upload(app_client, token, project_id, "fam_parent.docx", _template_bytes())
    parent_manifest_id = _approve_a_manifest_for(
        parent_template["current_version_id"], org_id, user_id, parent_template["id"]
    )
    seed = app_client.post(
        f"/api/v1/templates/{parent_template['id']}/inherit-manifest",
        headers=_auth(token), json={},
    ).json()
    assert seed["family_id"]

    revision = _upload(app_client, token, project_id, "fam_revision.docx",
                       _template_bytes(with_countersignature=False))
    res = app_client.post(
        f"/api/v1/templates/{revision['id']}/inherit-manifest", headers=_auth(token), json={},
    )

    assert res.status_code == 201, res.text
    body = res.json()
    assert body["decision"]["outcome"] == "reused", body["decision"]["reason"]
    assert body["decision"]["parent_manifest_id"] == parent_manifest_id
    # No similarity and no threshold on the wire, in the reason or anywhere else.
    import re as _re
    _dump = json.dumps(body)
    assert "similarity" not in _dump, _re.findall(r'.{150}similarity.{150}', _dump)
    assert not any(ch.isdigit() for ch in body["decision"]["reason"])
    assert all("similarity" not in c for c in body["decision"]["considered"])

    # The measurement is still taken and kept: on the row, for the operator.
    from app.db import SessionLocal
    from app.models import TemplateManifest

    db = SessionLocal()
    try:
        row = db.get(TemplateManifest, body["manifest"]["id"])
        assert row.prescan_summary["inheritance"]["branch"] == "REUSE_MANIFEST"
        assert row.confidence >= 0.9
    finally:
        db.close()

    # The draft is a draft. Nothing here is approved by inheriting it.
    assert body["manifest"]["status"] == "draft"
    assert {f["status"] for f in body["manifest"]["fields"]} == {"PROPOSED"}

    # The countersignature paragraph is gone from this revision, so the mapping
    # that filled it comes back flagged rather than disappearing.
    flagged = {o["object_id"] for o in body["inheritance"]["needs_review"]}
    assert flagged == {"regional_head"}
    assert body["inheritance"]["objects_inherited"] == 2

    # And the diff is what makes "review changed objects only" real.
    assert body["diff"]["counts"]["changed"] == 2
    assert body["diff"]["counts"]["added"] == 0


def test_the_diff_endpoint_compares_two_manifests_of_the_same_tenant(app_client, org_a):
    token, project_id, org_id, user_id = org_a
    template = _upload(app_client, token, project_id, "fam_diff.docx", _template_bytes())
    first = _approve_a_manifest_for(
        template["current_version_id"], org_id, user_id, template["id"]
    )
    second = _approve_a_manifest_for(
        template["current_version_id"], org_id, user_id, template["id"]
    )

    res = app_client.get(
        f"/api/v1/template-manifests/{first}/diff", headers=_auth(token),
        params={"against": second},
    )

    assert res.status_code == 200, res.text
    body = res.json()
    assert body["identical"] is True
    assert body["counts"] == {"added": 0, "removed": 0, "changed": 0, "unchanged": 2}
    assert body["review_object_ids"] == []


def test_the_diff_endpoint_scopes_the_other_manifest_too(app_client, two_orgs, org_a):
    """Scoping only the path parameter would turn `against` into a read of any
    manifest in the estate -- the diff prints field mappings, source references
    and anchor tokens."""
    from app.db import SessionLocal
    from app.models import TemplateManifest, User

    token_a, project_a, org_id, user_id = org_a
    _ta, _pa, token_b, _project_b = two_orgs

    template = _upload(app_client, token_a, project_a, "fam_leak.docx", _template_bytes())
    mine = _approve_a_manifest_for(
        template["current_version_id"], org_id, user_id, template["id"]
    )

    db = SessionLocal()
    try:
        other_user = db.query(User).filter(User.email == "user-b@tenant.test").one()
        theirs = TemplateManifest(
            org_id=other_user.org_id, template_file_id=None, template_version_id="tv_theirs",
            version_no=1, status="approved", fields=[], conditions=[], blocks=[],
            delete_always=[], confidence=1.0, compiled_by="rule_based", prescan_summary={},
            created_by=other_user.id,
        )
        db.add(theirs)
        db.commit()
        theirs_id = theirs.id
    finally:
        db.close()

    res = app_client.get(
        f"/api/v1/template-manifests/{theirs_id}/diff", headers=_auth(token_a),
        params={"against": mine},
    )
    assert res.status_code == 404

    res = app_client.get(
        f"/api/v1/template-manifests/{mine}/diff", headers=_auth(token_a),
        params={"against": theirs_id},
    )
    assert res.status_code == 404, "the `against` manifest was not tenant-scoped"


def test_inheriting_from_an_unapproved_parent_is_refused_by_the_endpoint(app_client, org_a):
    from app.db import SessionLocal
    from app.models import TemplateManifest

    token, project_id, org_id, user_id = org_a
    template = _upload(app_client, token, project_id, "fam_draft_parent.docx", _template_bytes())
    db = SessionLocal()
    try:
        draft = TemplateManifest(
            org_id=org_id, template_file_id=template["id"],
            template_version_id="tv_elsewhere", version_no=1, status="draft",
            fields=[], conditions=[], blocks=[], delete_always=[], confidence=0.5,
            compiled_by="rule_based", prescan_summary={}, created_by=user_id,
        )
        db.add(draft)
        db.commit()
        draft_id = draft.id
    finally:
        db.close()

    res = app_client.post(
        f"/api/v1/templates/{template['id']}/inherit-manifest", headers=_auth(token),
        json={"parent_manifest_id": draft_id},
    )

    assert res.status_code == 409
    assert res.json()["detail"]["error"]["code"] == "PARENT_NOT_APPROVED"


def test_a_template_with_no_parsed_version_cannot_be_fingerprinted(app_client, org_a):
    """Fingerprinting reads the stored binary. A template that never parsed has
    no structure to compare, and answering "no family match" for it would bill
    the customer a from-scratch compile for a file that cannot be compiled."""
    from app.db import SessionLocal
    from app.models import TemplateFile

    token, project_id, org_id, user_id = org_a
    db = SessionLocal()
    try:
        orphan = TemplateFile(org_id=org_id, project_id=project_id, name="unparsed.docx",
                              status="uploaded", created_by=user_id)
        db.add(orphan)
        db.commit()
        orphan_id = orphan.id
    finally:
        db.close()

    res = app_client.post(
        f"/api/v1/templates/{orphan_id}/inherit-manifest", headers=_auth(token), json={},
    )
    assert res.status_code == 400
    assert res.json()["detail"]["error"]["code"] == "TEMPLATE_NOT_PARSED"


def test_a_template_cannot_inherit_from_its_own_template_version(app_client, org_a):
    """The result would be a draft that duplicates the manifest already running
    against these exact bytes -- two live manifests for one template version, and
    no way to say which one production means."""
    token, project_id, org_id, user_id = org_a
    template = _upload(app_client, token, project_id, "fam_self.docx", _template_bytes())
    mine = _approve_a_manifest_for(
        template["current_version_id"], org_id, user_id, template["id"]
    )

    res = app_client.post(
        f"/api/v1/templates/{template['id']}/inherit-manifest", headers=_auth(token),
        json={"parent_manifest_id": mine},
    )

    assert res.status_code == 409
    assert res.json()["detail"]["error"]["code"] == "PARENT_IS_SELF"


def test_a_signature_object_is_carried_across_rather_than_refusing_the_inheritance(
        app_client, org_a):
    """This used to answer 422 MANIFEST_OBJECTS_UNSTORABLE.

    `template_manifests` held three object lists and §6 defines ten types, so an
    approved SIGNATURE object had nowhere to go. Writing the row anyway would
    have dropped it on the floor and reported a successful inheritance -- a
    letter that inherits everything except the part that makes it binding -- so
    refusing was the honest answer available at the time.

    The row now carries an `objects` column and the write is lossless, which
    matters because a real offer letter has a signature block: the refusal was
    the common case, not the edge one. The object comes back PROPOSED, as every
    inherited object does -- an approval given to one document is not an
    approval of another."""
    from app.db import SessionLocal
    from app.models import TemplateManifest

    token, project_id, org_id, user_id = org_a
    parent_template = _upload(app_client, token, project_id, "fam_sig_parent.docx",
                              _template_bytes())
    db = SessionLocal()
    try:
        manifest = TemplateManifest(
            org_id=org_id, template_file_id=parent_template["id"],
            template_version_id=parent_template["current_version_id"], version_no=1,
            status="approved",
            fields=[{"id": "signature", "object_type": "SIGNATURE",
                     "signer_source_ref": "source.signer", "image_policy": "NONE",
                     "esign_ref": None, "status": "APPROVED"}],
            conditions=[], blocks=[], delete_always=[], confidence=1.0,
            compiled_by="rule_based", prescan_summary={}, created_by=user_id,
        )
        db.add(manifest)
        db.commit()
        parent_manifest_id = manifest.id
    finally:
        db.close()

    revision = _upload(app_client, token, project_id, "fam_sig_revision.docx",
                       _template_bytes(with_countersignature=False))
    res = app_client.post(
        f"/api/v1/templates/{revision['id']}/inherit-manifest", headers=_auth(token),
        json={"parent_manifest_id": parent_manifest_id},
    )

    assert res.status_code == 201, res.text
    body = res.json()

    from app.db import SessionLocal as _Session
    from app.models import TemplateManifest as _Manifest

    db = _Session()
    try:
        row = db.get(_Manifest, body["manifest"]["id"])
        stored = {o["object_id"]: o for o in row.objects}
    finally:
        db.close()

    assert "signature" in stored, "the SIGNATURE object was dropped, not carried"
    assert stored["signature"]["object_type"] == "SIGNATURE"
    assert stored["signature"]["signer_source_ref"] == "source.signer"
    # Inherited, therefore proposed: the approval on the parent was given to a
    # different document.
    assert stored["signature"]["status"] == "PROPOSED"


def test_a_family_whose_stored_fingerprint_cannot_be_read_fails_loudly(app_client, org_a):
    """Scoring the half of a fingerprint this build understands would produce a
    similarity that is quietly wrong, and a wrong similarity is what decides
    whether a mapping is carried across unreviewed."""
    from app.db import SessionLocal
    from app.models import TemplateFamily

    token, project_id, org_id, user_id = org_a
    template = _upload(app_client, token, project_id, "fam_unreadable.docx", _template_bytes())

    db = SessionLocal()
    try:
        broken = TemplateFamily(
            org_id=org_id, name="from a newer build",
            fingerprint={"footnote_count": 3},
            representative_template_version_id=template["current_version_id"],
        )
        db.add(broken)
        db.commit()
        broken_id = broken.id
    finally:
        db.close()

    try:
        res = app_client.post(
            f"/api/v1/templates/{template['id']}/inherit-manifest", headers=_auth(token), json={},
        )
        assert res.status_code == 500
        assert res.json()["detail"]["error"]["code"] == "FAMILY_FINGERPRINT_UNREADABLE"
    finally:
        # Leaving it behind would make every later match in this database raise.
        db = SessionLocal()
        try:
            db.delete(db.get(TemplateFamily, broken_id))
            db.commit()
        finally:
            db.close()
