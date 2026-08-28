"""The gate that reads the letter for what is *not* in it.

The document these tests were written against told a part-time colleague they
were "employed on a full time basis", lost its offer sentence, its date, its
hours of work and nine other paragraphs, and reported `qa_passed: True`. Eight
fields the user had filled in resolved to real values and reached nothing,
because the paragraphs holding them had been deleted before the fill ran -- so
every leftover-scaffolding gate was correct to stay silent.

The true negatives here matter at least as much as the positives. This is the
check most able to stop a letter that is entirely correct: a value is absent for
a good reason whenever its block lost a condition, and the string on the page is
routinely not the string in the record. A gate that fires on a correct document
gets overridden, and an overridden gate protects nothing.
"""

import docx
import pytest

from app.qa.content_loss import (
    RESOLVED_VALUE_ABSENT,
    content_loss_failures,
    findings,
    lost_values,
)
from app.qa.policy import BLOCKING, REGISTRY, QaCheck, QaPolicy, default_policy


@pytest.fixture(autouse=True)
def registered_check():
    """This check is registered by the QA wiring, which is a separate change.

    Registering it here when it is absent keeps the tests meaningful before that
    lands and a no-op after it does -- `policy.runs()` rejects a name the
    registry does not know, by design.
    """
    added = RESOLVED_VALUE_ABSENT not in REGISTRY
    if added:
        REGISTRY[RESOLVED_VALUE_ABSENT] = QaCheck(
            RESOLVED_VALUE_ABSENT, "A field resolved to a value the document does not carry.", BLOCKING
        )
    try:
        yield
    finally:
        if added:
            REGISTRY.pop(RESOLVED_VALUE_ABSENT, None)


def _body(paragraphs=(), table_rows=()):
    d = docx.Document()
    for text in paragraphs:
        d.add_paragraph(text)
    if table_rows:
        table = d.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, text in enumerate(row):
                table.cell(r, c).text = text
    return d.element.body


def _runs_body(runs):
    """One paragraph whose text is split across several runs."""
    d = docx.Document()
    p = d.add_paragraph()
    for text in runs:
        p.add_run(text)
    return d.element.body


def _slot(paragraph_index, span_index=0):
    return {"kind": "blue_placeholder", "text": "<x>", "paragraph_index": paragraph_index, "span_index": span_index}


def _field(fid, slots=(), **over):
    return {"id": fid, "type": "string", "slots": list(slots), **over}


def _manifest(fields=(), blocks=()):
    return {"fields": list(fields), "blocks": list(blocks), "conditions": [], "delete_always": []}


def _details(body, manifest, field_values, drop_block_ids=()):
    return content_loss_failures(body, manifest, field_values, set(drop_block_ids))


# ------------------------------------------------------------ true positives

def test_a_value_the_document_lost_is_reported():
    body = _body(["Dear Anna,", "Yours sincerely,"])
    manifest = _manifest([_field("annual_salary", [_slot(7)])])

    lost = lost_values(body, manifest, {"annual_salary": ("$76,800.00", "source_record")}, set())

    assert [lv.field_id for lv in lost] == ["annual_salary"]
    assert lost[0].value == "$76,800.00"
    assert lost[0].expected_paragraphs == (7,)


def test_the_note_names_the_field_the_value_and_where_it_should_have_been():
    """A reviewer must be able to act on this without opening a debugger."""
    body = _body(["Dear Anna,"])
    manifest = _manifest([_field("annual_salary", [_slot(7)])])

    (detail,) = _details(body, manifest, {"annual_salary": ("$76,800.00", "source_record")})

    assert "'annual_salary'" in detail
    assert "$76,800.00" in detail
    assert "paragraph 7" in detail
    assert "source_record" in detail


def test_every_lost_field_is_reported_not_only_the_first():
    """The letter that prompted this lost eight fields at once. A reviewer who
    fixes one and regenerates must not discover the next one the same way."""
    body = _body(["Dear Anna,", "Yours sincerely,"])
    manifest = _manifest([
        _field("offer_sentence", [_slot(3)]),
        _field("start_date", [_slot(4)]),
        _field("hours_of_work", [_slot(5)]),
    ])
    values = {
        "offer_sentence": ("We are pleased to offer you the position", "source_record"),
        "start_date": ("14 September 2026", "source_record"),
        "hours_of_work": ("22.8 hours per week", "source_record"),
    }

    assert [lv.field_id for lv in lost_values(body, manifest, values, set())] == [
        "offer_sentence", "start_date", "hours_of_work",
    ]


def test_a_value_is_not_satisfied_by_two_adjacent_paragraphs():
    """The trailing space is the point. Concatenating the body into one string
    -- or newline-joining it, which whitespace normalisation then flattens --
    finds "Anna Smith" across this boundary, and the document does not contain
    it. The name was lost and the letter greets nobody."""
    body = _body(["Dear Anna ", "Smith, welcome to the team."])
    manifest = _manifest([_field("colleague_name", [_slot(0)])])

    assert len(_details(body, manifest, {"colleague_name": ("Anna Smith", "source_record")})) == 1


def test_a_live_slot_is_still_expected_when_a_sibling_slot_was_dropped():
    """The salary appears in the letter body and in the remuneration table. The
    table's block losing its condition does not excuse the body sentence."""
    body = _body(["Dear Anna,"])
    manifest = _manifest(
        [_field("annual_salary", [_slot(4), _slot(11)])],
        [{"id": "ft_table", "start_paragraph": 10, "end_paragraph": 14}],
    )

    (detail,) = _details(body, manifest, {"annual_salary": ("$76,800.00", "system")}, {"ft_table"})

    assert "paragraph 4" in detail
    assert "11" not in detail  # the dropped slot is not offered as a place to look


# ------------------------------------------------------------ true negatives

def test_a_value_the_document_carries_is_not_reported():
    body = _body(["Dear Anna,", "Your salary is $76,800.00 per annum."])
    manifest = _manifest([_field("annual_salary", [_slot(1)])])

    assert _details(body, manifest, {"annual_salary": ("$76,800.00", "source_record")}) == []


def test_a_value_inside_a_table_cell_counts_as_present():
    """The remuneration table is the part of an offer letter a reader checks
    first, and it is not in `Document.paragraphs`."""
    body = _body(["Dear Anna,"], table_rows=[["Base salary", "$76,800.00"]])
    manifest = _manifest([_field("annual_salary", [_slot(4)])])

    assert _details(body, manifest, {"annual_salary": ("$76,800.00", "source_record")}) == []


def test_a_value_split_across_runs_counts_as_present():
    """Word fragments a line at every formatting change, so a value can arrive
    in the document as two text nodes."""
    body = _runs_body(["Welcome ", "Anna ", "Smith", " to the team."])
    manifest = _manifest([_field("colleague_name", [_slot(0)])])

    assert _details(body, manifest, {"colleague_name": ("Anna Smith", "source_record")}) == []


def test_a_value_that_moved_to_another_paragraph_is_not_a_loss():
    """Every dropped block renumbers the paragraphs after it. Requiring the value
    at its original index would accuse most correct documents."""
    body = _body(["Your salary is $76,800.00 per annum."])
    manifest = _manifest([_field("annual_salary", [_slot(31)])])

    assert _details(body, manifest, {"annual_salary": ("$76,800.00", "source_record")}) == []


def test_a_dropped_block_excuses_the_value_it_contained():
    """A letter that chose "with recruitment" correctly omits the
    without-recruitment email address."""
    body = _body(["Dear Anna,"])
    manifest = _manifest(
        [_field("recruitment_email", [_slot(9)])],
        [{"id": "without_recruitment", "start_paragraph": 8, "end_paragraph": 12}],
    )

    values = {"recruitment_email": ("noreply@example.com", "source_record")}
    assert _details(body, manifest, values, {"without_recruitment"}) == []
    # ...and the same absence with the block kept is a loss.
    assert len(_details(body, manifest, values, set())) == 1


def test_a_dropped_inline_switch_excuses_the_value_in_its_span():
    """A span-scoped block is one arm of a switch on a single line. Dropping it
    blanks the spans and leaves the paragraph, so a paragraph-range match would
    never see it."""
    body = _body(["You are employed on a part time basis."])
    manifest = _manifest(
        [_field("full_time_hours", [_slot(0, span_index=3)])],
        [{"id": "ft_arm", "start_paragraph": 0, "start_span": 2, "end_span": 4,
          "end_paragraph": 0}],
    )

    assert _details(body, manifest, {"full_time_hours": ("38", "source_record")}, {"ft_arm"}) == []


def test_a_span_scoped_drop_does_not_excuse_the_rest_of_its_paragraph():
    """The losing arm of a switch is blanked; the static text around it is not."""
    body = _body(["You are employed on a part time basis."])
    manifest = _manifest(
        [_field("hours_of_work", [_slot(0, span_index=9)])],
        [{"id": "ft_arm", "start_paragraph": 0, "start_span": 2, "end_span": 4, "end_paragraph": 0}],
    )

    assert len(_details(body, manifest, {"hours_of_work": ("22.8", "source_record")}, {"ft_arm"})) == 1


def test_a_condition_variable_is_never_expected_in_the_document():
    """`transaction_type = "with recruitment"` decides which branch survives and
    is printed nowhere, correctly. It carries no slot."""
    body = _body(["Dear Anna,"])
    manifest = _manifest([_field("transaction_type", [])])

    assert _details(body, manifest, {"transaction_type": ("with recruitment", "source_record")}) == []


def test_a_field_that_resolved_to_nothing_is_left_to_the_missing_value_gate():
    """`required_value_missing` owns that defect. Two notes on one defect is how
    a reviewer learns to skim `qa_notes`."""
    body = _body(["Dear Anna,"])
    manifest = _manifest([_field("manager_name", [_slot(2)], required=True)])

    assert _details(body, manifest, {"manager_name": (None, "missing")}) == []
    assert _details(body, manifest, {"manager_name": ("", "source_record")}) == []
    assert _details(body, manifest, {"manager_name": ("   ", "source_record")}) == []


def test_a_sentence_removal_excuses_a_neighbours_value_on_the_same_line():
    """`REMOVE_SENTENCE` deletes the line the missing field sat in, taking any
    resolved value on it. The manifest asked for that deletion."""
    body = _body(["Dear Anna,"])
    manifest = _manifest([
        _field("bonus_scheme", [_slot(6)], on_missing="REMOVE_SENTENCE"),
        _field("bonus_percent", [_slot(6)]),
    ])
    values = {"bonus_scheme": (None, "missing"), "bonus_percent": ("12%", "source_record")}

    assert _details(body, manifest, values) == []
    # The exclusion is owed to the *missing* field. Once it resolves, the line
    # stays and the neighbour's absence is a real loss again.
    values["bonus_scheme"] = ("Short Term Incentive", "source_record")
    assert len(_details(body, manifest, values)) == 2


# ------------------------------------------------- formatting is not deletion

def test_the_formatted_value_is_compared_not_the_raw_record_value():
    """`78450` reaches the page as `$78,450.00` through `value_format`. Comparing
    against the record would accuse every currency field in the estate."""
    body = _body(["Your salary is $78,450.00 per annum."])
    manifest = _manifest([_field("annual_salary", [_slot(1)], type="currency")])

    assert _details(body, manifest, {"annual_salary": ("$78,450.00", "source_record")}) == []


def test_collapsed_whitespace_is_not_a_loss():
    """The inline-switch repair collapses runs of spaces after the fill wrote
    the value, so the string on the page can differ from the one recorded."""
    body = _body(["Located at Level 3 Sydney."])
    manifest = _manifest([_field("work_location", [_slot(0)])])

    assert _details(body, manifest, {"work_location": ("Level 3   Sydney", "source_record")}) == []


def test_a_non_breaking_space_in_a_currency_value_is_not_a_loss():
    """babel puts U+00A0 inside some locales' currency forms."""
    body = _body(["Votre salaire est de 78 450,00 EUR par an."])
    manifest = _manifest([_field("annual_salary", [_slot(0)], type="currency")])

    assert _details(body, manifest, {"annual_salary": ("78 450,00 EUR", "source_record")}) == []


def test_case_and_punctuation_are_not_normalised_away():
    """Folding either would start hiding a real substitution: a letter that says
    "PART TIME" where the record said "Part Time" has a defect worth seeing."""
    body = _body(["You are employed on a PART TIME basis."])
    manifest = _manifest([_field("employment_basis", [_slot(0)])])

    assert len(_details(body, manifest, {"employment_basis": ("Part Time", "source_record")})) == 1


# --------------------------------------------------------------------- edges

def test_a_bare_value_mapping_is_read_as_a_value_not_unpacked_as_a_pair():
    """The renderer builds `{field_id: value}` for the overflow gate. Unpacking
    a two-character value as `(value, source)` raises nothing and is wrong."""
    body = _body(["Dear Anna,"])
    manifest = _manifest([_field("country_code", [_slot(0)])])

    (detail,) = _details(body, manifest, {"country_code": "AU"})

    assert "'AU'" in detail
    assert _details(_body(["Country: AU"]), manifest, {"country_code": "AU"}) == []


def test_a_field_absent_from_the_values_mapping_is_not_reported():
    body = _body(["Dear Anna,"])
    manifest = _manifest([_field("never_resolved", [_slot(0)])])

    assert _details(body, manifest, {}) == []


def test_a_block_id_that_is_dropped_but_undeclared_excuses_nothing():
    """A condition naming a block the manifest does not define is
    `manifest_validation`'s `condition_targets_unknown_block`. It must not
    silently excuse a value here."""
    body = _body(["Dear Anna,"])
    manifest = _manifest([_field("annual_salary", [_slot(4)])])

    values = {"annual_salary": ("$76,800.00", "source_record")}
    assert len(_details(body, manifest, values, {"a_block_that_does_not_exist"})) == 1


def test_a_slot_with_no_recorded_paragraph_is_treated_as_live():
    """An absence can only be excused when it is provably deliberate."""
    body = _body(["Dear Anna,"])
    manifest = _manifest(
        [_field("annual_salary", [{"kind": "blue_placeholder", "text": "<Salary>"}])],
        [{"id": "ft_table", "start_paragraph": 0, "end_paragraph": 40}],
    )

    (detail,) = _details(body, manifest, {"annual_salary": ("$76,800.00", "source_record")}, {"ft_table"})
    assert "does not record" in detail


def test_a_mergefield_slot_is_dropped_by_paragraph_not_by_span():
    """A mergefield slot records no span index, and the renderer likewise only
    skips it for a paragraph-scoped drop."""
    body = _body(["Dear Anna,"])
    manifest = _manifest(
        [_field("ft_salary", [{"kind": "mergefield", "code": "LAB__FT_SALARY__38_HR_", "paragraph_index": 12}])],
        [
            {"id": "inline_arm", "start_paragraph": 12, "end_paragraph": 12, "start_span": 0, "end_span": 5},
            {"id": "ft_table", "start_paragraph": 10, "end_paragraph": 14},
        ],
    )
    values = {"ft_salary": ("$76,800.00", "source_record")}

    assert len(_details(body, manifest, values, {"inline_arm"})) == 1
    assert _details(body, manifest, values, {"ft_table"}) == []


def test_an_empty_manifest_produces_nothing():
    assert _details(_body(["Dear Anna,"]), _manifest(), {}) == []


# -------------------------------------------------------------------- policy

def test_findings_carry_the_check_name_the_registry_knows():
    body = _body(["Dear Anna,"])
    manifest = _manifest([_field("annual_salary", [_slot(4)])])

    out = findings(body, manifest, {"annual_salary": ("$76,800.00", "source_record")}, set(), default_policy())

    assert [f.check for f in out] == [RESOLVED_VALUE_ABSENT]


def test_the_check_does_no_work_when_the_policy_does_not_run_it():
    body = _body(["Dear Anna,"])
    manifest = _manifest([_field("annual_salary", [_slot(4)])])
    off = QaPolicy(severities={name: BLOCKING for name in REGISTRY}, enabled=frozenset())

    assert findings(body, manifest, {"annual_salary": ("$76,800.00", "source_record")}, set(), off) == []
