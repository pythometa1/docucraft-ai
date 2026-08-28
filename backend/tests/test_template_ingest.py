"""Template ingestion must produce a mappable template on *both* upload paths.

Bulk onboarding is how an estate of thousands of templates actually arrives, so
a template that lands without its section tree is not a corner case -- it is the
common case, and it is silent: the mapping wizard just shows an empty list.
"""

import io

import docx
import pytest


def _offer_letter_bytes() -> bytes:
    d = docx.Document()
    d.add_heading("Position", level=1)
    d.add_paragraph("You are offered the role of {position_title}.")
    d.add_heading("Remuneration", level=1)
    d.add_paragraph("Your salary will be {salary}.")
    d.add_heading("Conditions", level=2)
    d.add_paragraph("Standard conditions apply.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def project(two_orgs):
    token_a, project_a, _tb, _pb = two_orgs
    return token_a, project_a


def test_single_upload_parses_sections(app_client, project):
    token, project_id = project
    res = app_client.post(
        f"/api/v1/projects/{project_id}/templates",
        headers=_auth(token),
        files={"file": ("offer.docx", _offer_letter_bytes(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["status"] == "ready"
    assert body["section_count"] == 3
    assert set(body["jinja_vars"]) == {"{position_title}", "{salary}"}


def test_bulk_onboard_also_parses_sections(app_client, project):
    """Previously this stored the blob and stopped, leaving section_count at 0."""
    token, project_id = project
    res = app_client.post(
        f"/api/v1/projects/{project_id}/templates:bulk-onboard",
        headers=_auth(token),
        files=[
            ("files", ("bulk_a.docx", _offer_letter_bytes(),
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("bulk_b.docx", _offer_letter_bytes(),
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ],
        data={"auto_compile": "false"},
    )
    assert res.status_code == 201, res.text
    assert res.json()["uploaded_count"] == 2

    listing = app_client.get(f"/api/v1/projects/{project_id}/templates", headers=_auth(token)).json()
    bulk = [t for t in listing["items"] if t["name"].startswith("bulk_")]
    assert len(bulk) == 2
    for template in bulk:
        assert template["status"] == "ready"
        assert template["section_count"] == 3, "bulk-onboarded template has no section tree"

        sections = app_client.get(
            f"/api/v1/template-versions/{template['current_version_id']}/sections",
            headers=_auth(token),
        )
        assert sections.status_code == 200
        titles = [node["title"] for node in sections.json()["items"]]
        assert titles == ["Position", "Remuneration"]  # tree: Conditions nests under Remuneration


def test_unparseable_template_is_kept_with_an_error(app_client, project):
    """A bad file must not vanish or abort the request -- the user still needs to
    see what they uploaded and why it failed."""
    token, project_id = project
    res = app_client.post(
        f"/api/v1/projects/{project_id}/templates",
        headers=_auth(token),
        files={"file": ("broken.docx", b"this is not a docx at all",
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["status"] == "failed"
    assert "Could not parse template" in body["parse_error"]


def test_source_listing_exposes_its_current_version(app_client, project):
    """Callers address a source by *version*, not by file. Without this field the
    mapping wizard sent an empty source_version_ids and generation silently fell
    back to every chunk in the project."""
    token, project_id = project
    csv = b"full_name,role\nJordan Ellis,Analyst\n"
    created = app_client.post(
        f"/api/v1/projects/{project_id}/sources",
        headers=_auth(token),
        files={"file": ("people.csv", csv, "text/csv")},
    )
    assert created.status_code == 201, created.text
    assert created.json()["current_version_id"]

    listing = app_client.get(f"/api/v1/projects/{project_id}/sources", headers=_auth(token)).json()
    source = next(s for s in listing["items"] if s["name"] == "people.csv")
    assert source["current_version_id"], "listing must expose current_version_id"

    records = app_client.get(
        f"/api/v1/source-versions/{source['current_version_id']}/records",
        headers=_auth(token),
    )
    assert records.status_code == 200
    assert records.json()["columns"] == ["full_name", "role"]
