"""End-to-end: what the fill engine writes into the document when a value is
missing, and whether it admits to it.

`test_missing_and_blocking.py` covers the policy decisions in isolation. These
run a real .docx through `fill_template`, because the failure being guarded
against is not a wrong return value -- it is a finished-looking letter with a
blank where the salary should be.
"""

import os
import tempfile

import docx
import pytest

from app.templates.parsers.docx_prescan import prescan
from app.generation.docx_renderer import fill_template


def _template(paragraphs) -> str:
    d = docx.Document()
    for text in paragraphs:
        d.add_paragraph(text)
    path = os.path.join(tempfile.mkdtemp(), "t.docx")
    d.save(path)
    return path


def _text(path) -> str:
    return "\n".join(p.text for p in docx.Document(path).paragraphs)


def _field(fid, scan, paragraph_index=0, **over):
    """A field bound to the `<fid>` bracket the prescan found."""
    slots = [
        {"paragraph_index": s.paragraph_index, "span_index": s.span_index, "text": f"<{fid}>", "field_id": fid}
        for s in scan.spans
        if f"<{fid}>" in (s.text or "")
    ]
    return {"id": fid, "type": "text", "slots": slots, **over}


def _manifest(fields=(), conditions=(), blocks=()):
    return {
        "fields": list(fields), "conditions": list(conditions),
        "blocks": list(blocks), "delete_always": [],
    }


# --------------------------------------------------------------------- BLOCK

def test_missing_required_field_blocks_the_document(tmp_path):
    """The most expensive bug class: the letter reads fine and is wrong."""
    path = _template(["Your manager is <manager_name>."])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")

    result = fill_template(path, out, _manifest([_field("manager_name", scan, required=True)]), {})

    assert result.qa_passed is False
    assert any("manager_name" in n for n in result.qa_notes), result.qa_notes
    lineage = {e["field_id"]: e for e in result.field_lineage}
    assert lineage["manager_name"]["source"] == "missing"
    assert lineage["manager_name"]["on_missing"] == "BLOCK"
    assert lineage["manager_name"]["blocking"] is True


def test_present_required_field_passes(tmp_path):
    path = _template(["Your manager is <manager_name>."])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")

    result = fill_template(
        path, out, _manifest([_field("manager_name", scan, required=True)]), {"manager_name": "Dana Ruiz"}
    )

    assert result.qa_passed is True
    assert "Dana Ruiz" in _text(out)


# --------------------------------------------------------------------- BLANK

def test_optional_field_renders_blank_without_blocking(tmp_path):
    """An optional field that is genuinely absent is not a defect."""
    path = _template(["Second line: <address_line2>."])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")

    result = fill_template(path, out, _manifest([_field("address_line2", scan)]), {})

    assert result.qa_passed is True
    assert _text(out).strip() == "Second line: ."
    assert result.field_lineage[0]["on_missing"] == "BLANK"


# ------------------------------------------------------------------- DEFAULT

def test_default_policy_writes_the_declared_default(tmp_path):
    path = _template(["Reporting to <manager_name>."])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")

    field = _field("manager_name", scan, required=True, on_missing="DEFAULT", default="your line manager")
    result = fill_template(path, out, _manifest([field]), {})

    assert result.qa_passed is True
    assert "your line manager" in _text(out)
    assert result.field_lineage[0]["source"] == "fallback"


def test_default_policy_with_no_default_declared_blocks(tmp_path):
    """A manifest that says DEFAULT and declares none is incoherent; guessing
    "" here is the silent blank in a different costume."""
    path = _template(["Reporting to <manager_name>."])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")

    field = _field("manager_name", scan, on_missing="DEFAULT")
    result = fill_template(path, out, _manifest([field]), {})

    assert result.qa_passed is False
    assert any("no default" in n for n in result.qa_notes), result.qa_notes


# ----------------------------------------------------------- REMOVE_SENTENCE

def test_remove_sentence_drops_only_its_own_sentence(tmp_path):
    path = _template(["You start on 1 March. You will report to <manager_name>. Parking is provided."])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")

    field = _field("manager_name", scan, on_missing="REMOVE_SENTENCE")
    result = fill_template(path, out, _manifest([field]), {})

    text = _text(out)
    assert "You start on 1 March." in text
    assert "Parking is provided." in text
    assert "report to" not in text
    assert result.qa_passed is True


def test_remove_sentence_leaves_the_document_when_the_value_is_present(tmp_path):
    path = _template(["You start on 1 March. You will report to <manager_name>. Parking is provided."])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")

    field = _field("manager_name", scan, on_missing="REMOVE_SENTENCE")
    result = fill_template(path, out, _manifest([field]), {"manager_name": "Dana Ruiz"})

    assert "report to Dana Ruiz" in _text(out)
    assert result.qa_passed is True


def test_no_removal_marker_survives_into_the_document(tmp_path):
    path = _template(["Alone: <manager_name>"])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")

    field = _field("manager_name", scan, on_missing="REMOVE_SENTENCE")
    fill_template(path, out, _manifest([field]), {})

    assert "\x00" not in _text(out)
    assert "RS0" not in _text(out)


# ----------------------------------------------------------------- conditions

def test_undecidable_condition_keeps_its_block_and_blocks_the_document(tmp_path):
    """The section the employee was entitled to see stays in; the letter does
    not go out until someone supplies the field."""
    path = _template(["Intro.", "Temporary assignment terms.", "Closing."])
    out = str(tmp_path / "out.docx")
    manifest = _manifest(
        conditions=[{"id": "c1", "expression": "assignment_type == 'TEMPORARY'", "keeps_blocks": ["b1"]}],
        blocks=[{"id": "b1", "start_paragraph": 1, "end_paragraph": 1, "boundary_method": "marker"}],
    )

    result = fill_template(path, out, manifest, {})

    assert "Temporary assignment terms." in _text(out)
    assert result.qa_passed is False
    assert any("assignment_type" in n for n in result.qa_notes), result.qa_notes
    assert result.condition_lineage[0]["result"] is None
    assert result.condition_lineage[0]["reason"] == "missing_input"


def test_a_decided_false_condition_still_drops_its_block(tmp_path):
    """The change must not make every condition undecidable."""
    path = _template(["Intro.", "Temporary assignment terms.", "Closing."])
    out = str(tmp_path / "out.docx")
    manifest = _manifest(
        conditions=[{"id": "c1", "expression": "assignment_type == 'TEMPORARY'", "keeps_blocks": ["b1"]}],
        blocks=[{"id": "b1", "start_paragraph": 1, "end_paragraph": 1, "boundary_method": "marker"}],
    )

    result = fill_template(path, out, manifest, {"assignment_type": "PERMANENT"})

    assert "Temporary assignment terms." not in _text(out)
    assert result.qa_passed is True
    assert result.condition_lineage[0]["result"] is False


def test_supplied_verdicts_are_still_honoured(tmp_path):
    """The resolution orchestrator sees context this layer cannot; a verdict it
    hands down is a decision, not a guess."""
    path = _template(["Intro.", "Temporary assignment terms.", "Closing."])
    out = str(tmp_path / "out.docx")
    manifest = _manifest(
        conditions=[{"id": "c1", "expression": "needs_judgement == 1", "keeps_blocks": ["b1"]}],
        blocks=[{"id": "b1", "start_paragraph": 1, "end_paragraph": 1, "boundary_method": "marker"}],
    )

    result = fill_template(path, out, manifest, {}, condition_verdicts={"c1": True})

    assert "Temporary assignment terms." in _text(out)
    assert result.qa_passed is True
    assert result.condition_lineage[0]["reason"] == "supplied"


# ------------------------------------------------------------------ QA is AND

def test_a_scaffolding_pass_cannot_clear_an_earlier_block(tmp_path):
    """`qa_passed` used to be *assigned* by the final scaffolding sweep, which
    discarded every block recorded during the fill."""
    path = _template(["Your manager is <manager_name>."])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")

    # Leaves no bracket, no mergefield, no control token behind -- so the final
    # sweep finds nothing and would previously have reported a clean document.
    result = fill_template(path, out, _manifest([_field("manager_name", scan, required=True)]), {})

    assert "<manager_name>" not in _text(out)
    assert result.qa_passed is False


# ---------------------------------------------------------------- mergefields
#
# A MERGEFIELD resolves down a separate code path from a bracket placeholder --
# it replaces a whole complex-field run sequence rather than editing a span --
# so every policy has to be proven twice.

def _mergefield_template(code: str, sentences: str) -> str:
    """A .docx whose paragraph carries a real `{ MERGEFIELD <code> }`."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    d = docx.Document()
    p = d.add_paragraph(sentences)

    def _run(*children):
        r = OxmlElement("w:r")
        for c in children:
            r.append(c)
        p._p.append(r)

    begin, sep, end = (OxmlElement("w:fldChar") for _ in range(3))
    begin.set(qn("w:fldCharType"), "begin")
    sep.set(qn("w:fldCharType"), "separate")
    end.set(qn("w:fldCharType"), "end")
    instr = OxmlElement("w:instrText")
    instr.text = f" MERGEFIELD {code} "
    _run(begin)
    _run(instr)
    _run(sep)
    _run(end)

    path = os.path.join(tempfile.mkdtemp(), "mf.docx")
    d.save(path)
    return path


def _mergefield_field(fid, code, **over):
    return {
        "id": fid, "type": "text",
        "slots": [{"kind": "mergefield", "paragraph_index": 0, "span_index": 0, "code": code}],
        **over,
    }


def test_mergefield_resolves_from_its_code(tmp_path):
    path = _mergefield_template("manager_name", "Reporting to ")
    out = str(tmp_path / "out.docx")

    result = fill_template(
        path, out, _manifest([_mergefield_field("manager_name", "manager_name")]), {"manager_name": "Dana Ruiz"}
    )

    assert "Dana Ruiz" in _text(out)
    assert result.field_lineage[0]["source"] == "source_record"
    assert result.qa_passed is True


def test_mergefield_with_a_null_cell_is_missing_not_blank(tmp_path):
    """`code in record` used to be enough, so a null cell formatted as a value
    and the field was reported resolved."""
    path = _mergefield_template("manager_name", "Reporting to ")
    out = str(tmp_path / "out.docx")

    field = _mergefield_field("manager_name", "manager_name", required=True)
    result = fill_template(path, out, _manifest([field]), {"manager_name": None})

    assert result.qa_passed is False
    assert result.field_lineage[0]["source"] == "missing"
    assert result.field_lineage[0]["blocking"] is True


def test_mergefield_remove_sentence(tmp_path):
    path = _mergefield_template("manager_name", "You start on 1 March. You will report to ")
    out = str(tmp_path / "out.docx")

    field = _mergefield_field("manager_name", "manager_name", on_missing="REMOVE_SENTENCE")
    result = fill_template(path, out, _manifest([field]), {})

    text = _text(out)
    assert "You start on 1 March." in text
    assert "report to" not in text
    assert "" not in text
    assert result.qa_passed is True


# ------------------------------------------------------------ helper edges

def test_sentence_bounds_falls_back_to_the_whole_string():
    from app.generation.docx_renderer import _sentence_bounds

    # An offset past the end has no sentence; the whole string is the safe answer.
    assert _sentence_bounds("No terminator here", 999) == (0, 18)


def test_sentence_removal_reports_a_marker_it_cannot_find(tmp_path):
    from app.generation.docx_renderer import _apply_sentence_removal, _strip_markers

    path = _template(["Nothing to remove here."])
    d = docx.Document(path)
    p_el = d.paragraphs[0]._p

    assert _apply_sentence_removal(p_el, "RS7") is False
    _strip_markers(p_el)  # no marker present: a no-op, not a crash
    assert d.paragraphs[0].text == "Nothing to remove here."


def test_stray_marker_is_stripped_rather_than_rendered(tmp_path):
    from app.generation.docx_renderer import _strip_markers

    path = _template(["Left over RS3 mid-sentence."])
    d = docx.Document(path)
    _strip_markers(d.paragraphs[0]._p)
    assert d.paragraphs[0].text == "Left over  mid-sentence."
