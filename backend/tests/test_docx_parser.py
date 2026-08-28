"""Heading detection across languages.

A missed heading is not a small error: the template parses to zero sections and
becomes invisible to the section-mapping engine, so the whole document silently
cannot be mapped.
"""

import docx
import pytest
from docx.oxml.ns import qn

from app.templates.parsers.docx_parser import heading_level, parse_docx_template

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _doc_with_style_name(tmp_path, style_name: str, filename: str):
    """Build a document whose Heading 1 style carries `style_name` as its display
    name -- what a template with author-defined localized styles looks like."""
    d = docx.Document()
    d.add_heading("Section One", level=1)
    d.add_paragraph("Body text under the first section.")
    d.add_heading("Section Two", level=1)
    d.add_paragraph("Body text under the second section.")
    for style in d.styles.element.findall(W + "style"):
        if style.get(W + "styleId") == "Heading1":
            style.find(W + "name").set(W + "val", style_name)
    path = tmp_path / filename
    d.save(str(path))
    return str(path)


@pytest.mark.parametrize("style_name,label", [
    ("heading 1", "english-builtin"),
    ("Überschrift 1", "german"),
    ("Titre 1", "french"),
    ("Título 1", "spanish"),
    ("見出し 1", "japanese"),
    ("Заголовок 1", "russian"),
])
def test_headings_are_found_regardless_of_style_language(tmp_path, style_name, label):
    path = _doc_with_style_name(tmp_path, style_name, f"{label}.docx")
    parsed = parse_docx_template(path)
    assert len(parsed.sections) == 2, f"{label}: found {len(parsed.sections)} sections"
    assert [s.title for s in parsed.sections] == ["Section One", "Section Two"]


def test_outline_level_is_used_when_the_style_name_is_meaningless(tmp_path):
    """A custom style with an opaque name still declares its outline level, which
    is what Word itself reads to build a table of contents."""
    d = docx.Document()
    heading = d.add_paragraph("Custom Styled Heading")
    p_pr = heading._p.get_or_add_pPr()
    lvl = p_pr.makeelement(qn("w:outlineLvl"), {qn("w:val"): "0"})
    p_pr.append(lvl)
    d.add_paragraph("Body copy.")
    path = tmp_path / "outline.docx"
    d.save(str(path))

    assert heading_level(docx.Document(str(path)).paragraphs[0]) == 1


def test_body_paragraphs_are_not_mistaken_for_headings(tmp_path):
    d = docx.Document()
    d.add_paragraph("Just ordinary body text.")
    path = tmp_path / "plain.docx"
    d.save(str(path))

    assert heading_level(docx.Document(str(path)).paragraphs[0]) is None
    assert parse_docx_template(str(path)).sections == []


def test_jinja_variables_are_still_detected(tmp_path):
    """Regression guard: the mapping wizard keys its field dropdown off these."""
    d = docx.Document()
    d.add_heading("Offer", level=1)
    d.add_paragraph("Dear {employee_name}, welcome to {department}.")
    path = tmp_path / "jinja.docx"
    d.save(str(path))

    parsed = parse_docx_template(str(path))
    assert set(parsed.jinja_vars) == {"{employee_name}", "{department}"}
    assert parsed.kind == "mixed"
