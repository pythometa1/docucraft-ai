"""The workbook a template asks for, and the failure it exists to prevent.

The bug that motivated this is on record. A Pfizer transfer letter compiled to
three condition variables -- `work_schedule`, `employment_type`,
`transaction_type` -- and the customer's HR export named a column
`DRV_Employment_Type` while filling it with `Full Time` / `Part Time`. The
binder matched on the name, so `employment_type` took that column, the real
contract column went unbound, and `work_schedule` was assigned a salary. One
letter lost its work-schedule paragraph; a part-time colleague received a letter
describing full-time employment, and it passed every QA gate.

None of that is reachable through a workbook generated here: the columns carry
the field ids, and a condition column will not accept a value no branch tests
for. These tests hold both halves of that claim.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

from app.generation.source_ingestion import extract_records
from app.generation.source_resolver import suggest_bindings
from app.generation.source_template import (
    DATA_SHEET,
    GUIDE_SHEET,
    MAX_INLINE_LIST,
    build_workbook,
    filename_for,
    permitted_values,
    plan_columns,
)

MANIFEST = {
    "fields": [
        {"id": "colleague_first_name", "type": "string", "required": True,
         "slots": [{"kind": "blue_placeholder", "text": "<Colleague First Name>"}]},
        {"id": "start_date", "type": "date",
         "slots": [{"kind": "blue_placeholder", "text": "<Start Date>"}]},
        {"id": "annual_salary", "type": "currency",
         "slots": [{"kind": "mergefield", "code": "LAB__FT_SALARY__38_HR_"}]},
    ],
    "conditions": [
        {"id": "c1", "expression": "work_schedule == 'Full time'", "keeps_blocks": ["b1"]},
        {"id": "c2", "expression": "work_schedule == 'Part time'", "keeps_blocks": ["b2"]},
        {"id": "c3", "expression": "employment_type == 'Fixed term'", "keeps_blocks": ["b3"]},
        {"id": "c4", "expression": "employment_type == 'Regular'", "keeps_blocks": ["b4"]},
    ],
    "blocks": [],
}


def _load(manifest=MANIFEST):
    return openpyxl.load_workbook(io.BytesIO(build_workbook(manifest, template_name="offer.docx")))


# ------------------------------------------------------------ what it contains

def test_condition_variables_become_columns_even_with_no_placeholder():
    """`work_schedule` appears in no placeholder anywhere in the document. A
    workbook built from `manifest["fields"]` would omit it, and its absence is
    what silently deletes conditional sections."""
    ids = {c.field_id for c in plan_columns(MANIFEST)}
    assert {"work_schedule", "employment_type"} <= ids
    assert {"colleague_first_name", "start_date", "annual_salary"} <= ids


def test_permitted_values_are_read_from_the_expressions():
    assert permitted_values(MANIFEST) == {
        "work_schedule": ["Full time", "Part time"],
        "employment_type": ["Fixed term", "Regular"],
    }


def test_a_non_equality_condition_constrains_nothing():
    """`salary > 100000` admits no finite set, so it must not produce a dropdown
    offering two arbitrary values."""
    manifest = {**MANIFEST, "conditions": [{"id": "x", "expression": "annual_salary > 100000"}]}
    assert permitted_values(manifest) == {}


def test_condition_columns_come_first():
    """They decide which sections exist, so they are answered before someone
    types prose into a paragraph a later answer removes."""
    columns = plan_columns(MANIFEST)
    leading = [c.field_id for c in columns[:2]]
    assert set(leading) == {"employment_type", "work_schedule"}


def test_the_data_sheet_is_first():
    """`extract_records` reads `worksheets[0]` when no sheet is named. A guide
    sheet in front of it would make the default read return documentation."""
    workbook = _load()
    assert workbook.sheetnames[0] == DATA_SHEET
    assert GUIDE_SHEET in workbook.sheetnames


def test_headers_are_field_ids_not_prettified_labels():
    """The header is what the binder matches on. A humanised header would put
    this file back through the fuzzy matching it exists to avoid."""
    sheet = _load()[DATA_SHEET]
    headers = {c.value for c in sheet[1]}
    assert "colleague_first_name" in headers
    assert "Colleague first name" not in headers


# ------------------------------------------------------------ the dropdowns

def test_condition_columns_carry_exactly_the_branch_values():
    sheet = _load()[DATA_SHEET]
    lists = {dv.formula1 for dv in sheet.data_validations.dataValidation}
    assert '"Full time,Part time"' in lists
    assert '"Fixed term,Regular"' in lists


def test_only_condition_columns_are_constrained():
    """A dropdown on a free-text column would stop a person entering their own
    colleague's name."""
    sheet = _load()[DATA_SHEET]
    constrained = set()
    for dv in sheet.data_validations.dataValidation:
        for ref in str(dv.sqref).split():
            constrained.add(ref.split("2:")[0])
    headers = {c.column_letter: c.value for c in sheet[1]}
    for letter in constrained:
        assert headers[letter] in ("work_schedule", "employment_type")


def test_a_branch_value_containing_a_comma_gets_no_dropdown():
    """Excel's inline list is comma-separated with no escape, so such a value
    would split into two wrong options. A dropdown offering strings no branch
    tests for is worse than none: it looks authoritative."""
    manifest = {**MANIFEST, "conditions": [
        {"id": "c", "expression": "region == 'Melbourne, VIC'"},
        {"id": "d", "expression": "region == 'Sydney'"},
    ]}
    sheet = openpyxl.load_workbook(io.BytesIO(build_workbook(manifest)))[DATA_SHEET]
    assert list(sheet.data_validations.dataValidation) == []


def test_an_over_long_value_list_gets_no_dropdown():
    values = [f"Value number {i:03d} spelled out at length" for i in range(20)]
    assert len(",".join(values)) > MAX_INLINE_LIST
    manifest = {**MANIFEST, "conditions": [
        {"id": f"c{i}", "expression": f"category == '{v}'"} for i, v in enumerate(values)
    ]}
    sheet = openpyxl.load_workbook(io.BytesIO(build_workbook(manifest)))[DATA_SHEET]
    assert list(sheet.data_validations.dataValidation) == []


# ------------------------------------------------------------ the round trip

def _fill(manifest, rows: list[dict]) -> str:
    """Generate the workbook, fill it in as a person would, save it."""
    import tempfile, os

    workbook = openpyxl.load_workbook(io.BytesIO(build_workbook(manifest)))
    sheet = workbook[DATA_SHEET]
    headers = [c.value for c in sheet[1]]
    for offset, row in enumerate(rows):
        for index, header in enumerate(headers, start=1):
            sheet.cell(row=2 + offset, column=index, value=row.get(header, f"{header} value"))
    path = os.path.join(tempfile.mkdtemp(), "filled.xlsx")
    workbook.save(path)
    return path


def test_a_filled_workbook_binds_every_field_to_its_own_column():
    """The whole point. Every target pairs with the column of its own name, so
    no reviewer has to work out which column meant what."""
    path = _fill(MANIFEST, [
        {"work_schedule": "Full time", "employment_type": "Fixed term"},
        {"work_schedule": "Part time", "employment_type": "Regular"},
    ])
    columns, records = extract_records(path, "xlsx", DATA_SHEET)
    assert len(records) == 2

    plan = suggest_bindings(MANIFEST, columns, records=records)
    mismatched = [(s.field_id, s.column) for s in plan.suggestions if s.column != s.field_id]
    assert mismatched == [], f"a field was bound to someone else's column: {mismatched}"
    assert plan.unmatched_fields == []


def test_a_filled_workbook_leaves_no_condition_value_without_a_branch():
    """`unmatched_condition_values` is the list the reviewer had to repair by
    hand, and each entry is a letter that would generate with a section missing.
    Values chosen from the dropdown cannot produce one."""
    path = _fill(MANIFEST, [
        {"work_schedule": "Full time", "employment_type": "Fixed term"},
        {"work_schedule": "Part time", "employment_type": "Regular"},
    ])
    columns, records = extract_records(path, "xlsx", DATA_SHEET)
    plan = suggest_bindings(MANIFEST, columns, records=records)
    assert [(u.field_id, u.observed_value) for u in plan.unmatched_condition_values] == []


def test_the_pfizer_failure_is_not_reproducible_through_this_workbook():
    """The regression, stated as the original defect.

    The customer's export held work-schedule values under a column named
    `DRV_Employment_Type`. Here the two are separate columns carrying their own
    names, so `employment_type` cannot absorb `Full time` and `work_schedule`
    cannot be left to take whatever column is still free.
    """
    path = _fill(MANIFEST, [{"work_schedule": "Part time", "employment_type": "Regular"}])
    columns, records = extract_records(path, "xlsx", DATA_SHEET)
    plan = suggest_bindings(MANIFEST, columns, records=records)
    bound = {s.field_id: s.column for s in plan.suggestions}
    assert bound["work_schedule"] == "work_schedule"
    assert bound["employment_type"] == "employment_type"
    assert records[0]["work_schedule"] == "Part time"
    assert records[0]["employment_type"] == "Regular"


# ------------------------------------------------------------ dates as text

def test_date_columns_are_held_as_text():
    """`source_ingestion._clean_cell` does `str(value)`, so a real Excel date
    arrives as "2024-07-01 00:00:00" and that is what would be printed in the
    letter. Holding the column as text keeps what the person typed."""
    sheet = _load()[DATA_SHEET]
    header = {c.value: c.column for c in sheet[1]}
    assert sheet.cell(row=2, column=header["start_date"]).number_format == "@"
    assert sheet.cell(row=2, column=header["colleague_first_name"]).number_format != "@"


# ------------------------------------------------------------ naming

@pytest.mark.parametrize("given,expected", [
    ("offer.docx", "offer_source_template.xlsx"),
    ("HCM.AU.CT.042 Salary.DOCX", "HCM.AU.CT.042_Salary_source_template.xlsx"),
    ("", "template_source_template.xlsx"),
    ("../../etc/passwd", ".._.._etc_passwd_source_template.xlsx"),
])
def test_filename_is_derived_safely(given, expected):
    """The name reaches a Content-Disposition header, so a path separator in a
    template name must not survive into it."""
    assert filename_for(given) == expected
