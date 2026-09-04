"""The referee: what the document says is wrong with the model's reading.

These are the loop's only objective signal, so each one is asserted twice --
that it fires on a manifest built to trip it, and that it stays silent on a clean
one. A check that cannot fire and a check that never has anything to fire on look
identical from the outside, and only the negative half tells them apart.
"""

from __future__ import annotations

import pathlib

import docx
import pytest

from app.compiler import assertions as A
from app.templates.parsers.docx_prescan import prescan


@pytest.fixture
def template(tmp_path):
    def _make(paragraphs):
        document = docx.Document()
        for text in paragraphs:
            document.add_paragraph(text)
        path = tmp_path / "t.docx"
        document.save(str(path))
        return prescan(str(path))
    return _make


def _slot(index, text):
    return {"kind": "text_match", "text": text, "paragraph_index": index, "span_index": 0}


def _manifest(fields=None, conditions=None, blocks=None, delete_always=None):
    return {
        "fields": fields or [], "conditions": conditions or [],
        "blocks": blocks or [], "delete_always": delete_always or [],
    }


PARAS = ["Dear <Name>,", "Your role is <Role>.", "Yours sincerely,"]
CLEAN = _manifest(fields=[
    {"id": "name", "slots": [_slot(0, "<Name>")]},
    {"id": "role", "slots": [_slot(1, "<Role>")]},
])


def test_a_correct_reading_raises_nothing(template):
    assert A.collect(template(PARAS), CLEAN) == []


def test_an_unclaimed_placeholder_is_reported_with_its_paragraph(template):
    """This is the failure that reaches the reader as visible scaffolding."""
    partial = _manifest(fields=[CLEAN["fields"][0]])
    faults = A.collect(template(PARAS), partial)

    assert [f.check for f in faults] == [A.UNCOVERED_PLACEHOLDER]
    assert faults[0].paragraph_index == 1
    assert "<Role>" in faults[0].detail


def test_a_ruled_line_is_not_treated_as_a_placeholder(template):
    """`<______>` is somewhere to write, not a field. Demanding the writer claim
    it makes the loop unsatisfiable: it correctly declines, the check fires
    again, and a perfectly-read template parks as failed."""
    scan = template(["Signed <__________________> on <Date>."])
    manifest = _manifest(fields=[{"id": "date", "slots": [_slot(0, "<Date>")]}])
    assert A.collect(scan, manifest) == []


def test_a_placeholder_claimed_in_a_neighbouring_paragraph_still_counts(template):
    """`_locate` searches a window around the index the model gave, so the slot
    may legitimately be recorded next door."""
    scan = template(PARAS)
    manifest = _manifest(fields=[
        CLEAN["fields"][0],
        {"id": "role", "slots": [_slot(0, "<Role>")]},
    ])
    assert [f.check for f in A.collect(scan, manifest)] == []


@pytest.mark.parametrize("expression", [
    "For happy colleagues:",          # prose, not a rule
    "",                                # nothing at all
    "'Permanent'",                     # a value with no field to test
])
def test_an_expression_that_cannot_be_executed_is_reported(template, expression):
    manifest = {**CLEAN, "conditions": [{"id": "c1", "expression": expression}]}
    faults = [f for f in A.collect(template(PARAS), manifest) if f.check == A.UNEXECUTABLE_CONDITION]
    assert [f.object_id for f in faults] == ["c1"]


def test_an_executable_expression_is_left_alone(template):
    manifest = {**CLEAN, "conditions": [{"id": "c1", "expression": "colleague_type == 'Fixed term'"}]}
    assert [f for f in A.collect(template(PARAS), manifest) if f.check == A.UNEXECUTABLE_CONDITION] == []


def test_a_field_with_no_occurrence_is_reported(template):
    manifest = _manifest(fields=[*CLEAN["fields"], {"id": "ghost", "slots": []}])
    faults = [f for f in A.collect(template(PARAS), manifest) if f.check == A.FIELD_WITHOUT_SLOT]
    assert [f.object_id for f in faults] == ["ghost"]


def test_a_field_only_in_deleted_paragraphs_is_reported(template):
    """The compiler declared a field, gave it a spreadsheet column, and then
    guaranteed the paragraph holding it would never survive."""
    manifest = {**CLEAN, "delete_always": [{"paragraph_index": 1, "span_index": None}]}
    faults = [f for f in A.collect(template(PARAS), manifest) if f.check == A.ORPHANED_FIELD]
    assert [f.object_id for f in faults] == ["role"]


def test_test_fill_failures_arrive_as_assertions(template):
    faults = A.collect(template(PARAS), CLEAN, test_fill_notes=["Leftover placeholder brackets: ['Yes/No']"])
    assert [f.check for f in faults] == [A.TEST_FILL_FAILURE]
    assert "Yes/No" in faults[0].detail


def test_rendering_is_readable_and_bounded(template):
    assert "No mechanical faults" in A.render([])
    many = [A.Assertion(A.UNCOVERED_PLACEHOLDER, f"fault {i}") for i in range(100)]
    rendered = A.render(many, limit=10)
    assert rendered.count("\n") == 10
    assert "+90 more" in rendered


# ---- MERGEFIELDs need a real Word file; python-docx cannot author one --------

FIXTURES = pathlib.Path(__file__).resolve().parent / "fixtures" / "templates"
HOSPIRA = FIXTURES / "hospira_offer.docx"


@pytest.mark.skipif(not HOSPIRA.exists(), reason="hospira_offer.docx fixture is missing")
def test_an_unclaimed_merge_field_is_reported():
    """A MERGEFIELD renders as its cached result -- usually blank -- so it is
    invisible in the paragraph text. Nothing but this check can see it."""
    scan = prescan(str(HOSPIRA))
    assert scan.mergefields, "fixture no longer carries merge fields"

    faults, warnings = A.uncovered_mergefields(scan, _manifest())
    assert faults, "every merge field is unclaimed by an empty manifest"
    assert all(f.check == A.UNCOVERED_MERGEFIELD for f in faults)
    assert warnings == [], "nothing is claimed, so nothing collides"


@pytest.mark.skipif(not HOSPIRA.exists(), reason="hospira_offer.docx fixture is missing")
def test_a_claimed_merge_field_is_not_reported():
    """A claim is a `mergefield` slot carrying the code -- matched on letters and
    digits alone, because a model returns the code with the guillemets, without
    them, with Word's quotes, or re-cased."""
    scan = prescan(str(HOSPIRA))
    code = scan.mergefields[0].code.strip().strip('"')
    manifest = _manifest(fields=[{
        "id": "salary",
        "slots": [{"kind": "mergefield", "code": code, "paragraph_index": scan.mergefields[0].paragraph_index}],
    }])

    faults, warnings = A.uncovered_mergefields(scan, manifest)
    assert not any(code in f.detail for f in faults)
    assert warnings == [], "one encoding only, so nothing is written twice"


def test_a_merge_field_in_a_table_cell_with_no_bracket_is_a_fault():
    """The counter-example that overturned the first rule.

    A remuneration row reads `$«LAB__FT_SALARY__38_HR_»` -- there is no bracket on
    that paragraph at all, so the merge field IS the slot. Treating it as an
    unresolvable duplicate left it unclaimed, and `unresolved_mergefield` then
    blocked every document the template could produce.
    """
    class _MergeField:
        paragraph_index, code = 69, '"LAB__FT_SALARY__38_HR_"'

    class _Scan:
        mergefields = [_MergeField()]
        spans = []

    faults, warnings = A.uncovered_mergefields(_Scan(), _manifest())
    assert [f.check for f in faults] == [A.UNCOVERED_MERGEFIELD]
    assert warnings == []


def test_two_merge_fields_beside_a_bracket_meaning_different_things_are_faults():
    """`This position is classified as <Grade> ... «FT_SALARY» ... «FT_TOTAL»` is
    three slots on one paragraph, not one slot three times. Structure cannot tell
    that apart from a true duplicate, which is why the model is asked."""
    class _MF:
        def __init__(self, code): self.paragraph_index, self.code = 66, code

    class _Scan:
        mergefields = [_MF('"LAB__FT_SALARY__38_HR_"'), _MF('"LAB__FT_TP_38_HR"')]
        spans = []

    manifest = _manifest(fields=[{"id": "grade", "slots": [_slot(66, "<Grade>")]}])
    faults, warnings = A.uncovered_mergefields(_Scan(), manifest)
    assert len(faults) == 2, "a bracket on the paragraph does not excuse them"
    assert warnings == []


def test_one_field_encoded_as_both_a_bracket_and_a_merge_field_is_a_warning():
    """The genuine duplicate, and the only one that is: the MODEL said these are
    one slot by putting both on one field. Filling both prints the value twice --
    `$76,800.0061,440.00`, measured on a real addendum letter -- so the template
    needs one encoding removed and a person has to choose which."""
    class _Scan:
        mergefields = []
        spans = []

    manifest = _manifest(fields=[{
        "id": "base_salary",
        "slots": [
            _slot(17, "<Base Salary>"),
            {"kind": "mergefield", "code": "PART_TIME_SALARY", "paragraph_index": 17},
        ],
    }])
    faults, warnings = A.uncovered_mergefields(_Scan(), manifest)

    assert faults == []
    assert len(warnings) == 1
    assert warnings[0]["code"] == "W-FIELD-CODE"
    assert warnings[0]["evidence"] == "base_salary"
    assert "prints twice" in warnings[0]["message"]


def test_instruction_text_nothing_removes_is_reported(template):
    """The feedback that makes `inline_branches` get used at all.

    The schema and the prompt for inline branches existed first, and the model
    still described the recruitment switch in prose and left the array empty --
    because no assertion ever said the instruction text was still there. The
    compile reported success and every document it produced was blocked.
    """
    scan = template(["Dear <Name>,", "Yours sincerely,"])

    class _Span:
        def __init__(self, p, i, colour, text):
            self.paragraph_index, self.span_index, self.color, self.text = p, i, colour, text
            self.in_hyperlink = False

    scan.spans = [
        _Span(9, 0, "black", "please return a copy of this letter to "),
        _Span(9, 1, "red", "For transactions initiated with recruitment (e.g. Change Job)"),
        _Span(9, 2, "blue", "<Contact>"),
    ]

    faults = A.surviving_instructions(scan, _manifest())
    assert [f.check for f in faults] == [A.SURVIVING_INSTRUCTION]
    assert faults[0].paragraph_index == 9
    assert "inline_branches" in faults[0].detail


def test_a_removed_instruction_is_not_reported(template):
    """Silence once the manifest deletes it -- span-scoped, so the sentence it
    sits inside survives."""
    scan = template(["x"])

    class _Span:
        def __init__(self, p, i, colour, text):
            self.paragraph_index, self.span_index, self.color, self.text = p, i, colour, text
            self.in_hyperlink = False

    scan.spans = [_Span(9, 1, "red", "For transactions initiated with recruitment")]
    manifest = _manifest()
    manifest["delete_always"] = [{"paragraph_index": 9, "span_index": 1, "scope": "span"}]

    assert A.surviving_instructions(scan, manifest) == []


def test_a_red_subheading_inside_a_block_is_not_an_instruction(template):
    """`Dates of Effect` is a heading the letter needs, coloured like an
    instruction. The fill engine makes the same allowance, for the same reason:
    deleting it silently costs the letter a heading."""
    scan = template(["x"])

    class _Span:
        def __init__(self, p, i, colour, text):
            self.paragraph_index, self.span_index, self.color, self.text = p, i, colour, text
            self.in_hyperlink = False

    scan.spans = [_Span(25, 0, "red", "Dates of Effect")]
    manifest = _manifest(blocks=[{"id": "b", "start_paragraph": 25, "end_paragraph": 27}])

    assert A.surviving_instructions(scan, manifest) == []


def _spans(*triples):
    class _Span:
        def __init__(self, p, i, colour, text):
            self.paragraph_index, self.span_index, self.color, self.text = p, i, colour, text
            self.in_hyperlink = False

    class _Scan:
        mergefields = []
        spans = [_Span(*t) for t in triples]
    return _Scan()


def test_wrapping_a_paragraph_in_blocks_does_not_silence_the_instruction_check():
    """The loophole a model actually took.

    Handed "instruction text nothing removes", it wrapped the paragraph in four
    single-paragraph blocks. That satisfied the sub-heading exemption, silenced
    this check, and produced a manifest that deletes the whole paragraph --
    "To signify your acceptance of this offer" included.
    """
    scan = _spans(
        (163, 0, "black", "please sign, date, and return a copy of this letter to "),
        (163, 1, "red", "For transactions initiated with recruitment (e.g. Change Job)"),
    )
    manifest = _manifest(blocks=[
        {"id": "a", "start_paragraph": 163, "end_paragraph": 163},
        {"id": "b", "start_paragraph": 163, "end_paragraph": 163},
    ])

    assert [f.check for f in A.surviving_instructions(scan, manifest)] == [A.SURVIVING_INSTRUCTION]


def test_a_multi_paragraph_block_still_exempts_a_red_subheading():
    """The exemption this was written for survives: `Dates of Effect` sits in a
    block covering 25..27 and is a heading the letter needs."""
    scan = _spans((25, 0, "red", "Dates of Effect"))
    manifest = _manifest(blocks=[{"id": "b", "start_paragraph": 25, "end_paragraph": 27}])

    assert A.surviving_instructions(scan, manifest) == []


def test_a_switch_modelled_as_whole_paragraph_blocks_is_reported():
    """Whole-paragraph blocks can only keep or delete the entire line, so a
    switch whose arms share a paragraph loses the sentence they sit in."""
    scan = _spans(
        (163, 0, "black", "please sign, date, and return a copy of this letter to "),
        (163, 2, "blue", "<Contact A>"),
        (163, 5, "blue", "<Contact B>"),
    )
    manifest = _manifest(blocks=[
        {"id": "with_rec", "start_paragraph": 163, "end_paragraph": 163},
        {"id": "without_rec", "start_paragraph": 163, "end_paragraph": 163},
    ])

    faults = A.paragraph_scoped_switches(scan, manifest)
    assert [f.check for f in faults] == [A.PARAGRAPH_SCOPED_SWITCH]
    assert "inline_branches" in faults[0].detail


def test_span_scoped_arms_are_the_correct_shape_and_raise_nothing():
    scan = _spans(
        (163, 0, "black", "please sign, date, and return a copy of this letter to "),
        (163, 2, "blue", "<Contact A>"),
        (163, 5, "blue", "<Contact B>"),
    )
    manifest = _manifest(blocks=[
        {"id": "with_rec", "start_paragraph": 163, "end_paragraph": 163, "start_span": 2, "end_span": 2},
        {"id": "without_rec", "start_paragraph": 163, "end_paragraph": 163, "start_span": 5, "end_span": 5},
    ])

    assert A.paragraph_scoped_switches(scan, manifest) == []


def test_one_block_on_a_paragraph_is_not_a_switch():
    """A single conditional paragraph is ordinary; only two or more arms on one
    line are the anti-pattern."""
    scan = _spans((100, 0, "black", "This clause applies only sometimes."))
    manifest = _manifest(blocks=[{"id": "solo", "start_paragraph": 100, "end_paragraph": 100}])

    assert A.paragraph_scoped_switches(scan, manifest) == []


def test_an_unmarked_parenthesised_instruction_is_caught():
    """Red is the convention, not a guarantee.

    A real contract writes "(remove this section if there is no higher duty
    allowance)" in ordinary black text beside a blue heading. Keying only on
    colour meant the compile reported success and every document was then
    blocked at generation by `instruction_text_remains` -- which reads the same
    pattern this now does, so the two agree by construction.
    """
    scan = _spans(
        (8, 0, "blue", "In case of Higher Duty Allowance"),
        (8, 1, "black", "(remove this section if there is no higher duty allowance)"),
    )
    faults = A.surviving_instructions(scan, _manifest())
    assert [f.check for f in faults] == [A.SURVIVING_INSTRUCTION]
    assert faults[0].paragraph_index == 8


def test_an_instruction_inside_the_block_it_introduces_is_still_caught():
    """The exemption is for sub-headings, and it hid this: the instruction sat
    on the first paragraph of the very block it introduces."""
    scan = _spans((8, 1, "black", "(remove this section if there is no higher duty allowance)"))
    manifest = _manifest(blocks=[{"id": "hda", "start_paragraph": 8, "end_paragraph": 9}])

    assert [f.check for f in A.surviving_instructions(scan, manifest)] == [A.SURVIVING_INSTRUCTION]


def test_ordinary_prose_in_a_block_is_not_mistaken_for_an_instruction():
    """The pattern must not fire on contract parentheticals -- "(includes
    superannuation)" is not an instruction to anybody."""
    scan = _spans(
        (30, 0, "black", "Your remuneration (includes superannuation) is reviewed annually."),
        (31, 0, "black", "The company may vary these terms (used for reference only)."),
    )
    assert A.surviving_instructions(scan, _manifest()) == []


def test_an_instruction_sharing_a_run_with_content_is_removed_on_its_own(tmp_path):
    """A run is only entirely an instruction when the markup said so, and markup
    is what a lossy conversion destroys.

    A legacy .doc converted to .docx arrives with every run black, so "To signify
    your acceptance of this offer ... (e.g. Change Job with Recruitment)" is ONE
    run holding a real sentence and an instruction. Deleting the span takes the
    sentence; leaving it prints the instruction. Neither is acceptable, and the
    loop had no third move -- it applied corrections for two rounds and the fault
    count never fell.
    """
    import dataclasses

    import docx as docx_mod

    from app.compiler.llm_compiler import assemble
    from app.generation.docx_renderer import fill_template
    from app.templates.parsers.docx_prescan import prescan

    d = docx_mod.Document()
    d.add_paragraph(
        "To signify your acceptance of this offer, return a copy to "
        "For transactions initiated with recruitment (e.g. Change Job with Recruitment) the team."
    )
    template = tmp_path / "t.docx"
    d.save(str(template))

    manifest = dataclasses.asdict(assemble(prescan(str(template)), {
        "fields": [], "conditions": [], "inline_branches": [], "scaffolding_paragraphs": [],
        "instruction_spans": [{
            "paragraph_index": 0,
            "match_text": "For transactions initiated with recruitment (e.g. Change Job with Recruitment) ",
            "reason": "instruction",
        }],
        "language": "en", "notes": [],
    }, model="test"))

    entry = manifest["delete_always"][0]
    assert "remove" in entry, "a substring must be removed as a substring, not by dropping the run"
    assert entry.get("scope") != "span"

    out = tmp_path / "out.docx"
    fill_template(str(template), str(out), manifest, {})
    text = "\n".join(p.text for p in docx_mod.Document(str(out)).paragraphs)
    assert "To signify your acceptance of this offer" in text, "the sentence must survive"
    assert "Change Job with Recruitment" not in text, "the instruction must go"


def test_a_run_that_is_wholly_an_instruction_is_still_dropped_whole(tmp_path):
    """The other half: where the markup did separate them, nothing changes."""
    import dataclasses

    import docx as docx_mod

    from app.compiler.llm_compiler import assemble
    from app.templates.parsers.docx_prescan import prescan

    d = docx_mod.Document()
    d.add_paragraph("Insert for Fixed term colleagues:")
    template = tmp_path / "t.docx"
    d.save(str(template))

    manifest = dataclasses.asdict(assemble(prescan(str(template)), {
        "fields": [], "conditions": [], "inline_branches": [], "scaffolding_paragraphs": [],
        "instruction_spans": [{"paragraph_index": 0,
                               "match_text": "Insert for Fixed term colleagues:", "reason": "i"}],
        "language": "en", "notes": [],
    }, model="test"))

    assert manifest["delete_always"][0].get("scope") == "span"


def test_a_partial_removal_silences_the_instruction_check():
    """The check must judge what the reader will see, not what the file stores.

    A `remove` entry strips a literal from a run the manifest otherwise keeps.
    Reading the run as stored meant the check never saw the removal, so an
    instruction sharing a sentence stayed a fault however many times the reviewer
    correctly removed it -- four corrections a round and the count never fell.
    Seven legacy contracts failed to compile on the same two paragraphs for this
    reason alone.
    """
    scan = _spans((160, 0, "black",
                   "To signify your acceptance of this offer, return a copy to "
                   "For transactions initiated with recruitment (e.g. Change Job with Recruitment) the team."))

    unhandled = _manifest()
    assert [f.check for f in A.surviving_instructions(scan, unhandled)] == [A.SURVIVING_INSTRUCTION]

    handled = _manifest()
    handled["delete_always"] = [{
        "paragraph_index": 160, "span_index": 0,
        "remove": ["For transactions initiated with recruitment (e.g. Change Job with Recruitment) "],
    }]
    assert A.surviving_instructions(scan, handled) == []


def test_a_removal_that_leaves_other_instruction_text_still_fires():
    """Removing one instruction from a run does not excuse a second."""
    scan = _spans((10, 0, "black",
                   "Body text (include if part time) and more (remove this section if unpaid)."))
    manifest = _manifest()
    manifest["delete_always"] = [{"paragraph_index": 10, "span_index": 0,
                                  "remove": ["(include if part time) "]}]

    assert [f.check for f in A.surviving_instructions(scan, manifest)] == [A.SURVIVING_INSTRUCTION]


def test_a_whole_line_instruction_naming_its_span_is_not_reported_as_surviving():
    """`rule_compiler` writes `{paragraph_index, span_index}` with no `scope` and
    no `remove` for an instruction that occupies a whole run. `fill_template`
    drops that paragraph -- it consults `span_index` only inside a kept block --
    and this check used to demand `span_index is None` before it would agree.

    The two producers differ in shape: `llm_compiler.assemble` always sets
    `scope` or `remove`, `rule_compiler` sets neither. `agentic_compiler` only
    ever feeds this the former, so the disagreement never fired in production --
    but had it, the loop would have been handed a fault no correction can clear,
    burned every round, and reported `llm_unconverged` for a template that
    compiles cleanly.
    """
    scan = _spans((0, 0, "black", "Hello."),
                  (1, 0, "red", "Delete this line before sending."))

    deleted = _manifest(delete_always=[{"paragraph_index": 1, "span_index": 0}])
    assert A.surviving_instructions(scan, deleted) == []

    # And it is still reported when the manifest does nothing about it at all.
    assert [f.check for f in A.surviving_instructions(scan, _manifest())] == [
        A.SURVIVING_INSTRUCTION]


def test_a_red_run_a_field_fills_is_not_reported_as_surviving():
    """The convention marks instructions red, and it also marks some
    placeholders red -- `rule_compiler` emits a `red_placeholder` slot for a
    bracket in a red span, and a real offer letter writes `<Date DD/MM/YYYY>`
    exactly that way. The value is written over it at fill time, so it is never
    printed; reporting it hands the review loop a fault no correction can clear,
    because the model cannot delete a span it correctly claimed."""
    scan = _spans((14, 0, "red", "<Date DD/MM/YYYY>"))
    claimed = _manifest(fields=[{"id": "date", "slots": [_slot(14, "<Date DD/MM/YYYY>")]}])

    assert A.surviving_instructions(scan, claimed) == []
    assert [f.check for f in A.surviving_instructions(scan, _manifest())] == [
        A.SURVIVING_INSTRUCTION]


def test_a_red_run_inside_a_hyperlink_is_not_reported_as_surviving():
    """`docx_renderer` protects hyperlink runs and never deletes one, so an
    instruction reported inside a link is a fault nothing can act on. A real
    offer letter links the word "FUSE" and colours the link red."""
    class _Span:
        def __init__(self):
            self.paragraph_index, self.span_index = 105, 1
            self.color, self.text, self.in_hyperlink = "red", "FUSE", True

    class _Scan:
        mergefields = []
        spans = [_Span()]

    assert A.surviving_instructions(_Scan(), _manifest()) == []


# ------------------------------------------- a placeholder Word split in half

class _Span:
    def __init__(self, paragraph_index, span_index, text):
        self.paragraph_index = paragraph_index
        self.span_index = span_index
        self.text = text


class _Scan:
    def __init__(self, spans):
        self.spans = spans


def test_a_placeholder_split_across_runs_is_a_warning_and_never_a_fault():
    """The Chinese master, and the compile it used to kill.

    Word splits a run wherever formatting or a spell-check mark changes, and
    around CJK text it does it constantly. The real file reads

        span 0: '\u6211\u8c28\u4ee3\u8868\u8f89\u745e\u4e2d\u56fd\uff0c\u5411\u60a8\u786e\u8ba4\u56e0<'
        span 1: 'Transaction Action Reason'
        span 2: '>\u800c\u4ea7\u751f\u7684\u4ee5\u4e0b\u53d8\u52a8...'

    which is one placeholder to a reader and none to any single span.

    It has to be *reported* -- nothing will fill it, so every document keeps the
    literal text. But it must not be a fault, because no correction can clear it:
    `_locate` finds a slot with `needle in span.text`, one span at a time, so a
    claim on a split token is dropped as unlocatable; and the fill engine
    replaces inside a single run, so it could not be filled even if it stuck.
    Measured: 9 fields merged, then three rounds each applying three corrections
    against the same two faults, ending `llm_unconverged` with an empty manifest
    on a template that had been read correctly.
    """
    from app.compiler.assertions import uncovered_placeholders

    scan = _Scan([
        _Span(15, 0, "We confirm the change arising from <"),
        _Span(15, 1, "Transaction Action Reason"),
        _Span(15, 2, "> which takes effect on "),
    ])
    faults, warnings = uncovered_placeholders(scan, {"fields": [], "conditions": []})

    assert faults == [], "a split placeholder must not drive the compile loop"
    assert [w["code"] for w in warnings] == ["W-SPLIT-PLACEHOLDER"]
    assert warnings[0]["paragraph_index"] == 15
    assert warnings[0]["evidence"] == "<Transaction Action Reason>"
    assert "split across several runs" in warnings[0]["message"]


def test_a_whole_placeholder_in_one_run_is_still_a_fault():
    """The claimable case keeps driving the loop -- that is what makes the
    compiler correct its own misses."""
    from app.compiler.assertions import UNCOVERED_PLACEHOLDER, uncovered_placeholders

    scan = _Scan([_Span(3, 0, "Dated <Date>.")])
    faults, warnings = uncovered_placeholders(scan, {"fields": [], "conditions": []})

    assert [a.check for a in faults] == [UNCOVERED_PLACEHOLDER]
    assert warnings == []


def test_a_split_placeholder_a_field_does_claim_is_not_reported():
    """Joining must not turn a covered placeholder into a fault."""
    from app.compiler.assertions import uncovered_placeholders

    scan = _Scan([
        _Span(19, 0, "Salary CNY<"),
        _Span(19, 1, "Pay Rate Monthly>"),
    ])
    manifest = {
        "fields": [{"id": "pay_rate_monthly",
                    "slots": [{"text": "<Pay Rate Monthly>", "paragraph_index": 19}]}],
        "conditions": [],
    }
    assert uncovered_placeholders(scan, manifest)[0] == []


def test_a_placeholder_answered_by_conditions_rather_than_a_field_is_not_a_fault():
    """`<Work Location/City Location>` is one placeholder offering a choice
    between two, and the right reading is two conditions keeping one branch
    each. No field ever claims it, and it is fully accounted for.

    Reporting it would be the expensive kind of wrong: the writer cannot satisfy
    the fault without inventing a field that should not exist, so the loop runs
    to its ceiling and parks a correctly-read template as `llm_unconverged`.
    """
    from app.compiler.assertions import uncovered_placeholders

    scan = _Scan([_Span(21, 0, "Your work location is <Work Location/City Location>.")])
    manifest = {
        "fields": [],
        "conditions": [
            {"id": "city_location", "expression": "city_location != ''",
             "compiled_from": "<Work Location/City Location>"},
            {"id": "work_location", "expression": "city_location == ''",
             "compiled_from": "<Work Location/City Location>"},
        ],
    }
    assert uncovered_placeholders(scan, manifest)[0] == []


def test_a_token_that_is_not_a_placeholder_is_still_ignored():
    """A ruled line for someone to hand-write on is not a field, and demanding
    one would make the loop unsatisfiable."""
    from app.compiler.assertions import uncovered_placeholders

    scan = _Scan([_Span(4, 0, "Signed: <__________________>  <>")])
    assert uncovered_placeholders(scan, {"fields": [], "conditions": []})[0] == []


def test_one_paragraph_reports_a_repeated_placeholder_once():
    from app.compiler.assertions import uncovered_placeholders

    scan = _Scan([
        _Span(7, 0, "<Name> and again <"),
        _Span(7, 1, "Name>"),
    ])
    found, _split = uncovered_placeholders(scan, {"fields": [], "conditions": []})
    assert len(found) == 1


# ------------------------------------------------- a block nothing governs

def test_a_block_no_condition_keeps_is_a_fault():
    """A block nothing governs is not a conditional section — it is ordinary
    text with a condition's name on it.

    `docx_renderer` drops a block only when a condition that *keeps* it decides
    False, so a block no condition references survives whatever the data says.
    Both halves of an IF/ELSE then print, one directly contradicting the other:

        This is a fixed-term contract commencing on 1 September 2026.
        This is a permanent contract commencing on 1 September 2026.

    Measured on a template with eight blocks and three conditions: the five
    ungoverned blocks were every positive arm. Nothing downstream caught it,
    because every other gate asks what was left over or what is missing, and this
    section is neither."""
    from app.compiler.assertions import UNGOVERNED_BLOCK, structural_faults

    manifest = {
        "fields": [],
        "blocks": [
            {"id": "blk_0_if_arm", "start_paragraph": 7, "end_paragraph": 7},
            {"id": "blk_1_else_arm", "start_paragraph": 9, "end_paragraph": 9},
        ],
        "conditions": [
            {"id": "c", "expression": "contract_type != 'Fixed Term'",
             "keeps_blocks": ["blk_1_else_arm"]},
        ],
    }
    faults = structural_faults(manifest)
    assert [f.check for f in faults] == [UNGOVERNED_BLOCK]
    assert faults[0].object_id == "blk_0_if_arm"
    assert faults[0].paragraph_index == 7


def test_a_block_every_arm_of_which_is_governed_is_silent():
    from app.compiler.assertions import structural_faults

    manifest = {
        "fields": [],
        "blocks": [
            {"id": "blk_0_if_arm", "start_paragraph": 7, "end_paragraph": 7},
            {"id": "blk_1_else_arm", "start_paragraph": 9, "end_paragraph": 9},
        ],
        "conditions": [
            {"id": "c0", "expression": "contract_type == 'Fixed Term'",
             "keeps_blocks": ["blk_0_if_arm"]},
            {"id": "c1", "expression": "contract_type != 'Fixed Term'",
             "keeps_blocks": ["blk_1_else_arm"]},
        ],
    }
    assert structural_faults(manifest) == []


def test_a_manifest_with_no_blocks_at_all_is_silent():
    """Most templates have no conditional sections. The check must not invent a
    fault for them."""
    from app.compiler.assertions import structural_faults

    assert structural_faults({"fields": [], "blocks": [], "conditions": []}) == []


def test_the_rule_compiler_cannot_produce_an_ungoverned_block():
    """`register_condition` runs for every block the deterministic path creates,
    so this fault is reachable only from the agentic path. Pinned so a refactor
    of the rule compiler cannot quietly start emitting them."""
    import inspect

    from app.compiler import rule_compiler

    src = inspect.getsource(rule_compiler)
    body = src[src.index("# ---- pass 3: resolve condition markers"):]
    # Every `blocks.append` in the marker and inline passes is followed by a
    # `register_condition` for that same block.
    assert body.count("blocks.append(") == body.count("register_condition(")
