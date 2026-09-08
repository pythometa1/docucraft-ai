"""What a CSR source file is allowed to become before it can be cited.

Every defect these guard against ends the same way: a sentence in an approved
report, with a citation on it, that the source does not support. A table split
across chunks with its column headings left on part 1; a number cut in half at
a chunk boundary; a table whose id was never recovered and so is retrieved for
the wrong section. None of them look wrong in the draft, which is why they are
asserted here rather than trusted to review.
"""

import importlib

import docx
import pytest

from app.docgen.chunking import (
    MAX_TABLE_TOKENS, OVERLAP_TOKENS, TARGET_TOKENS, chunk_extraction, estimate_tokens,
)
from app.docgen.extraction import (
    ExtractedPage, ExtractedTable, Extraction, ExtractorUnavailable, UnsupportedSource,
    extract, table_identity,
)

def _importable(*names) -> bool:
    for name in names:
        try:
            importlib.import_module(name)
            return True
        except ImportError:
            continue
    return False


HAS_PYMUPDF = _importable("pymupdf", "fitz")
HAS_STRIPRTF = _importable("striprtf.striprtf")

VOCABULARY = (
    "patients", "randomised", "received", "placebo", "treatment", "period",
    "assessments", "baseline", "compliance", "reported", "discontinued",
    "investigator", "population", "analysis", "exposure", "incidence",
)


def _paragraph(seed: int, words: int = 60) -> str:
    """A paragraph built only from VOCABULARY, so any token in a chunk that is
    not a vocabulary word is a word the chunker broke."""
    return " ".join(VOCABULARY[(seed + index) % len(VOCABULARY)] for index in range(words))


def _prose_pages(count: int, paragraphs_per_page: int = 6) -> list:
    pages = []
    for number in range(1, count + 1):
        blocks = [f"marker{number}"] + [
            _paragraph(number * 10 + index) for index in range(paragraphs_per_page)]
        pages.append(ExtractedPage(page=number, text="\n\n".join(blocks)))
    return pages


def _tokens_of(text: str) -> list:
    return [word.strip(".,;:()[]") for word in text.split()]


# ------------------------------------------------------------ table identity

def test_table_identity_reads_the_id_and_title_from_a_caption():
    """The id is how a section prompt, a citation and a QC reviewer all name
    the table; lose it and the table is retrievable only by luck."""
    assert table_identity(
        [["Age", "54.2"]], caption="Table 14.2.1: Demographics and Baseline") == (
            "14.2.1", "Demographics and Baseline")


def test_table_identity_reads_the_id_from_the_tables_own_first_rows():
    """A TLF pasted into Word carries its caption as a merged first row and no
    caption paragraph at all."""
    rows = [
        ["Table 14.1.1 Summary of Adverse Events", "", ""],
        ["System Organ Class", "Placebo", "Active"],
        ["Nervous system disorders", "4", "9"],
    ]
    assert table_identity(rows, caption=None) == ("14.1.1", "Summary of Adverse Events")


def test_table_identity_takes_the_title_from_the_row_below_a_bare_id():
    rows = [["Listing 16.2.1"], ["Discontinued Patients"], ["101", "Withdrew consent"]]
    assert table_identity(rows, caption=None) == ("16.2.1", "Discontinued Patients")


def test_table_identity_returns_no_id_rather_than_a_guessed_one():
    """A table filed under a neighbour's id is retrieved INSTEAD of that
    neighbour, so an absent id stays absent."""
    rows = [["Subject", "Age", "Sex"], ["101", "54", "F"]]
    assert table_identity(rows, caption="Demographics by Arm") == (None, "Demographics by Arm")
    # No caption: the column headings are data, not a name for the table.
    assert table_identity(rows, caption=None) == (None, None)
    # "Table 14" alone is a section cross-reference far more often than a
    # post-text table, and a wrong id makes the wrong table retrievable.
    assert table_identity([["x"]], caption="see Table 14 for details")[0] is None


# -------------------------------------------------------------- table chunks

def test_a_table_becomes_exactly_one_chunk_that_opens_with_its_header_line():
    table = ExtractedTable(
        table_id="14.1.1", title="Demographics", page=7,
        rows=[["Characteristic", "Placebo (N=50)", "Active (N=52)"],
              ["Age, mean (SD)", "54.2 (11.1)", "55.0 (10.4)"]])
    chunks = chunk_extraction(
        Extraction(pages=[], tables=[table], page_count=0), doc_type="tlf")

    assert len(chunks) == 1
    chunk = chunks[0]
    assert chunk["is_table"] is True
    assert chunk["table_id"] == "14.1.1"
    assert chunk["page"] == 7
    assert chunk["content"].startswith("Table 14.1.1 -- Demographics")
    lines = chunk["content"].splitlines()
    assert lines[1] == "Characteristic | Placebo (N=50) | Active (N=52)"
    assert lines[2] == "Age, mean (SD) | 54.2 (11.1) | 55.0 (10.4)"


def test_a_table_is_never_merged_with_the_prose_around_it():
    """Prose and a grid in one chunk is a table flattened into text, which is
    the form no reviewer can check a number against."""
    extraction = Extraction(
        pages=[ExtractedPage(page=1, text=_paragraph(1))],
        tables=[ExtractedTable(table_id="14.1.1", title="Demographics", page=1,
                               rows=[["Arm", "N"], ["Placebo", "50"]])],
        page_count=1)
    chunks = chunk_extraction(extraction, doc_type="tlf")

    assert [chunk["is_table"] for chunk in chunks] == [False, True]
    assert "Placebo | 50" not in chunks[0]["content"]
    assert "randomised" not in chunks[1]["content"]


def test_an_oversized_table_splits_by_rows_and_every_part_repeats_its_headings():
    """A half-table whose columns are unlabelled is a block of numbers that
    cannot be attributed to an arm -- so the header line and the column header
    row are repeated in every part, and no row is lost or duplicated."""
    body = [[f"S{index:04d}", "Active", "Headache and dizziness", str(index % 4)]
            for index in range(300)]
    table = ExtractedTable(
        table_id="14.3.1", title="Adverse Events by Patient", page=12,
        rows=[["Subject", "Arm", "Event", "Grade"], *body])

    parts = chunk_extraction(
        Extraction(pages=[], tables=[table], page_count=0), doc_type="tlf")

    assert len(parts) > 1
    recovered = []
    for index, part in enumerate(parts, start=1):
        lines = part["content"].splitlines()
        assert lines[0] == f"Table 14.3.1 -- Adverse Events by Patient (part {index} of {len(parts)})"
        assert lines[1] == "Subject | Arm | Event | Grade"
        assert part["table_id"] == "14.3.1"
        assert part["is_table"] is True
        assert part["token_count"] <= MAX_TABLE_TOKENS
        recovered.extend(lines[2:])

    assert recovered == [f"S{index:04d} | Active | Headache and dizziness | {index % 4}"
                         for index in range(300)]


def test_a_row_is_never_split_even_when_it_alone_exceeds_the_budget():
    """Over budget is readable; a row cut in half puts values under the wrong
    headings, which reads as a real result."""
    huge = " ".join(["narrative"] * 3000)
    table = ExtractedTable(table_id="14.9.9", title="Wide", page=1,
                           rows=[["Subject", "Comment"], ["101", huge], ["102", "short"]])
    parts = chunk_extraction(
        Extraction(pages=[], tables=[table], page_count=0), doc_type="tlf")

    assert any(huge in part["content"] for part in parts)
    assert all(part["content"].splitlines()[1] == "Subject | Comment" for part in parts)


def test_a_merged_caption_row_does_not_become_the_column_header_row():
    """Otherwise the real column names appear on part 1 only, and every later
    part's numbers sit under the table's title instead of under an arm."""
    body = [[f"S{index:04d}", "Active", "Headache", "2"] for index in range(300)]
    rows = [["Table 14.3.2 Adverse Events", "", "", ""],
            ["Subject", "Arm", "Event", "Grade"], *body]
    table_id, title = table_identity(rows, caption=None)
    parts = chunk_extraction(
        Extraction(pages=[], tables=[ExtractedTable(table_id, title, 3, rows)], page_count=0),
        doc_type="tlf")

    assert len(parts) > 1
    for part in parts:
        assert part["content"].splitlines()[1] == "Subject | Arm | Event | Grade"


def test_a_listing_or_figure_caption_row_is_dropped_like_a_table_one():
    """The header line always renders the word "Table", but ICH post-text
    sources write "Listing 16.2.1" and "Figure 14.2.3" -- and table_identity
    accepts all three. Compared verbatim the listing caption matches nothing,
    survives as the column header row, and every part after the first sits
    under "Listing 16.2.1" instead of under Subject | Day | Reason."""
    body = [[f"S{index:04d}", str(index), "Withdrew consent"] for index in range(300)]
    for caption in ("Listing 16.2.1 Discontinued Patients",
                    "Figure 14.2.3 Kaplan-Meier Survival"):
        rows = [[caption, "", ""], ["Subject", "Day", "Reason"], *body]
        table_id, title = table_identity(rows, caption=None)
        parts = chunk_extraction(
            Extraction(pages=[], tables=[ExtractedTable(table_id, title, 3, rows)],
                       page_count=0),
            doc_type="tlf")

        assert len(parts) > 1, caption
        for part in parts:
            assert part["content"].splitlines()[1] == "Subject | Day | Reason", caption


def test_an_id_row_and_a_title_row_are_both_dropped():
    """A TLF routinely prints the id and the title as two separate rows; if
    only the first is dropped the second becomes the column header row."""
    body = [[f"S{index:04d}", "Active", "2"] for index in range(300)]
    rows = [["Table 14.3.2", "", ""], ["Adverse Events", "", ""],
            ["Subject", "Arm", "Grade"], *body]
    table_id, title = table_identity(rows, caption=None)
    parts = chunk_extraction(
        Extraction(pages=[], tables=[ExtractedTable(table_id, title, 3, rows)], page_count=0),
        doc_type="tlf")

    assert len(parts) > 1
    for part in parts:
        assert part["content"].splitlines()[1] == "Subject | Arm | Grade"


def test_a_real_first_row_is_not_dropped_as_a_caption():
    """The caption rule is anchored at one end of the header line, because a
    coincidental match in the middle of two short strings costs a real row."""
    table = ExtractedTable(table_id=None, title=None, page=1,
                           rows=[["Table", "", ""], ["Arm", "N", "Pct"], ["A", "1", "2"]])
    content = chunk_extraction(
        Extraction(pages=[], tables=[table], page_count=0), doc_type="tlf")[0]["content"]

    assert content.splitlines()[1] == "Table |  | "


def test_a_cell_containing_a_pipe_cannot_fake_an_extra_column():
    """An unescaped pipe shifts every value after it one column left, and a
    placebo number read under the active arm is the worst thing this module
    can emit."""
    table = ExtractedTable(table_id="14.1.2", title="Exposure", page=1,
                           rows=[["Arm", "Dose"], ["Active", "10|20 mg"]])
    content = chunk_extraction(
        Extraction(pages=[], tables=[table], page_count=0), doc_type="tlf")[0]["content"]

    assert content.splitlines()[2] == r"Active | 10\|20 mg"


# ---------------------------------------------------------- narrative chunks

def test_narrative_chunks_respect_the_token_target_and_carry_a_section_hint():
    pages = _prose_pages(4)
    pages[0].text = "9.4.6 Blinding\n\n" + pages[0].text
    pages[3].text = "9.4.7 Prior and Concomitant Therapy\n\n" + pages[3].text
    chunks = chunk_extraction(
        Extraction(pages=pages, tables=[], page_count=4), doc_type="protocol")

    assert len(chunks) > 1
    for chunk in chunks:
        assert chunk["is_table"] is False
        assert chunk["table_id"] is None
        assert chunk["page"] in (1, 2, 3, 4)
        # The target bounds the chunk's OWN text; the carried overlap sits on
        # top of it, and the //4 estimate itself is one token coarse.
        assert chunk["token_count"] <= TARGET_TOKENS + OVERLAP_TOKENS + 2
        assert chunk["token_count"] == estimate_tokens(chunk["content"])
    assert chunks[0]["token_count"] > TARGET_TOKENS // 2, "chunks far under target"
    assert chunks[0]["section_hint"] == "9.4.6 Blinding"
    assert chunks[-1]["section_hint"] == "9.4.7 Prior and Concomitant Therapy"


def test_consecutive_narrative_chunks_overlap():
    """A result and the sentence qualifying it must reach a model together at
    least once, or the qualifier is retrievable only on its own."""
    chunks = chunk_extraction(
        Extraction(pages=_prose_pages(4), tables=[], page_count=4), doc_type="protocol")

    assert len(chunks) > 1
    for earlier, later in zip(chunks, chunks[1:]):
        carried = later["content"].split("\n\n")[0]
        assert carried, "no overlap carried into the next chunk"
        assert carried in earlier["content"]
        assert estimate_tokens(carried) <= OVERLAP_TOKENS + 1


def test_no_chunk_ever_splits_a_word():
    """A truncated number ("12.4" arriving as "12.") reads as a real value."""
    long_paragraph = " ".join(_paragraph(index, words=200) for index in range(30))
    extraction = Extraction(
        pages=[ExtractedPage(page=1, text=long_paragraph)], tables=[], page_count=1)
    chunks = chunk_extraction(extraction, doc_type="protocol")

    assert len(chunks) > 1, "the oversized paragraph was not split at all"
    for chunk in chunks:
        for word in _tokens_of(chunk["content"]):
            assert word in VOCABULARY, f"{word!r} is not a whole source word"


def test_a_page_boundary_is_hard_for_table_oriented_documents():
    """A TLF page is one table plus ITS footnotes: a chunk spanning the break
    attaches page 12's footnote to page 13's table."""
    chunks = chunk_extraction(
        Extraction(pages=_prose_pages(3), tables=[], page_count=3), doc_type="tlf")

    for chunk in chunks:
        seen = {marker for marker in ("marker1", "marker2", "marker3")
                if marker in chunk["content"]}
        assert len(seen) == 1, f"chunk spans pages {seen}"
        assert seen == {f"marker{chunk['page']}"}


def test_whitespace_only_pages_produce_no_chunks():
    """An empty chunk embeds to a vector that matches everything weakly, and
    is retrieved in place of a source that says something."""
    extraction = Extraction(
        pages=[ExtractedPage(page=1, text="   \n\n \t "), ExtractedPage(page=2, text="")],
        tables=[ExtractedTable(table_id=None, title=None, page=3, rows=[["", ""], [None, ""]])],
        page_count=2)

    assert chunk_extraction(extraction, doc_type="protocol") == []


def test_estimate_tokens_never_reports_a_costless_chunk():
    assert estimate_tokens("") == 1
    assert estimate_tokens("a" * 4000) == 1000


# ------------------------------------------------------------------- extract

def test_extract_reads_a_docx_paragraphs_and_tables(tmp_path):
    document = docx.Document()
    document.add_paragraph("The study was double-blind and placebo-controlled.")
    document.add_paragraph("Table 14.1.1 Demographics")
    table = document.add_table(rows=2, cols=3)
    for row, values in enumerate([["Characteristic", "Placebo", "Active"],
                                  ["Age, mean", "54.2", "55.0"]]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    path = tmp_path / "csr_section.docx"
    document.save(str(path))

    extraction = extract(str(path))

    assert extraction.page_count == 1
    assert "double-blind and placebo-controlled" in extraction.pages[0].text
    assert len(extraction.tables) == 1
    found = extraction.tables[0]
    # The caption is the paragraph ABOVE the table, which is where Word keeps
    # it and the only place the id can be recovered from.
    assert (found.table_id, found.title) == ("14.1.1", "Demographics")
    assert found.rows == [["Characteristic", "Placebo", "Active"],
                          ["Age, mean", "54.2", "55.0"]]


def test_extract_reads_a_csv_whatever_delimiter_it_was_exported_with(tmp_path):
    """A semicolon export read with a comma parses as one column: every number
    lands in the first cell and the table looks empty rather than broken."""
    path = tmp_path / "Table 14.2.1 Disposition.csv"
    path.write_text("Arm;Randomised;Completed\nPlacebo;50;47\nActive;52;49\n", encoding="utf-8")

    extraction = extract(str(path))

    assert len(extraction.tables) == 1
    assert extraction.tables[0].rows[1] == ["Placebo", "50", "47"]
    assert extraction.tables[0].table_id == "14.2.1"
    # No pages, not page 0: a spreadsheet has no page, and "p.0" in a citation
    # is a number a reviewer would try to look up.
    assert extraction.pages == [] and extraction.tables[0].page is None


def test_extract_reads_plain_text_as_one_page(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("9.4.6 Blinding\n\nThe randomisation code was held centrally.",
                    encoding="utf-8")

    extraction = extract(str(path))

    assert extraction.page_count == 1
    assert "randomisation code" in extraction.pages[0].text


def test_extract_refuses_a_file_type_it_cannot_read(tmp_path):
    path = tmp_path / "scan.dcm"
    path.write_bytes(b"not a document")

    with pytest.raises(UnsupportedSource) as raised:
        extract(str(path))
    assert ".dcm" in str(raised.value)


def test_a_legacy_binary_format_is_refused_with_the_way_out(tmp_path):
    """"Unsupported" alone sends the writer back to upload the same .doc."""
    path = tmp_path / "protocol.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0")

    with pytest.raises(UnsupportedSource) as raised:
        extract(str(path))
    assert ".docx" in str(raised.value)


def test_the_mime_type_rescues_a_file_stored_without_its_extension(tmp_path):
    """Uploads land in storage under a generated name, and a document that
    silently produced no chunks is discovered only when a section comes back
    thin."""
    path = tmp_path / "a3f9c2"
    path.write_text("The protocol was amended twice.", encoding="utf-8")

    extraction = extract(str(path), mime_type="text/plain; charset=utf-8")

    assert extraction.pages[0].text.startswith("The protocol")


def test_a_missing_optional_reader_names_the_package_to_install(tmp_path, monkeypatch):
    """Distinct from UnsupportedSource on purpose: this one is fixed by an
    install, so the message has to name what to install. Forced rather than
    skipped when the library happens to be present, because the path that
    only runs on machines without PyMuPDF is the path nobody ever sees fail."""
    def missing(name):
        raise ImportError(f"No module named {name!r}")

    monkeypatch.setattr(importlib, "import_module", missing)
    path = tmp_path / "protocol.pdf"
    path.write_bytes(b"%PDF-1.7\n")

    with pytest.raises(ExtractorUnavailable) as raised:
        extract(str(path))
    assert "PyMuPDF" in str(raised.value) and "pip install" in str(raised.value)


@pytest.mark.skipif(not HAS_PYMUPDF, reason="PyMuPDF is not installed here")
def test_extract_reads_every_page_of_a_pdf_and_keeps_its_captions(tmp_path):
    """Page numbers come from the PDF and nowhere else: a citation is checked
    by opening the source at the page it names."""
    pymupdf = importlib.import_module("pymupdf" if _importable("pymupdf") else "fitz")
    document = pymupdf.open()
    for number in (1, 2):
        page = document.new_page()
        page.insert_text((72, 72), f"Table 14.1.{number} Demographics")
        page.insert_text((72, 96), f"Prose belonging to page {number}.")
    path = tmp_path / "tlf.pdf"
    document.save(str(path))
    document.close()

    extraction = extract(str(path))

    assert extraction.page_count == 2
    assert [page.page for page in extraction.pages] == [1, 2]
    assert "belonging to page 2" in extraction.pages[1].text
    # Whether this build's table finder sees a ruled grid is its business; a
    # table it DOES return must carry the caption printed above it.
    for table in extraction.tables:
        assert table.table_id in ("14.1.1", "14.1.2")
        assert table.page in (1, 2)


@pytest.mark.skipif(not HAS_STRIPRTF, reason="striprtf is not installed here")
def test_extract_reads_an_rtf_as_one_page(tmp_path):
    path = tmp_path / "narrative.rtf"
    path.write_text(r"{\rtf1\ansi Patient 101 discontinued on day 14.}", encoding="utf-8")

    extraction = extract(str(path))

    assert extraction.page_count == 1
    assert "discontinued on day 14" in extraction.pages[0].text
