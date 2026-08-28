"""Control-token dialects: `[[IF x > 0]] ... [[ENDIF]]`.

A template can be fully colour-coded and still express every condition in a
syntax the rule compiler has never seen. These tests cover the deterministic
half of handling that — locating and stripping the tokens — which needs no
model and must work identically whoever compiled the manifest.
"""

import io

import docx
import pytest

from app.templates.parsers.docx_prescan import prescan
from app.generation.docx_renderer import fill_template
from app.compiler.rule_compiler import CONTROL_TOKEN_RE, find_control_markers


def _template(paragraphs) -> str:
    d = docx.Document()
    for text in paragraphs:
        d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    import tempfile, os
    path = os.path.join(tempfile.mkdtemp(), "t.docx")
    with open(path, "wb") as fh:
        fh.write(buf.getvalue())
    return path


@pytest.mark.parametrize("text,expected", [
    ("[[IF joining_bonus > 0]]", ["[[IF joining_bonus > 0]]"]),
    ("[[ELSE]]", ["[[ELSE]]"]),
    ("[[ELSE IF x = 1]]", ["[[ELSE IF x = 1]]"]),
    ("[[ENDIF]]", ["[[ENDIF]]"]),
    ("[[PROMPT: write one sentence]]", ["[[PROMPT: write one sentence]]"]),
    ("[[IF a]]<addr>[[ENDIF]]", ["[[IF a]]", "[[ENDIF]]"]),
    ("no markers here", []),
])
def test_every_control_token_is_located(text, expected):
    assert CONTROL_TOKEN_RE.findall(text) == expected


def test_markers_are_reported_with_span_coordinates():
    path = _template(["[[IF x > 0]]", "Body text.", "[[ENDIF]]"])
    markers = find_control_markers(prescan(path))
    assert {m["paragraph_index"] for m in markers} == {0, 2}
    assert all("span_index" in m and m["remove"] for m in markers)


def _manifest(scan, **kw):
    base = {"fields": [], "conditions": [], "blocks": [], "delete_always": find_control_markers(scan)}
    base.update(kw)
    return base


def _text(path):
    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs)


def test_markers_on_their_own_line_take_the_line_with_them(tmp_path):
    """Stripping the token but keeping the paragraph leaves a blank line
    wherever the template had an [[IF]] — the letter reads with gaps."""
    path = _template(["[[IF x > 0]]", "Body text.", "[[ENDIF]]"])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")
    fill_template(path, out, _manifest(scan), {"x": 1})

    d = docx.Document(out)
    assert [p.text for p in d.paragraphs] == ["Body text."]


def test_an_inline_marker_is_stripped_without_losing_its_content(tmp_path):
    """`[[IF address_line2 is not blank]]<address_line2>[[ENDIF]]` has to keep
    the address and lose both markers -- deleting the paragraph would drop a
    line of the recipient's address."""
    path = _template(["[[IF a]]Near City Centre Mall[[ENDIF]]"])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")
    fill_template(path, out, _manifest(scan), {"a": "y"})

    assert _text(out).strip() == "Near City Centre Mall"


def test_a_genuinely_blank_line_is_not_swallowed(tmp_path):
    """Only paragraphs that carried a marker are cleaned up; the template's own
    spacing has to survive."""
    path = _template(["First.", "", "[[ENDIF]]", "Second."])
    scan = prescan(path)
    out = str(tmp_path / "out.docx")
    fill_template(path, out, _manifest(scan), {})

    assert [p.text for p in docx.Document(out).paragraphs] == ["First.", "", "Second."]


def test_qa_fails_while_control_tokens_remain(tmp_path):
    """The gates must treat a leftover [[ENDIF]] the same as a leftover
    <placeholder>: both are scaffolding reaching the reader."""
    path = _template(["[[IF x]]", "Body.", "[[ENDIF]]"])
    out = str(tmp_path / "out.docx")
    # deliberately compiled WITHOUT marker handling
    result = fill_template(path, out, {"fields": [], "conditions": [], "blocks": [], "delete_always": []}, {})
    assert "[[" in _text(out)
    assert result.qa_passed is False
    assert any("control token" in n.lower() for n in result.qa_notes), result.qa_notes


# ------------------------------------------------------- condition polarity
def test_a_delete_when_true_condition_is_negated_at_compile_time():
    """Manifests are keep-when-true throughout. Templates are not: alongside
    `[[IF bonus > 0]] ... [[ENDIF]]` they write
    `[[IF bonus = 0]] Delete this entire section [[ENDIF]]`.

    Leaving the model to invert that produced a manifest whose bonus section was
    kept only for people with *no* bonuses -- so fourteen of twenty letters
    silently lost the section their data called for, and every one passed QA,
    because the gates check for leftover scaffolding and cannot know what should
    have been present. `effect` carries the polarity and the negation happens
    here, deterministically.
    """
    from app.compiler.llm_compiler import compile_manifest_llm
    from app.compiler import llm_compiler

    class _Result:
        model = "test"
        error = None
        input_tokens = output_tokens = 0
        data = {
            "fields": [],
            "conditions": [{
                "id": "hide_bonus_section",
                "expression": "joining_bonus == 0 and esop_units == 0",
                "effect": "delete",
                "start_paragraph": 1, "end_paragraph": 2,
                "compiled_from": "[[IF ...]] Delete this entire section [[ENDIF]]",
            }],
            "scaffolding_paragraphs": [], "language": "en", "notes": [],
        }

    class _Provider:
        def structured(self, **kw):
            return _Result()

        def generate(self, **kw):
            raise AssertionError("not used")

    path = _template(["Heading", "Joining bonus text", "ESOP text"])
    scan = prescan(path)
    original = llm_compiler.get_llm_provider
    llm_compiler.get_llm_provider = lambda *a, **k: _Provider()
    try:
        manifest = compile_manifest_llm(scan, ["Heading", "Joining bonus text", "ESOP text"])
    finally:
        llm_compiler.get_llm_provider = original

    expression = manifest.conditions[0]["expression"]
    assert expression == "not (joining_bonus == 0 and esop_units == 0)"

    from app.expressions.token_parser import safe_eval_condition
    assert safe_eval_condition(expression, {"joining_bonus": "100000", "esop_units": "0"}) is True
    assert safe_eval_condition(expression, {"joining_bonus": "0", "esop_units": "0"}) is False


def test_a_keep_when_true_condition_is_left_alone():
    from app.expressions.token_parser import safe_eval_condition

    assert safe_eval_condition("joining_bonus > 0", {"joining_bonus": "100000"}) is True
    assert safe_eval_condition("joining_bonus > 0", {"joining_bonus": "0"}) is False
